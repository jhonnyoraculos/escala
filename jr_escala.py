
import customtkinter as ctk
from datetime import date, datetime, timedelta
import locale
import os
import shutil
import sqlite3
from pathlib import Path
from tkinter import filedialog, messagebox
from uuid import uuid4
from PIL import Image, ImageDraw, ImageFont, ImageColor, ImageOps

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:  # openpyxl pode não estar disponível
    Workbook = None
    Font = None

from disponibilidade import configurar_base_disponibilidade, verificar_disponibilidade


try:
    locale.setlocale(locale.LC_TIME, "pt_BR.utf8")
except locale.Error:
    try:
        locale.setlocale(locale.LC_TIME, "pt_BR")
    except locale.Error:
        try:
            locale.setlocale(locale.LC_TIME, "")
        except locale.Error:
            pass

def estilo_padrao():
    ctk.set_appearance_mode("light")
    ctk.set_default_color_theme("blue")


def criar_gradiente_horizontal(largura: int, altura: int, cor_inicio: str, cor_fim: str) -> Image.Image:
    img = Image.new("RGB", (largura, altura), cor_inicio)
    r1, g1, b1 = ImageColor.getrgb(cor_inicio)
    r2, g2, b2 = ImageColor.getrgb(cor_fim)
    for x in range(largura):
        ratio = x / max(largura - 1, 1)
        r = int(r1 + (r2 - r1) * ratio)
        g = int(g1 + (g2 - g1) * ratio)
        b = int(b1 + (b2 - b1) * ratio)
        ImageDraw.Draw(img).line([(x, 0), (x, altura)], fill=(r, g, b))
    return img


def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def _rgb_to_hex(rgb: tuple[int, int, int]) -> str:
    return "#{:02X}{:02X}{:02X}".format(*rgb)


def ajustar_tom(hex_color: str, fator: float) -> str:
    r, g, b = _hex_to_rgb(hex_color)
    r = max(0, min(255, int(r * fator)))
    g = max(0, min(255, int(g * fator)))
    b = max(0, min(255, int(b * fator)))
    return _rgb_to_hex((r, g, b))


def animacao_pulse(widget, cor_flash="#D6F8D6", duracao=160):
    if widget is None:
        return
    try:
        original = widget.cget("fg_color")
        widget.configure(fg_color=cor_flash)
        widget.after(duracao, lambda: widget.configure(fg_color=original))
    except Exception:
        pass


def aplicar_hover_cartao(widget, cor_normal: str | None = None):
    if cor_normal is None:
        cor_normal = CARD_BG_DEFAULT

    def on_enter(_):
        try:
            widget.configure(
                fg_color=CARD_BG_ALT,
                border_color=ajustar_tom(COR_AZUL_CLARO, 1.15),
                border_width=1,
            )
        except ctk.TclError:
            widget.configure(fg_color=CARD_BG_ALT)

    def on_leave(_):
        try:
            widget.configure(fg_color=cor_normal, border_color=BORDER_COR, border_width=1)
        except ctk.TclError:
            widget.configure(fg_color=cor_normal)

    widget.bind("<Enter>", on_enter)
    widget.bind("<Leave>", on_leave)


estilo_padrao()


def estilizar_card(frame, pad_x=10, pad_y=10):
    try:
        frame.configure(
            fg_color=COR_PAINEL,
            corner_radius=16,
            border_color=BORDER_COR,
            border_width=1,
        )
    except ctk.TclError:
        pass
    try:
        frame.pack_configure(padx=pad_x, pady=pad_y)
    except Exception:
        try:
            frame.grid_configure(padx=pad_x, pady=pad_y)
        except Exception:
            pass


def estilizar_scrollable(frame):
    frame.configure(fg_color=COR_PAINEL, corner_radius=16, border_color=BORDER_COR, border_width=1)


def estilizar_entry(entry):
    entry.configure(
        fg_color="#FFFFFF",
        border_width=2,
        border_color=COR_CINZA,
        text_color=COR_TEXTO,
        corner_radius=14,
        font=FONT_TEXTO,
        height=36,
        placeholder_text_color="#9AA3B8",
    )


def estilizar_optionmenu(menu):
    menu.configure(
        fg_color="#FFFFFF",
        button_color=COR_AZUL_CLARO,
        button_hover_color=COR_AZUL_HOVER,
        text_color=COR_TEXTO,
        corner_radius=16,
        font=FONT_TEXTO,
    )
    try:
        menu._canvas.configure(
            highlightthickness=1,
            highlightcolor=COR_CINZA,
            highlightbackground=COR_CINZA,
        )
        menu._canvas.configure(
            borderwidth=0,
        )
    except AttributeError:
        pass


def estilizar_botao(botao, variant="primary", small=False):
    base_font = ("Segoe UI Semibold", 13 if not small else 11)
    base_height = 40 if not small else 32
    base_corner = 20 if not small else 14
    base = {
        "corner_radius": base_corner,
        "font": base_font,
        "height": base_height,
        "border_width": 0,
    }
    cor_base = COR_AZUL_CLARO
    cor_hover = COR_AZUL_HOVER
    texto = "white"
    if variant == "primary":
        cor_base = COR_AZUL_CLARO
        cor_hover = COR_AZUL_HOVER
        texto = "white"
    elif variant == "danger":
        cor_base = COR_VERMELHA
        cor_hover = COR_VERMELHA_HOVER
        texto = "white"
    elif variant == "ghost":
        base["border_width"] = 1
        base["border_color"] = COR_NEUTRO
        cor_base = "#FFFFFF"
        cor_hover = "#EFF2F7"
        texto = COR_TEXTO
    elif variant == "danger-ghost":
        base["border_width"] = 1
        base["border_color"] = COR_VERMELHA
        cor_base = "#FFFFFF"
        cor_hover = "#FFE5E5"
        texto = COR_VERMELHA
    botao.configure(fg_color=cor_base, hover_color=cor_hover, text_color=texto, **base)
    botao._base_fg = cor_base
    botao._base_border = base.get("border_width", 0)

    def _hover_enter(_event, button=botao):
        if button.cget("state") == "disabled":
            return
        border = button._base_border + 1
        button.configure(border_width=border)

    def _hover_leave(_event, button=botao):
        button.configure(border_width=button._base_border)

    botao.bind("<Enter>", _hover_enter)
    botao.bind("<Leave>", _hover_leave)

    def _animar_click(_event, button=botao):
        cor_flash = "#C7F9CC" if variant in ("primary", "ghost") else "#FFDADA"
        animacao_pulse(button, cor_flash)

    botao.bind("<ButtonRelease-1>", _animar_click, add="+")


def adicionar_divisor_horizontal(parent, pady=(8, 12), padx=20):
    linha = ctk.CTkFrame(parent, fg_color=COR_CINZA, height=2, corner_radius=0)
    linha.pack(fill="x", padx=padx, pady=pady)
    linha.pack_propagate(False)
    return linha

COR_AZUL = "#1B5FAF"
COR_AZUL_CLARO = "#1990FF"
COR_AZUL_HOVER = "#0E75D0"
COR_AZUL_GRADIENTE_FIM = "#4A8EDB"
COR_VERMELHA = "#C8102E"
COR_VERMELHA_HOVER = "#A30D26"
COR_FUNDO = "#F5F6FA"
COR_TEXTO = "#1c1c1c"
COR_PAINEL = "#FFFFFF"
COR_CINZA = "#D6DEE9"
COR_NEUTRO = "#9BA7BA"
EMPTY_STATE_COLOR = "#6B7280"
CARD_BG_DEFAULT = "#FFFFFF"
CARD_BG_ALT = "#F0F4FF"
CARD_BG_EDICAO = "#E6EEFF"
BORDER_COR = "#E0E6EF"
NAV_BG = "#0D4A92"
NAV_PILL = "#1D4C85"
NAV_PILL_HOVER = "#2B64A7"
DEFAULT_EMPTY_MESSAGE = "Nenhum registro encontrado."
DISPLAY_VAZIO = "—"
VALOR_SEM_MOTORISTA = "Sem motorista"
VALOR_SEM_AJUDANTE = "Sem ajudante"
VALOR_SEM_CAMINHAO = "Sem caminhão"
MOTORISTA_AJUDANTE_TAG = "(.mot)"
AVISO_NENHUM_COLAB = "Nenhum colaborador disponível para este dia."
OBSERVACAO_OPCOES = [
    "0",
    "ROTA 1 DIA (BATE E VOLTA)",
    "ROTA 2 DIAS",
    "ROTA 3 DIAS",
    "ROTA 4 DIAS",
    "ROTA 5 DIAS",
]
OBSERVACAO_DURACAO = {
    "0": 0,
    "ROTA 1 DIA (BATE E VOLTA)": 0,
    "ROTA 2 DIAS": 1,
    "ROTA 3 DIAS": 2,
    "ROTA 4 DIAS": 3,
    "ROTA 5 DIAS": 4,
}
OBS_MARCADORES = [
    ("Sem cor", ""),
    ("Amarelo", "#FFF59D"),
    ("Verde", "#C8E6C9"),
    ("Azul", "#BBDEFB"),
    ("Vermelho", "#FFCDD2"),
    ("Laranja", "#FFE0B2"),
]
OBS_MARCADORES_MAP = {label: cor for label, cor in OBS_MARCADORES}
AVISO_NENHUM_COLAB = "Nenhum colaborador disponível para este dia."
FONT_TITULO = ("Segoe UI Semibold", 19)
FONT_SUBTITULO = ("Segoe UI Semibold", 13)
FONT_TEXTO = ("Segoe UI", 11)
TAB_ICONS = {
    "Carregamentos": "",
    "Escala (CD)": "",
    "Folga": "",
    "Oficinas": "",
    "Rotas Semanais": "",
    "Caminhões": "",
    "Férias": "",
    "Colaboradores": "",
    "LOG": "",
}


aba_estilo_atual = ""


def atualizar_estilo_abas(tabview=None, force: bool = False):
    view = tabview or globals().get("abas")
    if view is None:
        return
    seg = getattr(view, "_segmented_button", None)
    if seg is None:
        return
    atual = view.get()
    global aba_estilo_atual
    botoes = getattr(seg, "_buttons_dict", {})
    if not force and atual == aba_estilo_atual:
        return
    if force or not aba_estilo_atual:
        for nome, botao in botoes.items():
            if nome == atual:
                botao.configure(fg_color=COR_AZUL_CLARO, text_color="white")
            else:
                botao.configure(fg_color=NAV_PILL, text_color="#F4F8FF")
    else:
        anterior = botoes.get(aba_estilo_atual)
        if anterior:
            anterior.configure(fg_color=NAV_PILL, text_color="#F4F8FF")
        atual_btn = botoes.get(atual)
        if atual_btn:
            atual_btn.configure(fg_color=COR_AZUL_CLARO, text_color="white")
    aba_estilo_atual = atual


def personalizar_barra_abas(tabview):
    seg = getattr(tabview, "_segmented_button", None)
    if seg is None:
        return
    seg.configure(
        fg_color=COR_FUNDO,
        selected_color=COR_AZUL_CLARO,
        selected_hover_color=COR_AZUL_HOVER,
        unselected_color=NAV_PILL,
        unselected_hover_color=NAV_PILL_HOVER,
        text_color="white",
        font=("Segoe UI Semibold", 14),
        corner_radius=26,
        border_width=0,
    )
    botoes = getattr(seg, "_buttons_dict", {})
    for nome, botao in botoes.items():
        texto = f"{TAB_ICONS.get(nome, '')} {nome}".strip()
        largura = max(110, min(170, 8 * len(texto) + 22))
        botao.configure(
            text=texto,
            corner_radius=22,
            border_width=0,
            width=largura,
            height=34,
        )
        try:
            botao.grid_configure(padx=6, pady=6)
        except Exception:
            pass

        def _enter(_event, b=botao, tab_nome=nome):
            if tabview.get() != tab_nome:
                b.configure(fg_color=NAV_PILL_HOVER)

        def _leave(_event, b=botao, tab_nome=nome):
            if tabview.get() != tab_nome:
                b.configure(fg_color=NAV_PILL)

        botao.bind("<Enter>", _enter)
        botao.bind("<Leave>", _leave)
    atualizar_estilo_abas(tabview, force=True)
carregamento_em_edicao_id: int | None = None
colaborador_em_edicao_id: int | None = None
colaborador_foto_rel_atual: str | None = None
colaborador_foto_rel_original: str | None = None
colaborador_foto_origem_temp: str | None = None
oficina_em_edicao_id: int | None = None
ferias_em_edicao_id: int | None = None
folga_em_edicao_id: int | None = None
btn_refresh_disponibilidade: ctk.CTkButton | None = None
btn_fullscreen: ctk.CTkButton | None = None
label_colaborador_foto_preview: ctk.CTkLabel | None = None
btn_limpar_foto_colaborador: ctk.CTkButton | None = None
preview_motorista_carreg_label: ctk.CTkLabel | None = None
preview_ajudante_carreg_label: ctk.CTkLabel | None = None
preview_oficina_motorista_label: ctk.CTkLabel | None = None
preview_cd_motorista_label: ctk.CTkLabel | None = None
preview_cd_ajudante_label: ctk.CTkLabel | None = None
preview_folga_colaborador_label: ctk.CTkLabel | None = None
preview_ferias_colaborador_label: ctk.CTkLabel | None = None
label_carreg_dia_semana: ctk.CTkLabel | None = None
btn_motorista_ajudante: ctk.CTkButton | None = None
permitir_motorista_ajudante: bool = False
combo_carregamento_dia: ctk.CTkOptionMenu | None = None
carregamento_dia_map: dict[str, dict] = {}
carregamento_dia_selecionado_id: int | None = None
carregamentos_revisados: set[int] = set()
CARREGAMENTO_DIA_PLACEHOLDER = "Selecionar carregamento"
CARREGAMENTO_DIA_SEM_DATA = "Informe uma data"
CARREGAMENTO_DIA_VAZIO = "Nenhum carregamento"
CARREGAMENTO_DIA_MARCADOR = "[VISTO]"
colaborador_option_map: dict[str, int] = {}
folga_colaborador_map: dict[str, int] = {}
LOG_STATUS_OPCOES = ["Todos", "Em andamento", "Finalizados"]
log_registros_cache: list[dict] = []
log_motorista_map: dict[str, int | None] = {}
log_placa_map: dict[str, str | None] = {}
DIAS_SEMANA = [
    ("segunda", "Segunda-feira"),
    ("terça", "Terça-feira"),
    ("quarta", "Quarta-feira"),
    ("quinta", "Quinta-feira"),
    ("sexta", "Sexta-feira"),
    ("sábado", "Sábado"),
    ("domingo", "Domingo"),
]
DIAS_SEMANA_LABELS = [label for _, label in DIAS_SEMANA]
DIA_LABEL_PARA_CHAVE = {label: chave for chave, label in DIAS_SEMANA}


def combinar_observacoes(observacao_padrao: str | None, observacao_extra: str | None) -> str:
    padrao = (observacao_padrao or "").strip()
    extra = (observacao_extra or "").strip()
    if padrao and extra:
        return f"{padrao} - {extra}"
    if padrao:
        return padrao
    if extra:
        return extra
    return DISPLAY_VAZIO


def data_iso_para_extenso(data_iso: str | None) -> str:
    if not data_iso:
        return DISPLAY_VAZIO
    try:
        data_dt = datetime.strptime(data_iso, "%Y-%m-%d").date()
    except ValueError:
        return data_iso
    nome_dia = DIAS_EXTENSO[data_dt.weekday()]
    nome_mes = MESES_EXTENSO[data_dt.month - 1]
    return f"{nome_dia}, {data_dt.day:02d} de {nome_mes} de {data_dt.year}"


def data_iso_para_br(data_iso: str | None) -> str:
    if not data_iso:
        return DISPLAY_VAZIO
    try:
        return datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return data_iso


def data_br_para_iso(data_br: str | None) -> str | None:
    if not data_br:
        return None
    try:
        return datetime.strptime(data_br.strip(), "%d/%m/%Y").strftime("%Y-%m-%d")
    except ValueError:
        return None


def data_iso_para_br_entrada(data_iso: str | None) -> str:
    if not data_iso:
        return ""
    try:
        return datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return data_iso


def dias_para_texto(valor: int) -> str:
    sufixo = "dia" if abs(valor) == 1 else "dias"
    return f"{valor} {sufixo}"


def calcular_data_saida_padrao(data_iso: str | None) -> str | None:
    if not data_iso:
        return None
    try:
        base = datetime.strptime(data_iso, "%Y-%m-%d").date()
    except ValueError:
        return None
    dias = 3 if base.weekday() == 4 else 1
    return (base + timedelta(days=dias)).isoformat()


def calcular_data_saida_carregamento(data_iso: str | None) -> str | None:
    if not data_iso:
        return None
    try:
        base = datetime.strptime(data_iso, "%Y-%m-%d").date()
    except ValueError:
        return None
    dias = 3 if base.weekday() == 4 else 1
    return (base + timedelta(days=dias)).isoformat()


def normalizar_cor_hex(cor: str | None) -> str | None:
    if not cor:
        return None
    cor = cor.strip()
    if not cor:
        return None
    if not cor.startswith("#"):
        return cor.upper()
    if len(cor) == 4:
        r, g, b = cor[1], cor[2], cor[3]
        cor = f"#{r}{r}{g}{g}{b}{b}"
    return cor.upper()


def label_cor_observacao(cor_hex: str | None) -> str:
    cor_norm = normalizar_cor_hex(cor_hex)
    for label, cor in OBS_MARCADORES:
        if normalizar_cor_hex(cor) == cor_norm:
            return label
    return OBS_MARCADORES[0][0]


def ajustar_cor_marcador(cor_hex: str | None) -> str | None:
    cor_norm = normalizar_cor_hex(cor_hex)
    if not cor_norm:
        return None
    try:
        r, g, b = ImageColor.getrgb(cor_norm)
    except ValueError:
        return cor_norm
    mix = lambda v: int(v * 0.9 + 255 * 0.1)
    r2, g2, b2 = mix(r), mix(g), mix(b)
    return f"#{r2:02X}{g2:02X}{b2:02X}"
DIA_CHAVE_PARA_LABEL = {chave: label for chave, label in DIAS_SEMANA}

DB_PATH = Path(r"jr_escala.db")
RELATORIOS_DIR = Path("relatorios")
FOTOS_COLAB_DIR = Path("fotos_colaboradores")
BACKUP_DIR = Path("backups")
BACKUP_PREFIX = "jr_escala_backup"
MAX_BACKUPS = 15
LOGO_PATH = Path("logo-jr.png")
FOTO_EXTENSOES_SUPORTADAS = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp")
FOTO_PREVIEW_PADRAO = (96, 96)
try:
    IMAGE_RESAMPLE = Image.Resampling.LANCZOS
except AttributeError:
    IMAGE_RESAMPLE = Image.LANCZOS
FUNCOES = ["Selecione...", "Motorista", "Ajudante"]
AVISO_CAMPOS_VAZIOS = "Um dos campos está vazio. Deseja continuar mesmo assim?"
label_rotas_auto_info: ctk.CTkLabel | None = None
COLABORADORES_CACHE: dict[int, dict] = {}
FOTO_CACHE: dict[tuple[str, tuple[int, int], bool], ctk.CTkImage] = {}

configurar_base_disponibilidade(DB_PATH, OBSERVACAO_DURACAO)

try:
    if DB_PATH.parent != Path("."):
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
except OSError:
    pass


def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn


def listar_backups():
    if not BACKUP_DIR.exists():
        return []
    backups = list(BACKUP_DIR.glob(f"{BACKUP_PREFIX}_*.db"))
    backups.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return backups


def limpar_backups_antigos():
    backups = listar_backups()
    for antigo in backups[MAX_BACKUPS:]:
        try:
            antigo.unlink()
        except OSError:
            pass


def criar_backup_automatico():
    if not DB_PATH.exists():
        return
    try:
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    except OSError:
        return
    backups = listar_backups()
    if backups:
        ultima = backups[0]
        data_ultima = ultima.stem.replace(f"{BACKUP_PREFIX}_", "")[:8]
        if data_ultima == datetime.now().strftime("%Y%m%d"):
            return
    destino = BACKUP_DIR / f"{BACKUP_PREFIX}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    try:
        shutil.copy2(DB_PATH, destino)
    except OSError:
        return
    limpar_backups_antigos()


def obter_ultimo_backup() -> Path | None:
    backups = listar_backups()
    return backups[0] if backups else None


def restaurar_ultimo_backup():
    backup = obter_ultimo_backup()
    if not backup:
        messagebox.showwarning("Backup", "Nenhum backup encontrado.")
        return
    if not messagebox.askyesno(
        "Restaurar backup",
        "Isso vai substituir o banco atual pelo ultimo backup. Deseja continuar?",
    ):
        return
    try:
        shutil.copy2(backup, DB_PATH)
    except OSError as exc:
        messagebox.showerror("Backup", f"Falha ao restaurar backup.\n{exc}")
        return

    COLABORADORES_CACHE.clear()
    FOTO_CACHE.clear()
    carregamentos_revisados.clear()
    limpar_form_carregamento()
    limpar_form_oficina()
    limpar_form_escala_cd()
    limpar_form_folga()
    limpar_form_rotas_semanais()
    limpar_form_colaborador()
    limpar_caminhao_form()
    limpar_form_ferias()
    refresh_carregamentos_ui()
    refresh_caminhoes_ui()
    refresh_colaboradores_ui()
    refresh_escala_cd_dropdowns()
    refresh_escala_cd_lista()
    refresh_oficinas_ui()
    atualizar_lista_folgas()
    atualizar_lista_ferias()
    atualizar_lista_rotas_semanais()
    aplicar_filtros_log()
    recarregar_disponibilidade()
    messagebox.showinfo("Backup", "Backup restaurado com sucesso.")


def mostrar_msg_lista(frame, texto: str = DEFAULT_EMPTY_MESSAGE):
    ctk.CTkLabel(
        frame,
        text=texto,
        text_color=EMPTY_STATE_COLOR,
        font=("Inter", 14, "italic"),
    ).pack(pady=12, padx=10)


def carregar_fonte(tamanho: int, bold: bool = False):
    candidatos: list[str] = []
    win_fonts = Path(os.environ.get("WINDIR", r"C:\\Windows")) / "Fonts"

    def registrar_fontes(*nomes: str):
        for nome in nomes:
            caminho_windows = win_fonts / nome
            for candidato in (caminho_windows, Path(nome)):
                candidato_str = str(candidato)
                if candidato_str not in candidatos:
                    candidatos.append(candidato_str)

    if bold:
        registrar_fontes("arialbd.ttf", "ARIALBD.TTF", "segoeuib.ttf", "SEGOEUIB.TTF")
    registrar_fontes("arial.ttf", "ARIAL.TTF", "segoeui.ttf", "SEGOEUI.TTF", "calibri.ttf", "CALIBRI.TTF")

    for caminho in candidatos:
        try:
            return ImageFont.truetype(caminho, tamanho)
        except OSError:
            continue
    return ImageFont.load_default()


def medir_texto(draw: ImageDraw.ImageDraw, texto: str, fonte: ImageFont.ImageFont):
    if hasattr(draw, "textbbox"):
        bbox = draw.textbbox((0, 0), texto, font=fonte)
        return bbox[2] - bbox[0], bbox[3] - bbox[1]
    if hasattr(draw, "textsize"):
        return draw.textsize(texto, font=fonte)
    if hasattr(fonte, "getbbox"):
        bbox = fonte.getbbox(texto)
        return bbox[2] - bbox[0], bbox[3] - bbox[1]
    largura = len(texto) * fonte.size * 0.6
    return largura, fonte.size


def atualizar_mensagem_rotas_auto(texto: str, cor: str = COR_TEXTO):
    if label_rotas_auto_info is not None:
        label_rotas_auto_info.configure(text=texto, text_color=cor)


def normalizar_dia_semana(valor: str | None) -> str:
    if not valor:
        return DIAS_SEMANA[0][0]
    valor = valor.strip().lower()
    for chave, label in DIAS_SEMANA:
        if valor.startswith(chave[:3]) or valor.startswith(chave) or valor.startswith(label.lower()[:3]):
            return chave
    return DIAS_SEMANA[0][0]


def obter_dia_semana_por_data(data_iso: str) -> str:
    try:
        dt = datetime.strptime(data_iso, "%Y-%m-%d")
    except ValueError:
        return DIAS_SEMANA[0][0]
    idx = dt.weekday()
    idx = idx % len(DIAS_SEMANA)
    return DIAS_SEMANA[idx][0]


def confirmar_campos_vazios(campos: list[str]) -> bool:
    if not campos:
        return True
    mensagem = AVISO_CAMPOS_VAZIOS
    if campos:
        mensagem += "\nCampos: " + ", ".join(campos)
    return messagebox.askyesno("Campos vazios", mensagem)


def exportar_relatorio_imagem(
    aba_nome: str,
    titulo: str,
    colunas: list[str],
    linhas: list[list[str]],
    data_referencia: str,
    subtitulos: list[str] | None = None,
    col_widths: list[float] | None = None,
    highlight_col: int | None = None,
    highlight_colors: list[str | None] | None = None,
):
    if not linhas:
        linhas = [[DISPLAY_VAZIO for _ in colunas]]

    largura = 1920
    altura = 1080
    header_altura = 60
    rodape_altura = 30
    margem = 35
    tabela_top = header_altura + 15
    tabela_bottom = altura - rodape_altura - 15
    largura_util = largura - 2 * margem

    if not col_widths:
        col_widths = [1 / len(colunas)] * len(colunas)

    imagem = Image.new("RGB", (largura, altura), "#E4EBF7")
    draw = ImageDraw.Draw(imagem)

    font_logo = carregar_fonte(24, bold=True)
    font_header = carregar_fonte(16, bold=True)
    font_sub = carregar_fonte(12, bold=False)
    font_table_header = carregar_fonte(12, bold=True)
    font_table = carregar_fonte(10, bold=False)

    # Header
    draw.rectangle([0, 0, largura, header_altura], fill=COR_AZUL)
    draw.rectangle([0, 0, 360, header_altura], fill=COR_VERMELHA)
    if LOGO_PATH.exists():
        try:
            with Image.open(LOGO_PATH) as logo_img:
                logo_rel = logo_img.convert("RGBA").resize((35, 35))
            imagem.paste(logo_rel, (25, 15), logo_rel)
        except OSError:
            draw.text((50, 30), "JR", fill="white", font=font_logo)
    else:
        draw.text((50, 30), "JR", fill="white", font=font_logo)
    draw.text((100, 15), titulo.upper(), fill="white", font=font_header)
    draw.text(
        (100, 30),
        f"Gerado em: {datetime.now().strftime('%d/%m/%Y')}",
        fill="white",
        font=font_sub,
    )
    draw.text(
        (100, 41),
        f"Base: {datetime.strptime(data_referencia, '%Y-%m-%d').strftime('%d/%m/%Y')}",
        fill="white",
        font=font_sub,
    )
    if subtitulos:
        offset = 41
        for linha in subtitulos:
            draw.text((1000, offset), linha, fill="white", font=font_sub)
            offset += 15

    def quebrar_texto(texto: str, largura_max: float):
        if not texto:
            return [DISPLAY_VAZIO]
        palavras = texto.split()
        linhas_q = []
        atual = ""
        for palavra in palavras:
            tentativa = (atual + " " + palavra).strip()
            largura_t, _ = medir_texto(draw, tentativa, font_table)
            if largura_t <= largura_max - 20:
                atual = tentativa
            else:
                if atual:
                    linhas_q.append(atual)
                atual = palavra
        if atual:
            linhas_q.append(atual)
        return linhas_q or [DISPLAY_VAZIO]

    y = tabela_top
    header_altura_linha = 40
    x = margem
    for idx, coluna in enumerate(colunas):
        largura_coluna = int(col_widths[idx] * largura_util)
        draw.rectangle([x, y, x + largura_coluna, y + header_altura_linha], fill="#1A3E78")
        draw.text((x + 10, y + 20), coluna.upper(), fill="white", font=font_table_header)
        x += largura_coluna
    y += header_altura_linha

    linha_cor_1 = "#FFFFFF"
    linha_cor_2 = "#F5F6FA"
    linha_altura_base = 44

    for linha_idx, linha in enumerate(linhas):
        x = margem
        altura_linha = linha_altura_base
        celulas_processadas = []
        for col_idx, valor in enumerate(linha):
            largura_coluna = int(col_widths[col_idx] * largura_util)
            linhas_texto = quebrar_texto(str(valor), largura_coluna)
            altura_linha = max(altura_linha, 24 * len(linhas_texto) + 16)
            celulas_processadas.append((largura_coluna, linhas_texto))

        if y + altura_linha > tabela_bottom:
            break

        fundo = linha_cor_1 if linha_idx % 2 == 0 else linha_cor_2

        for col_idx, (largura_coluna, linhas_texto) in enumerate(celulas_processadas):
            bg = fundo
            if highlight_col is not None and col_idx == highlight_col:
                cor_personalizada = None
                if highlight_colors and linha_idx < len(highlight_colors):
                    cor_personalizada = ajustar_cor_marcador(highlight_colors[linha_idx])
                if cor_personalizada:
                    bg = cor_personalizada
                elif (linhas_texto[0] or "").strip().upper() not in ("0", DISPLAY_VAZIO):
                    bg = "#FFF2B2"
            draw.rectangle([x, y, x + largura_coluna, y + altura_linha], fill=bg, outline="#C1C7D6")
            text_y = y + 8
            for linha_texto in linhas_texto:
                draw.text((x + 10, text_y), linha_texto, fill=COR_TEXTO, font=font_table)
                text_y += 24
            x += largura_coluna
        y += altura_linha

    draw.text(
        (margem, altura - rodape_altura),
        "© JR Ferragens e Madeiras | Sistema JR Escala",
        fill=COR_TEXTO,
        font=font_sub,
    )

    nome_arquivo = f"relatorio_JR_{aba_nome}_{data_referencia}.png"
    caminho = RELATORIOS_DIR / nome_arquivo
    imagem.save(caminho)
    messagebox.showinfo(
        "Relatório gerado",
        f"Relatório salvo em: {caminho}",
    )


def init_db():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA foreign_keys = ON;")
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS colaboradores (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL,
                funcao TEXT NOT NULL,
                observacao TEXT DEFAULT '',
                foto TEXT DEFAULT '',
                ativo INTEGER NOT NULL DEFAULT 1
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS folgas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                data_fim TEXT,
                colaborador_id INTEGER NOT NULL,
                observacao_padrao TEXT,
                observacao_extra TEXT,
                observacao_cor TEXT,
                FOREIGN KEY(colaborador_id) REFERENCES colaboradores(id),
                UNIQUE(data, colaborador_id)
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS ferias (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                colaborador_id INTEGER NOT NULL,
                data_inicio TEXT NOT NULL,
                data_fim TEXT NOT NULL,
                observacao TEXT,
                FOREIGN KEY(colaborador_id) REFERENCES colaboradores(id)
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS carregamentos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                data_saida TEXT,
                rota TEXT NOT NULL,
                placa TEXT,
                motorista_id INTEGER,
                ajudante_id INTEGER,
                observacao TEXT,
                observacao_extra TEXT,
                observacao_cor TEXT,
                FOREIGN KEY(motorista_id) REFERENCES colaboradores(id),
                FOREIGN KEY(ajudante_id) REFERENCES colaboradores(id),
                UNIQUE(data, rota, placa)
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS oficinas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                motorista_id INTEGER,
                placa TEXT NOT NULL,
                observacao TEXT,
                observacao_extra TEXT,
                data_saida TEXT,
                observacao_cor TEXT,
                FOREIGN KEY(motorista_id) REFERENCES colaboradores(id),
                UNIQUE(data, placa)
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS caminhoes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                placa TEXT UNIQUE NOT NULL,
                modelo TEXT,
                observacao TEXT,
                ativo INTEGER NOT NULL DEFAULT 1
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS bloqueios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                colaborador_id INTEGER NOT NULL,
                data_inicio TEXT NOT NULL,
                data_fim TEXT NOT NULL,
                motivo TEXT,
                carregamento_id INTEGER,
                FOREIGN KEY(colaborador_id) REFERENCES colaboradores(id)
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS rotas_semanais (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dia_semana TEXT NOT NULL,
                rota TEXT NOT NULL,
                destino TEXT,
                observacao TEXT
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS escala_cd (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                motorista_id INTEGER,
                ajudante_id INTEGER,
                observacao TEXT,
                FOREIGN KEY(motorista_id) REFERENCES colaboradores(id),
                FOREIGN KEY(ajudante_id) REFERENCES colaboradores(id)
            );
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS ajustes_rotas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                carregamento_id INTEGER NOT NULL,
                data_ajuste TEXT NOT NULL,
                duracao_anterior INTEGER NOT NULL,
                duracao_nova INTEGER NOT NULL,
                observacao_ajuste TEXT,
                FOREIGN KEY(carregamento_id) REFERENCES carregamentos(id)
            );
            """
        )
        conn.commit()


def atualizar_carregamentos_para_campos_opcionais():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(carregamentos);")
        colunas = cur.fetchall()
        precisa_migrar = False
        for coluna in colunas:
            nome = coluna[1]
            notnull = coluna[3]
            if nome in ("motorista_id", "ajudante_id", "placa") and notnull == 1:
                precisa_migrar = True
                break
        if not precisa_migrar:
            return

        cur.execute("PRAGMA foreign_keys = OFF;")
        cur.execute("ALTER TABLE carregamentos RENAME TO carregamentos_old;")
        cur.execute(
            """
            CREATE TABLE carregamentos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                data_saida TEXT,
                rota TEXT NOT NULL,
                placa TEXT,
                motorista_id INTEGER,
                ajudante_id INTEGER,
                observacao TEXT,
                FOREIGN KEY(motorista_id) REFERENCES colaboradores(id),
                FOREIGN KEY(ajudante_id) REFERENCES colaboradores(id),
                UNIQUE(data, rota, placa)
            );
            """
        )
        cur.execute(
            """
            INSERT INTO carregamentos (id, data, data_saida, rota, placa, motorista_id, ajudante_id, observacao)
            SELECT id, data, NULL, rota, placa, motorista_id, ajudante_id, observacao
            FROM carregamentos_old;
            """
        )
        cur.execute("DROP TABLE carregamentos_old;")
        cur.execute("PRAGMA foreign_keys = ON;")
        conn.commit()


def atualizar_oficinas_para_motorista_opcional():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(oficinas);")
        colunas = cur.fetchall()
        precisa = False
        for coluna in colunas:
            nome = coluna[1]
            notnull = coluna[3]
            if nome == "motorista_id" and notnull == 1:
                precisa = True
                break
        if not precisa:
            return
        cur.execute("PRAGMA foreign_keys = OFF;")
        cur.execute("ALTER TABLE oficinas RENAME TO oficinas_old;")
        cur.execute(
            """
            CREATE TABLE oficinas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                motorista_id INTEGER,
                placa TEXT NOT NULL,
                observacao TEXT,
                FOREIGN KEY(motorista_id) REFERENCES colaboradores(id),
                UNIQUE(data, placa)
            );
            """
        )
        cur.execute(
            """
            INSERT INTO oficinas (id, data, motorista_id, placa, observacao)
            SELECT id, data, motorista_id, placa, observacao
            FROM oficinas_old;
            """
        )
        cur.execute("DROP TABLE oficinas_old;")
        cur.execute("PRAGMA foreign_keys = ON;")
        conn.commit()


def garantir_coluna_observacao_colaboradores():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(colaboradores);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "observacao" for coluna in colunas)
        if not possui:
            cur.execute(
                "ALTER TABLE colaboradores ADD COLUMN observacao TEXT DEFAULT '';"
            )
            conn.commit()


def garantir_coluna_foto_colaboradores():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(colaboradores);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "foto" for coluna in colunas)
        if not possui:
            cur.execute("ALTER TABLE colaboradores ADD COLUMN foto TEXT DEFAULT '';")
            conn.commit()


def garantir_coluna_observacao_extra_carregamentos():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(carregamentos);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "observacao_extra" for coluna in colunas)
        if not possui:
            cur.execute("ALTER TABLE carregamentos ADD COLUMN observacao_extra TEXT;")
            conn.commit()


def garantir_coluna_observacao_cor_carregamentos():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(carregamentos);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "observacao_cor" for coluna in colunas)
        if not possui:
            cur.execute("ALTER TABLE carregamentos ADD COLUMN observacao_cor TEXT;")
            conn.commit()


def garantir_coluna_data_saida_carregamentos():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(carregamentos);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "data_saida" for coluna in colunas)
        if not possui:
            cur.execute("ALTER TABLE carregamentos ADD COLUMN data_saida TEXT;")
            conn.commit()


def garantir_colunas_oficinas_complementares():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(oficinas);")
        colunas = cur.fetchall()
        nomes = {coluna[1] for coluna in colunas}
        if "observacao_extra" not in nomes:
            cur.execute("ALTER TABLE oficinas ADD COLUMN observacao_extra TEXT;")
        if "data_saida" not in nomes:
            cur.execute("ALTER TABLE oficinas ADD COLUMN data_saida TEXT;")
        if "observacao_cor" not in nomes:
            cur.execute("ALTER TABLE oficinas ADD COLUMN observacao_cor TEXT;")
        conn.commit()


def garantir_colunas_folgas_complementares():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(folgas);")
        colunas = cur.fetchall()
        nomes = {coluna[1] for coluna in colunas}
        if "data_fim" not in nomes:
            cur.execute("ALTER TABLE folgas ADD COLUMN data_fim TEXT;")
        if "observacao_padrao" not in nomes:
            cur.execute("ALTER TABLE folgas ADD COLUMN observacao_padrao TEXT;")
        if "observacao_extra" not in nomes:
            cur.execute("ALTER TABLE folgas ADD COLUMN observacao_extra TEXT;")
        if "observacao_cor" not in nomes:
            cur.execute("ALTER TABLE folgas ADD COLUMN observacao_cor TEXT;")
        conn.commit()


def garantir_coluna_observacao_cor_oficinas():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(oficinas);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "observacao_cor" for coluna in colunas)
        if not possui:
            cur.execute("ALTER TABLE oficinas ADD COLUMN observacao_cor TEXT;")
            conn.commit()




def garantir_coluna_carregamento_bloqueios():
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(bloqueios);")
        colunas = cur.fetchall()
        possui = any(coluna[1] == "carregamento_id" for coluna in colunas)
        if not possui:
            cur.execute(
                "ALTER TABLE bloqueios ADD COLUMN carregamento_id INTEGER;"
            )
            conn.commit()


def limpar_bloqueios_expirados():
    hoje = date.today().isoformat()
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM bloqueios WHERE data_fim < ?;", (hoje,))
        conn.commit()


def _atualizar_cache_colaboradores(colaboradores: list[dict]):
    for colaborador in colaboradores:
        COLABORADORES_CACHE[colaborador["id"]] = colaborador


def obter_colaborador_por_id(colaborador_id: int | None) -> dict | None:
    if not colaborador_id:
        return None
    if colaborador_id in COLABORADORES_CACHE:
        return COLABORADORES_CACHE[colaborador_id]
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            "SELECT id, nome, funcao, observacao, foto FROM colaboradores WHERE id = ?;",
            (colaborador_id,),
        )
        row = cur.fetchone()
        if row:
            dados = dict(row)
            COLABORADORES_CACHE[colaborador_id] = dados
            return dados
    return None


def formatar_ajudante_nome(nome: str, colaborador_id: int | None) -> str:
    if not colaborador_id:
        return nome
    if not nome or nome == DISPLAY_VAZIO:
        return nome or DISPLAY_VAZIO
    dados = obter_colaborador_por_id(colaborador_id)
    if not dados:
        return nome
    funcao = (dados.get("funcao") or "").lower()
    if funcao.startswith("motor") and MOTORISTA_AJUDANTE_TAG not in nome:
        return f"{nome} {MOTORISTA_AJUDANTE_TAG}"
    return nome


def caminho_absoluto_foto(rel_path: str | None) -> Path | None:
    if not rel_path:
        return None
    caminho = Path(rel_path)
    if not caminho.is_absolute():
        caminho = Path.cwd() / caminho
    return caminho if caminho.exists() else None


def _cache_key(caminho: Path, size: tuple[int, int], circular: bool) -> tuple[str, tuple[int, int], bool]:
    return (str(caminho), size, circular)


def carregar_imagem_em_cache(
    caminho: Path, size: tuple[int, int], circular: bool = False
) -> ctk.CTkImage | None:
    chave = _cache_key(caminho, size, circular)
    if chave in FOTO_CACHE:
        return FOTO_CACHE[chave]
    try:
        with Image.open(caminho) as pil_img:
            imagem = pil_img.copy()
    except OSError:
        return None
    if circular:
        imagem = imagem.convert("RGBA")
        imagem = ImageOps.fit(imagem, size, IMAGE_RESAMPLE)
        mask = Image.new("L", size, 0)
        draw = ImageDraw.Draw(mask)
        draw.ellipse((0, 0, size[0], size[1]), fill=255)
        imagem.putalpha(mask)
    img = ctk.CTkImage(light_image=imagem, dark_image=imagem, size=size)
    FOTO_CACHE[chave] = img
    return img


def obter_imagem_colaborador_por_id(
    colaborador_id: int | None,
    size: tuple[int, int] = FOTO_PREVIEW_PADRAO,
    circular: bool = False,
) -> ctk.CTkImage | None:
    if not colaborador_id:
        return None
    dados = obter_colaborador_por_id(colaborador_id)
    if not dados:
        return None
    caminho = caminho_absoluto_foto(dados.get("foto"))
    if not caminho:
        return None
    return carregar_imagem_em_cache(caminho, size, circular=circular)


def salvar_foto_colaborador_local(origem: str) -> str | None:
    if not origem:
        return None
    FOTOS_COLAB_DIR.mkdir(parents=True, exist_ok=True)
    destino = FOTOS_COLAB_DIR / f"{uuid4().hex}.png"
    try:
        with Image.open(origem) as img:
            img = img.convert("RGBA")
            img.save(destino, format="PNG")
    except OSError:
        try:
            shutil.copy2(origem, destino)
        except OSError:
            return None
    return str(destino)


def limpar_cache_foto(rel_path: str | None):
    if not rel_path:
        return
    caminho = caminho_absoluto_foto(rel_path)
    if not caminho:
        return
    alvo = str(caminho)
    for chave in list(FOTO_CACHE.keys()):
        if chave[0] == alvo:
            FOTO_CACHE.pop(chave, None)


def remover_arquivo_foto(rel_path: str | None):
    caminho = caminho_absoluto_foto(rel_path)
    if caminho and caminho.exists():
        try:
            caminho.unlink()
        except OSError:
            pass
    limpar_cache_foto(rel_path)

def add_colaborador(nome: str, funcao: str, observacao: str = "", foto: str | None = None) -> int:
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO colaboradores (nome, funcao, observacao, foto, ativo) VALUES (?, ?, ?, ?, 1);",
            (nome.strip(), funcao.strip(), observacao.strip(), foto or ""),
        )
        conn.commit()
        return cur.lastrowid


def listar_colaboradores():
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            "SELECT id, nome, funcao, observacao, foto FROM colaboradores WHERE ativo = 1 ORDER BY nome;"
        )
        registros = []
        for row in cur.fetchall():
            dados = dict(row)
            dados["foto"] = dados.get("foto") or None
            registros.append(dados)
    _atualizar_cache_colaboradores(registros)
    return registros


def atualizar_colaborador(
    colaborador_id: int, nome: str, funcao: str, observacao: str, foto: str | None = None
):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE colaboradores
            SET nome = ?, funcao = ?, observacao = ?, foto = ?
            WHERE id = ?;
            """,
            (nome.strip(), funcao.strip(), observacao.strip(), foto or "", colaborador_id),
        )
        conn.commit()
    if colaborador_id in COLABORADORES_CACHE:
        COLABORADORES_CACHE[colaborador_id].update(
            {
                "nome": nome.strip(),
                "funcao": funcao.strip(),
                "observacao": observacao.strip(),
                "foto": foto or None,
            }
        )


def desativar_colaborador(colaborador_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "UPDATE colaboradores SET ativo = 0 WHERE id = ?;",
            (colaborador_id,),
        )
        conn.commit()


def listar_colaboradores_por_funcao(
    funcao: str, data_iso: str | None = None, ignorar: dict[str, int] | None = None
):
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, nome, funcao, observacao, foto
            FROM colaboradores
            WHERE ativo = 1 AND funcao = ?
            ORDER BY nome;
            """,
            (funcao,),
        )
        colaboradores = []
        for row in cur.fetchall():
            dados = dict(row)
            dados["foto"] = dados.get("foto") or None
            colaboradores.append(dados)

    _atualizar_cache_colaboradores(colaboradores)

    if not data_iso:
        return colaboradores

    chave = "motoristas" if funcao.lower().startswith("motor") else "ajudantes"
    indisponiveis = verificar_disponibilidade(data_iso, ignorar).get(chave, set())
    return [col for col in colaboradores if col["id"] not in indisponiveis]


def remover_bloqueios_por_carregamento(carregamento_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "DELETE FROM bloqueios WHERE carregamento_id = ?;",
            (carregamento_id,),
        )
        conn.commit()


def criar_bloqueios_para_carregamento(
    carregamento_id: int,
    data_iso: str,
    colaborador_ids: list[int | None],
    observacao: str,
):
    dias = OBSERVACAO_DURACAO.get(observacao, 0)
    data_inicio = datetime.strptime(data_iso, "%Y-%m-%d").date()
    data_fim = data_inicio + timedelta(days=dias)
    with get_connection() as conn:
        cur = conn.cursor()
        for colaborador_id in colaborador_ids:
            if not colaborador_id:
                continue
            cur.execute(
                """
                INSERT INTO bloqueios (colaborador_id, data_inicio, data_fim, motivo, carregamento_id)
                VALUES (?, ?, ?, ?, ?);
                """,
                (
                    colaborador_id,
                    data_inicio.isoformat(),
                    data_fim.isoformat(),
                    observacao,
                    carregamento_id,
                ),
            )
        conn.commit()


def colaborador_escalado_na_data(colaborador_id: int, data_iso: str, ignorar_id: int | None = None) -> bool:
    if not colaborador_id:
        return False
    ignorar = {"carregamento_id": ignorar_id} if ignorar_id else None
    indisponiveis = verificar_disponibilidade(data_iso, ignorar)
    return (
        colaborador_id in indisponiveis.get("motoristas", set())
        or colaborador_id in indisponiveis.get("ajudantes", set())
    )


def recarregar_disponibilidade():
    refresh_disponibilidade_carregamentos()
    refresh_oficina_motoristas()
    refresh_oficina_caminhoes()
    refresh_escala_cd_dropdowns()
    atualizar_dropdown_folga()
    if "atualizar_filtros_log" in globals():
        try:
            atualizar_filtros_log()
        except Exception:
            pass


def atualizar_label_dia_semana_carreg():
    if label_carreg_dia_semana is None:
        return
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        label_carreg_dia_semana.configure(text="")
        return
    try:
        data_dt = datetime.strptime(data_iso, "%Y-%m-%d").date()
    except ValueError:
        label_carreg_dia_semana.configure(text="")
        return
    nome_dia = DIAS_EXTENSO[data_dt.weekday()].capitalize()
    label_carreg_dia_semana.configure(text=f"Dia da semana: {nome_dia}")


def atualizar_botao_motorista_ajudante():
    if btn_motorista_ajudante is None:
        return
    if permitir_motorista_ajudante:
        btn_motorista_ajudante.configure(text="Motorista como ajudante: ON")
        estilizar_botao(btn_motorista_ajudante, "primary", small=True)
    else:
        btn_motorista_ajudante.configure(text="Motorista como ajudante: OFF")
        estilizar_botao(btn_motorista_ajudante, "ghost", small=True)


def alternar_motorista_como_ajudante():
    global permitir_motorista_ajudante
    permitir_motorista_ajudante = not permitir_motorista_ajudante
    atualizar_botao_motorista_ajudante()
    refresh_disponibilidade_carregamentos()


def salvar_folga(
    data_inicio: str,
    colaborador_id: int,
    data_fim: str | None = None,
    observacao_padrao: str | None = None,
    observacao_extra: str | None = None,
    observacao_cor: str | None = None,
) -> int:
    data_fim = data_fim or None
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO folgas (data, data_fim, colaborador_id, observacao_padrao, observacao_extra, observacao_cor)
            VALUES (?, ?, ?, ?, ?, ?);
            """,
            (
                data_inicio,
                data_fim,
                colaborador_id,
                observacao_padrao,
                observacao_extra,
                observacao_cor,
            ),
        )
        conn.commit()
        return cur.lastrowid


def listar_folgas(data_iso: str):
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT
                f.id AS folga_id,
                c.id AS colaborador_id,
                c.nome,
                c.funcao,
                f.data,
                f.data_fim,
                f.observacao_padrao,
                f.observacao_extra,
                f.observacao_cor
            FROM folgas f
            INNER JOIN colaboradores c ON c.id = f.colaborador_id
            WHERE f.data = ?
            ORDER BY c.nome;
            """,
            (data_iso,),
        )
        return [dict(row) for row in cur.fetchall()]


def editar_folga(
    folga_id: int,
    data_inicio: str,
    data_fim: str | None,
    colaborador_id: int,
    observacao_padrao: str | None,
    observacao_extra: str | None,
    observacao_cor: str | None,
):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE folgas
            SET
                data = ?,
                data_fim = ?,
                colaborador_id = ?,
                observacao_padrao = ?,
                observacao_extra = ?,
                observacao_cor = ?
            WHERE id = ?;
            """,
            (
                data_inicio,
                data_fim,
                colaborador_id,
                observacao_padrao,
                observacao_extra,
                observacao_cor,
                folga_id,
            ),
        )
        conn.commit()


def remover_folga(folga_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM folgas WHERE id = ?;", (folga_id,))
        conn.commit()


def validar_periodo(data_inicio: str, data_fim: str):
    dt_inicio = datetime.strptime(data_inicio, "%Y-%m-%d").date()
    dt_fim = datetime.strptime(data_fim, "%Y-%m-%d").date()
    if dt_inicio > dt_fim:
        raise ValueError("A data de início não pode ser posterior à data de fim.")


def adicionar_ferias(colaborador_id: int, data_inicio: str, data_fim: str, observacao: str | None):
    validar_periodo(data_inicio, data_fim)
    observacao_db = (observacao or "").strip() or None
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO ferias (colaborador_id, data_inicio, data_fim, observacao)
            VALUES (?, ?, ?, ?);
            """,
            (colaborador_id, data_inicio, data_fim, observacao_db),
        )
        conn.commit()
        return cur.lastrowid


def atualizar_ferias(
    registro_id: int,
    colaborador_id: int,
    data_inicio: str,
    data_fim: str,
    observacao: str | None,
):
    validar_periodo(data_inicio, data_fim)
    observacao_db = (observacao or "").strip() or None
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE ferias
            SET colaborador_id = ?, data_inicio = ?, data_fim = ?, observacao = ?
            WHERE id = ?;
            """,
            (colaborador_id, data_inicio, data_fim, observacao_db, registro_id),
        )
        conn.commit()


def remover_ferias(registro_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM ferias WHERE id = ?;", (registro_id,))
        conn.commit()


def listar_ferias():
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT f.id,
                   f.colaborador_id,
                   c.nome,
                   f.data_inicio,
                   f.data_fim,
                   f.observacao
            FROM ferias f
            INNER JOIN colaboradores c ON c.id = f.colaborador_id
            ORDER BY f.data_inicio DESC, c.nome;
            """
        )
        return [dict(row) for row in cur.fetchall()]


def salvar_carregamento(
    data_iso: str,
    rota_texto: str,
    placa: str | None,
    motorista_id: int | None,
    ajudante_id: int | None,
    observacao: str,
    observacao_extra: str | None = None,
    observacao_cor: str | None = None,
    data_saida: str | None = None,
) -> int:
    if motorista_id and ajudante_id and motorista_id == ajudante_id:
        raise ValueError("Motorista e ajudante devem ser pessoas diferentes.")

    placa_db = placa.strip().upper() if placa else None
    observacao_db = observacao.strip() if observacao else None
    observacao_extra_db = observacao_extra.strip() if observacao_extra else None
    observacao_cor_db = observacao_cor.strip() if observacao_cor else None
    data_saida_db = data_saida or calcular_data_saida_carregamento(data_iso) or data_iso

    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO carregamentos (data, data_saida, rota, placa, motorista_id, ajudante_id, observacao, observacao_extra, observacao_cor)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
            """,
            (
                data_iso,
                data_saida_db,
                rota_texto,
                placa_db,
                motorista_id,
                ajudante_id,
                observacao_db,
                observacao_extra_db,
                observacao_cor_db,
            ),
        )
        conn.commit()
        return cur.lastrowid


def atualizar_carregamento(
    carregamento_id: int,
    data_iso: str,
    data_saida: str | None,
    rota_texto: str,
    placa: str | None,
    motorista_id: int | None,
    ajudante_id: int | None,
    observacao: str,
    observacao_extra: str | None = None,
    observacao_cor: str | None = None,
):
    placa_db = placa.strip().upper() if placa else None
    observacao_db = observacao.strip() if observacao else None
    observacao_extra_db = observacao_extra.strip() if observacao_extra else None
    observacao_cor_db = observacao_cor.strip() if observacao_cor else None
    data_saida_db = data_saida or calcular_data_saida_carregamento(data_iso) or data_iso
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE carregamentos
            SET data = ?, data_saida = ?, rota = ?, placa = ?, motorista_id = ?, ajudante_id = ?, observacao = ?, observacao_extra = ?, observacao_cor = ?
            WHERE id = ?;
            """,
            (
                data_iso,
                data_saida_db,
                rota_texto,
                placa_db,
                motorista_id,
                ajudante_id,
                observacao_db,
                observacao_extra_db,
                observacao_cor_db,
                carregamento_id,
            ),
        )
        conn.commit()


def remover_carregamento(carregamento_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM carregamentos WHERE id = ?;", (carregamento_id,))
        conn.commit()


def listar_carregamentos(data_iso: str):
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT car.id,
                   car.data,
                   car.data_saida,
                   car.rota,
                   car.placa,
                   car.observacao,
                   car.observacao_extra,
                   car.observacao_cor,
                   car.motorista_id,
                   car.ajudante_id,
                   mot.nome AS motorista_nome,
                   aj.nome AS ajudante_nome
            FROM carregamentos car
            LEFT JOIN colaboradores mot ON mot.id = car.motorista_id
            LEFT JOIN colaboradores aj ON aj.id = car.ajudante_id
            WHERE car.data = ?
            ORDER BY car.rota ASC, car.id ASC;
            """,
            (data_iso,),
        )
        return [dict(row) for row in cur.fetchall()]


def obter_data_saida_registro(registro: dict) -> str:
    valor = (registro.get("data_saida") or "").strip()
    if valor:
        return valor
    data_base = registro.get("data")
    return calcular_data_saida_carregamento(data_base) or data_base


def listar_ajustes_por_carregamentos(carregamento_ids: list[int]) -> dict[int, list[dict]]:
    if not carregamento_ids:
        return {}
    placeholders = ",".join("?" for _ in carregamento_ids)
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT carregamento_id,
                   data_ajuste,
                   duracao_anterior,
                   duracao_nova,
                   observacao_ajuste
            FROM ajustes_rotas
            WHERE carregamento_id IN ({placeholders})
            ORDER BY data_ajuste ASC, id ASC;
            """,
            carregamento_ids,
        )
        rows = [dict(row) for row in cur.fetchall()]
    agrupado: dict[int, list[dict]] = {}
    for row in rows:
        agrupado.setdefault(row["carregamento_id"], []).append(row)
    return agrupado


def registrar_ajuste_rota(
    carregamento_id: int,
    duracao_anterior: int,
    duracao_nova: int,
    observacao_ajuste: str | None = None,
):
    data_ajuste = datetime.now().strftime("%Y-%m-%d %H:%M")
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO ajustes_rotas (
                carregamento_id,
                data_ajuste,
                duracao_anterior,
                duracao_nova,
                observacao_ajuste
            )
            VALUES (?, ?, ?, ?, ?);
            """,
            (
                carregamento_id,
                data_ajuste,
                duracao_anterior,
                duracao_nova,
                (observacao_ajuste or "").strip() or None,
            ),
        )
        conn.commit()


def atualizar_bloqueios_para_ajuste(
    carregamento_id: int,
    nova_data_fim_iso: str,
    liberar_imediato: bool = False,
):
    with get_connection() as conn:
        cur = conn.cursor()
        if liberar_imediato:
            cur.execute("DELETE FROM bloqueios WHERE carregamento_id = ?;", (carregamento_id,))
        else:
            cur.execute(
                "UPDATE bloqueios SET data_fim = ? WHERE carregamento_id = ?;",
                (nova_data_fim_iso, carregamento_id),
            )
        conn.commit()


def remover_ajustes_por_carregamento(carregamento_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM ajustes_rotas WHERE carregamento_id = ?;", (carregamento_id,))
        conn.commit()


def remover_carregamento_completo(carregamento_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "DELETE FROM bloqueios WHERE carregamento_id = ?;",
            (carregamento_id,),
        )
        cur.execute(
            "DELETE FROM ajustes_rotas WHERE carregamento_id = ?;",
            (carregamento_id,),
        )
        cur.execute(
            "DELETE FROM carregamentos WHERE id = ?;",
            (carregamento_id,),
        )
        conn.commit()


def montar_resumo_ajustes(duracao_planejada: int, ajustes: list[dict]) -> str:
    if not ajustes:
        return f"Planejado: {dias_para_texto(duracao_planejada)}"
    resumo = ""
    anterior = duracao_planejada
    for ajuste in ajustes:
        novo = ajuste.get("duracao_nova", anterior)
        if novo > anterior:
            diff = novo - anterior
            resumo = f"Era {dias_para_texto(anterior)}, ficou +{diff} = {dias_para_texto(novo)}"
        elif novo < anterior:
            resumo = f"Era {dias_para_texto(anterior)}, voltaram em {dias_para_texto(novo)}"
        else:
            resumo = f"Ajustado e mantido em {dias_para_texto(novo)}"
        anterior = novo
    return resumo


def consultar_log_carregamentos(filtros: dict) -> list[dict]:
    query = [
        """
        SELECT car.id,
               car.data,
               car.data_saida,
               car.rota,
               car.placa,
               car.observacao,
               car.observacao_extra,
               car.motorista_id,
               car.ajudante_id,
               mot.nome AS motorista_nome,
               aj.nome AS ajudante_nome
        FROM carregamentos car
        LEFT JOIN colaboradores mot ON mot.id = car.motorista_id
        LEFT JOIN colaboradores aj ON aj.id = car.ajudante_id
        WHERE 1 = 1
        """
    ]
    params: list = []
    if filtros.get("data_inicio"):
        query.append("AND car.data >= ?")
        params.append(filtros["data_inicio"])
    if filtros.get("data_fim"):
        query.append("AND car.data <= ?")
        params.append(filtros["data_fim"])
    if filtros.get("motorista_id"):
        query.append("AND car.motorista_id = ?")
        params.append(filtros["motorista_id"])
    if filtros.get("placa"):
        query.append("AND UPPER(car.placa) = ?")
        params.append(filtros["placa"].upper())
    query.append("ORDER BY car.data DESC, car.id DESC")

    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(" ".join(query), tuple(params))
        registros = [dict(row) for row in cur.fetchall()]

    ajustes_map = listar_ajustes_por_carregamentos([reg["id"] for reg in registros])
    hoje = date.today()
    resultado: list[dict] = []

    for registro in registros:
        observacao_padrao = (registro.get("observacao") or "0").strip()
        duracao_planejada = OBSERVACAO_DURACAO.get(observacao_padrao, 0)
        ajustes = ajustes_map.get(registro["id"], [])
        duracao_efetiva = ajustes[-1]["duracao_nova"] if ajustes else duracao_planejada
        data_inicio_iso = obter_data_saida_registro(registro)
        try:
            data_inicio_dt = datetime.strptime(data_inicio_iso, "%Y-%m-%d").date()
        except ValueError:
            data_inicio_dt = hoje
            data_inicio_iso = hoje.isoformat()
        data_fim_dt = data_inicio_dt + timedelta(days=duracao_efetiva)
        data_fim_iso = data_fim_dt.isoformat()

        finalizado_manual = bool(ajustes) and duracao_efetiva <= 0
        if finalizado_manual:
            status = "Finalizado"
        elif hoje < data_inicio_dt:
            status = "Em andamento"
        elif hoje < data_fim_dt:
            status = "Em andamento"
        else:
            status = "Finalizado"

        status_filtro = filtros.get("status")
        if status_filtro and status_filtro != "Todos":
            if status_filtro == "Em andamento" and status != "Em andamento":
                continue
            if status_filtro == "Finalizados" and status != "Finalizado":
                continue

        restante = max((data_fim_dt - hoje).days, 0)
        andamento_texto = ""
        if status == "Em andamento":
            if restante > 0:
                andamento_texto = f"{observacao_padrao or 'ROTA'} — faltando {restante}"
            else:
                andamento_texto = f"{observacao_padrao or 'ROTA'} — retorna hoje"

        ajudante_nome = formatar_ajudante_nome(
            registro.get("ajudante_nome") or DISPLAY_VAZIO,
            registro.get("ajudante_id"),
        )
        resultado.append(
            {
                "id": registro["id"],
                "data": registro.get("data"),
                "data_br": data_iso_para_br(registro.get("data")),
                "data_saida": data_inicio_iso,
                "data_saida_br": data_iso_para_br(data_inicio_iso),
                "data_fim": data_fim_iso,
                "data_fim_br": data_iso_para_br(data_fim_iso),
                "rota": registro.get("rota") or DISPLAY_VAZIO,
                "placa": (registro.get("placa") or "").upper() or DISPLAY_VAZIO,
                "motorista": registro.get("motorista_nome") or DISPLAY_VAZIO,
                "ajudante": ajudante_nome,
                "observacao": observacao_padrao,
                "duracao_planejada": duracao_planejada,
                "duracao_efetiva": duracao_efetiva,
                "status": status,
                "status_texto": andamento_texto if status == "Em andamento" else "",
                "resumo": montar_resumo_ajustes(duracao_planejada, ajustes),
                "ajustes": ajustes,
            }
        )

    return resultado


def exportar_log_para_excel(registros: list[dict]):
    if not registros:
        messagebox.showinfo("LOG", "Nenhum registro para exportar.")
        return
    if Workbook is None:
        messagebox.showerror(
            "Exportação",
            "Biblioteca 'openpyxl' não disponível. Instale-a para exportar em Excel.",
        )
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "LOG"
    cabecalho = [
        "Data Carregamento",
        "Data Saída",
        "Motorista",
        "Ajudante",
        "Placa",
        "Rota",
        "Duração planejada (dias)",
        "Duração efetiva (dias)",
        "Status",
        "Resumo histórico",
    ]
    ws.append(cabecalho)
    if Font:
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for item in registros:
        ws.append(
            [
                item.get("data_br"),
                item.get("data_saida_br"),
                item.get("motorista"),
                item.get("ajudante"),
                item.get("placa"),
                item.get("rota"),
                item.get("duracao_planejada"),
                item.get("duracao_efetiva"),
                item.get("status"),
                item.get("resumo"),
            ]
        )
    relatorio_nome = f"log_jr_escala_{date.today().isoformat()}.xlsx"
    caminho = RELATORIOS_DIR / relatorio_nome
    wb.save(caminho)
    messagebox.showinfo(
        "Exportação concluída",
        f"? Log exportado com sucesso!\nArquivo salvo em: {caminho}",
    )


def carregamento_existe_para_rota(data_iso: str, rota_texto: str) -> bool:
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT 1 FROM carregamentos WHERE data = ? AND rota = ? LIMIT 1;",
            (data_iso, rota_texto),
        )
        return cur.fetchone() is not None


def preencher_carregamentos_automaticos():
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        atualizar_mensagem_rotas_auto("Informe uma data válida para carregar rotas.", COR_VERMELHA)
        return

    rotas = listar_rotas_para_data(data_iso)
    if not rotas:
        atualizar_mensagem_rotas_auto("Nenhuma rota automática cadastrada para este dia.")
        return

    inseridos = 0
    for rota in rotas:
        texto_rota = rota["rota"].strip()
        destino = rota.get("destino", "").strip()
        if destino:
            texto_rota = f"{texto_rota} - {destino}"
        if carregamento_existe_para_rota(data_iso, texto_rota):
            continue
        try:
            data_saida_raw = carreg_saida_var.get().strip() if "carreg_saida_var" in globals() else ""
            data_saida_iso = validar_data(data_saida_raw) if data_saida_raw else None
            observacao_extra = (rota.get("observacao") or "").strip() or None
            salvar_carregamento(
                data_iso,
                texto_rota,
                None,
                None,
                None,
                OBSERVACAO_OPCOES[0],
                observacao_extra=observacao_extra,
                data_saida=data_saida_iso,
            )
            inseridos += 1
        except sqlite3.Error:
            continue

    if inseridos:
        atualizar_mensagem_rotas_auto(
            f"{inseridos} rota(s) automática(s) adicionada(s) para {data_iso_para_br_entrada(data_iso)}.",
            COR_AZUL,
        )
    else:
        atualizar_mensagem_rotas_auto("Todas as rotas automáticas já estavam carregadas.")
    refresh_carregamentos_ui()
    recarregar_disponibilidade()


def limpar_form_carregamento():
    global carregamento_em_edicao_id
    carregamento_em_edicao_id = None
    entry_rota_num.delete(0, "end")
    entry_rota_destino.delete(0, "end")
    entry_carreg_placa.delete(0, "end")
    entry_observacao_extra.delete(0, "end")
    if VALOR_SEM_CAMINHAO in combo_carreg_placa.cget("values"):
        combo_carreg_placa.set(VALOR_SEM_CAMINHAO)
    valores_motoristas = combo_motorista.cget("values")
    if valores_motoristas:
        combo_motorista.set(valores_motoristas[0])
        atualizar_preview_combo_colaborador(
            combo_motorista, motorista_option_map, preview_motorista_carreg_label
        )
    valores_ajudantes = combo_ajudante.cget("values")
    if valores_ajudantes:
        combo_ajudante.set(valores_ajudantes[0])
        atualizar_preview_combo_colaborador(
            combo_ajudante, ajudante_option_map, preview_ajudante_carreg_label
        )
    combo_observacao.configure(values=OBSERVACAO_OPCOES)
    combo_observacao.set(OBSERVACAO_OPCOES[0])
    combo_cor_observacao.set(OBS_MARCADORES[0][0])
    btn_salvar_carreg.configure(text="Salvar carregamento")
    btn_cancelar_edicao_carreg.configure(state="disabled")


def preparar_edicao_carregamento(item: dict):
    global carregamento_em_edicao_id, carregamento_dia_selecionado_id
    carregamento_em_edicao_id = item["id"]
    carregamento_dia_selecionado_id = item.get("id")
    if combo_carregamento_dia is not None:
        for display, registro in carregamento_dia_map.items():
            if registro.get("id") == carregamento_dia_selecionado_id:
                combo_carregamento_dia.set(display)
                break
    rota_texto = item["rota"]
    partes = rota_texto.split("-", 1)
    entry_rota_num.delete(0, "end")
    entry_rota_num.insert(0, partes[0].strip())
    entry_rota_destino.delete(0, "end")
    entry_rota_destino.insert(0, partes[1].strip() if len(partes) > 1 else "")
    entrada_obs = item.get("observacao") or OBSERVACAO_OPCOES[0]
    valores_obs = list(combo_observacao.cget("values"))
    if entrada_obs not in valores_obs:
        valores_obs.append(entrada_obs)
        combo_observacao.configure(values=valores_obs)
    combo_observacao.set(entrada_obs)
    entry_observacao_extra.delete(0, "end")
    entry_observacao_extra.insert(0, item.get("observacao_extra") or "")
    combo_cor_observacao.set(label_cor_observacao(item.get("observacao_cor")))
    placa_texto = item.get("placa") or ""
    entry_carreg_placa.delete(0, "end")
    entry_carreg_placa.insert(0, placa_texto.upper())
    selecionar_combo_caminhao_por_placa(placa_texto)
    selecionar_combo_colaborador(
        combo_motorista,
        motorista_option_map,
        item.get("motorista_id"),
        item.get("motorista_nome"),
        VALOR_SEM_MOTORISTA,
        preview_motorista_carreg_label,
    )
    selecionar_combo_colaborador(
        combo_ajudante,
        ajudante_option_map,
        item.get("ajudante_id"),
        item.get("ajudante_nome"),
        VALOR_SEM_AJUDANTE,
        preview_ajudante_carreg_label,
    )
    data_saida_iso = (
        item.get("data_saida")
        or calcular_data_saida_carregamento(item.get("data"))
        or ""
    )
    carreg_saida_var.set(data_iso_para_br_entrada(data_saida_iso))
    btn_salvar_carreg.configure(text="Atualizar carregamento")
    btn_cancelar_edicao_carreg.configure(state="normal")


def duplicar_carregamento_dia(item: dict):
    global carregamento_em_edicao_id, carregamento_dia_selecionado_id
    carregamento_em_edicao_id = None
    carregamento_dia_selecionado_id = None
    if combo_carregamento_dia is not None:
        combo_carregamento_dia.set(CARREGAMENTO_DIA_PLACEHOLDER)
    rota_texto = item["rota"]
    partes = rota_texto.split("-", 1)
    entry_rota_num.delete(0, "end")
    entry_rota_num.insert(0, partes[0].strip())
    entry_rota_destino.delete(0, "end")
    entry_rota_destino.insert(0, partes[1].strip() if len(partes) > 1 else "")
    entrada_obs = item.get("observacao") or OBSERVACAO_OPCOES[0]
    valores_obs = list(combo_observacao.cget("values"))
    if entrada_obs not in valores_obs:
        valores_obs.append(entrada_obs)
        combo_observacao.configure(values=valores_obs)
    combo_observacao.set(entrada_obs)
    entry_observacao_extra.delete(0, "end")
    entry_observacao_extra.insert(0, item.get("observacao_extra") or "")
    combo_cor_observacao.set(label_cor_observacao(item.get("observacao_cor")))
    entry_carreg_placa.delete(0, "end")
    if VALOR_SEM_CAMINHAO in combo_carreg_placa.cget("values"):
        combo_carreg_placa.set(VALOR_SEM_CAMINHAO)
    valores_motoristas = combo_motorista.cget("values")
    if VALOR_SEM_MOTORISTA in valores_motoristas:
        combo_motorista.set(VALOR_SEM_MOTORISTA)
        atualizar_preview_combo_colaborador(
            combo_motorista, motorista_option_map, preview_motorista_carreg_label
        )
    valores_ajudantes = combo_ajudante.cget("values")
    if VALOR_SEM_AJUDANTE in valores_ajudantes:
        combo_ajudante.set(VALOR_SEM_AJUDANTE)
        atualizar_preview_combo_colaborador(
            combo_ajudante, ajudante_option_map, preview_ajudante_carreg_label
        )
    data_saida_iso = (
        item.get("data_saida")
        or calcular_data_saida_carregamento(item.get("data"))
        or ""
    )
    carreg_saida_var.set(data_iso_para_br_entrada(data_saida_iso))
    btn_salvar_carreg.configure(text="Salvar carregamento")
    btn_cancelar_edicao_carreg.configure(state="disabled")


def atualizar_preview_combo_colaborador(
    combo,
    mapping,
    preview_label,
    size: tuple[int, int] = FOTO_PREVIEW_PADRAO,
    selecionado: str | None = None,
    circular: bool = False,
):
    if preview_label is None or combo is None:
        return
    escolha = selecionado if selecionado is not None else combo.get()
    if escolha is None:
        return
    colaborador_id = mapping.get(escolha) if mapping else None
    imagem = obter_imagem_colaborador_por_id(colaborador_id, size=size, circular=circular)
    if imagem:
        preview_label.configure(image=imagem, text="")
        preview_label.image = imagem
    else:
        preview_label.configure(image=None, text="Sem foto")
        preview_label.image = None


def selecionar_combo_colaborador(combo, mapping, colaborador_id, nome, padrao, preview_label=None):
    if not colaborador_id:
        combo.set(padrao)
        if preview_label is not None:
            atualizar_preview_combo_colaborador(combo, mapping, preview_label)
        return
    for display, cid in mapping.items():
        if cid == colaborador_id:
            combo.set(display)
            if preview_label is not None:
                atualizar_preview_combo_colaborador(combo, mapping, preview_label)
            return
    if nome:
        novo_display = f"{nome} (#{colaborador_id})"
    else:
        novo_display = f"#{colaborador_id}"
    combo_ajudante_ref = globals().get("combo_ajudante")
    combo_cd_ajudante_ref = globals().get("combo_cd_ajudante")
    if combo in (combo_ajudante_ref, combo_cd_ajudante_ref):
        novo_display = formatar_ajudante_nome(novo_display, colaborador_id)
    valores = list(combo.cget("values"))
    valores.append(novo_display)
    combo.configure(values=valores)
    mapping[novo_display] = colaborador_id
    combo.set(novo_display)
    if preview_label is not None:
        atualizar_preview_combo_colaborador(combo, mapping, preview_label)


def selecionar_combo_caminhao_por_placa(placa: str):
    if not placa:
        if VALOR_SEM_CAMINHAO in combo_carreg_placa.cget("values"):
            combo_carreg_placa.set(VALOR_SEM_CAMINHAO)
        return
    placa_upper = placa.upper()
    for display, valor in caminhao_option_map.items():
        if valor.upper() == placa_upper:
            combo_carreg_placa.set(display)
            return
    combo_carreg_placa.set(VALOR_SEM_CAMINHAO)


def set_oficina_placa_field(placa: str):
    placa_upper = (placa or "").upper()
    if (
        combo_oficina_placa.winfo_manager()
        and combo_oficina_placa.cget("state") != "disabled"
    ):
        valores = list(combo_oficina_placa.cget("values"))
        for display in valores:
            if caminhao_option_map.get(display, "").upper() == placa_upper:
                combo_oficina_placa.set(display)
                entry_oficina_placa.delete(0, "end")
                return
        if valores:
            combo_oficina_placa.set(valores[0])
        entry_oficina_placa.delete(0, "end")
        entry_oficina_placa.insert(0, placa_upper)
    else:
        entry_oficina_placa.delete(0, "end")
        entry_oficina_placa.insert(0, placa_upper)


def handle_excluir_carregamento(item: dict):
    global carregamento_dia_selecionado_id
    if not messagebox.askyesno(
        "Excluir carregamento",
        f"Deseja remover a rota '{item['rota']}' de {item['data']}?",
    ):
        return
    try:
        remover_carregamento_completo(item["id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    carregamentos_revisados.discard(item["id"])
    if carregamento_dia_selecionado_id == item["id"]:
        carregamento_dia_selecionado_id = None
    limpar_form_carregamento()
    refresh_carregamentos_ui()
    atualizar_mensagem_rotas_auto("Carregamento removido.", COR_AZUL)
    recarregar_disponibilidade()
    if "aplicar_filtros_log" in globals():
        aplicar_filtros_log()


def duplicar_carregamentos_para_data(destino: str) -> str:
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT COUNT(1) FROM carregamentos WHERE data = ?;",
            (destino,),
        )
        if cur.fetchone()[0]:
            raise ValueError("Já existem carregamentos para esta data.")

        cur.execute(
            """
            SELECT data FROM carregamentos
            WHERE data < ?
            ORDER BY data DESC
            LIMIT 1;
            """,
            (destino,),
        )
        origem = cur.fetchone()
        if not origem:
            cur.execute(
                "SELECT data FROM carregamentos WHERE data <> ? ORDER BY data DESC LIMIT 1;",
                (destino,),
            )
            origem = cur.fetchone()
        if not origem:
            raise LookupError("Não há carregamentos anteriores para duplicar.")

        origem_data = origem[0]
        cur.execute(
            """
            INSERT INTO carregamentos (data, data_saida, rota, placa, motorista_id, ajudante_id, observacao, observacao_extra, observacao_cor)
            SELECT ?, data_saida, rota, placa, motorista_id, ajudante_id, observacao, observacao_extra, observacao_cor
            FROM carregamentos
            WHERE data = ?;
            """,
            (destino, origem_data),
        )
        conn.commit()
        return origem_data


def salvar_oficina(
    data_iso: str,
    motorista_id: int | None,
    placa: str,
    observacao: str,
    observacao_extra: str | None = None,
    data_saida: str | None = None,
    observacao_cor: str | None = None,
) -> int:
    disponibilidade = verificar_disponibilidade(data_iso)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )
    if motorista_id and motorista_id in indis_colaboradores:
        raise ValueError("Este motorista está indisponível nesta data.")

    data_saida_iso = data_saida or calcular_data_saida_padrao(data_iso)
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO oficinas (
                data,
                motorista_id,
                placa,
                observacao,
                observacao_extra,
                data_saida,
                observacao_cor
            )
            VALUES (?, ?, ?, ?, ?, ?, ?);
            """,
            (
                data_iso,
                motorista_id,
                placa,
                observacao,
                (observacao_extra or "").strip() or None,
                data_saida_iso,
                observacao_cor.strip() if observacao_cor else None,
            ),
        )
        conn.commit()
        return cur.lastrowid


def listar_oficinas(data_iso: str):
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT ofi.id,
                   ofi.data,
                   ofi.motorista_id,
                   ofi.placa,
                   ofi.observacao,
                   ofi.observacao_extra,
                   ofi.data_saida,
                   ofi.observacao_cor,
                   col.nome AS motorista_nome
            FROM oficinas ofi
            LEFT JOIN colaboradores col ON col.id = ofi.motorista_id
            WHERE ofi.data = ?
            ORDER BY ofi.id ASC;
            """,
            (data_iso,),
        )
        return [dict(row) for row in cur.fetchall()]


def editar_oficina(
    oficina_id: int,
    motorista_id: int | None,
    placa: str,
    observacao: str,
    observacao_extra: str | None = None,
    data_saida: str | None = None,
    observacao_cor: str | None = None,
):
    data_saida_iso = data_saida or None
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE oficinas
            SET motorista_id = ?, placa = ?, observacao = ?, observacao_extra = ?, data_saida = ?, observacao_cor = ?
            WHERE id = ?;
            """,
            (
                motorista_id,
                placa,
                observacao,
                (observacao_extra or "").strip() or None,
                data_saida_iso,
                observacao_cor.strip() if observacao_cor else None,
                oficina_id,
            ),
        )
        conn.commit()


def excluir_oficina(oficina_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM oficinas WHERE id = ?;", (oficina_id,))
        conn.commit()


def listar_rotas_semanais(dia_semana: str):
    dia = normalizar_dia_semana(dia_semana)
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, dia_semana, rota, destino, observacao
            FROM rotas_semanais
            WHERE dia_semana = ?
            ORDER BY rota COLLATE NOCASE ASC;
            """,
            (dia,),
        )
        return [dict(row) for row in cur.fetchall()]


def adicionar_rota_semana(dia_semana: str, rota: str, destino: str, observacao: str):
    dia = normalizar_dia_semana(dia_semana)
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO rotas_semanais (dia_semana, rota, destino, observacao)
            VALUES (?, ?, ?, ?);
            """,
            (dia, rota.strip(), (destino or "").strip(), (observacao or "").strip()),
        )
        conn.commit()
        return cur.lastrowid


def editar_rota_semana(rota_id: int, dia_semana: str, rota: str, destino: str, observacao: str):
    dia = normalizar_dia_semana(dia_semana)
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE rotas_semanais
            SET dia_semana = ?, rota = ?, destino = ?, observacao = ?
            WHERE id = ?;
            """,
            (dia, rota.strip(), (destino or "").strip(), (observacao or "").strip(), rota_id),
        )
        conn.commit()


def remover_rota_semana(rota_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM rotas_semanais WHERE id = ?;", (rota_id,))
        conn.commit()


def listar_rotas_para_data(data_iso: str):
    dia = obter_dia_semana_por_data(data_iso)
    return listar_rotas_semanais(dia)


def adicionar_escala_cd(data_iso: str, motorista_id: int | None, ajudante_id: int | None, observacao: str):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO escala_cd (data, motorista_id, ajudante_id, observacao)
            VALUES (?, ?, ?, ?);
            """,
            (data_iso, motorista_id, ajudante_id, observacao),
        )
        conn.commit()
        return cur.lastrowid


def listar_escala_cd(data_iso: str):
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(
            """
            SELECT e.id,
                   e.data,
                   e.observacao,
                   mot.id AS motorista_id,
                   mot.nome AS motorista_nome,
                   aj.id AS ajudante_id,
                   aj.nome AS ajudante_nome
            FROM escala_cd e
            LEFT JOIN colaboradores mot ON mot.id = e.motorista_id
            LEFT JOIN colaboradores aj ON aj.id = e.ajudante_id
            WHERE e.data = ?
            ORDER BY e.id ASC;
            """,
            (data_iso,),
        )
        return [dict(row) for row in cur.fetchall()]


def editar_escala_cd(escala_id: int, motorista_id: int | None, ajudante_id: int | None, observacao: str):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE escala_cd
            SET motorista_id = ?, ajudante_id = ?, observacao = ?
            WHERE id = ?;
            """,
            (motorista_id, ajudante_id, observacao, escala_id),
        )
        conn.commit()


def excluir_escala_cd(escala_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM escala_cd WHERE id = ?;", (escala_id,))
        conn.commit()


def add_caminhao(placa: str, modelo: str, observacao: str) -> int:
    placa_norm = placa.strip().upper()
    if not placa_norm:
        raise ValueError("Informe a placa do caminhão.")
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO caminhoes (placa, modelo, observacao, ativo)
            VALUES (?, ?, ?, 1);
            """,
            (placa_norm, modelo.strip(), observacao.strip()),
        )
        conn.commit()
        return cur.lastrowid


def listar_caminhoes(ativos_only: bool = True):
    query = "SELECT id, placa, modelo, observacao, ativo FROM caminhoes"
    if ativos_only:
        query += " WHERE ativo = 1"
    query += " ORDER BY placa;"
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute(query)
        return [dict(row) for row in cur.fetchall()]


def listar_caminhoes_ativos():
    return listar_caminhoes(ativos_only=True)


def editar_caminhao(caminhao_id: int, placa: str, modelo: str, observacao: str):
    placa_norm = placa.strip().upper()
    if not placa_norm:
        raise ValueError("Informe a placa do caminhão.")
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            """
            UPDATE caminhoes
            SET placa = ?, modelo = ?, observacao = ?
            WHERE id = ?;
            """,
            (placa_norm, modelo.strip(), observacao.strip(), caminhao_id),
        )
        conn.commit()


def remover_caminhao(caminhao_id: int):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM caminhoes WHERE id = ?;", (caminhao_id,))
        conn.commit()


def placa_em_manutencao(placa: str, data_iso: str) -> bool:
    if not placa:
        return False
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT 1 FROM oficinas WHERE placa = ? AND data = ? LIMIT 1;",
            (placa.upper(), data_iso),
        )
        return cur.fetchone() is not None


def validar_data(data_texto: str) -> str | None:
    texto = (data_texto or "").strip()
    if not texto:
        return None
    if "/" in texto:
        return data_br_para_iso(texto)
    try:
        return datetime.strptime(texto, "%Y-%m-%d").date().isoformat()
    except ValueError:
        return None


DIAS_EXTENSO = [
    "segunda-feira",
    "terça-feira",
    "quarta-feira",
    "quinta-feira",
    "sexta-feira",
    "sábado",
    "domingo",
]
MESES_EXTENSO = [
    "janeiro",
    "fevereiro",
    "março",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


def data_por_extenso(dia: date | None = None) -> str:
    dia = dia or date.today()
    nome_dia = DIAS_EXTENSO[dia.weekday()]
    nome_mes = MESES_EXTENSO[dia.month - 1]
    return f"{nome_dia.capitalize()}, {dia.day:02d} de {nome_mes} de {dia.year}"
init_db()
garantir_coluna_observacao_colaboradores()
garantir_coluna_foto_colaboradores()
garantir_coluna_observacao_extra_carregamentos()
garantir_coluna_observacao_cor_carregamentos()
garantir_coluna_data_saida_carregamentos()
garantir_colunas_oficinas_complementares()
garantir_colunas_folgas_complementares()
garantir_coluna_observacao_cor_oficinas()
atualizar_carregamentos_para_campos_opcionais()
atualizar_oficinas_para_motorista_opcional()
garantir_coluna_carregamento_bloqueios()
limpar_bloqueios_expirados()

try:
    RELATORIOS_DIR.mkdir(parents=True, exist_ok=True)
except OSError:
    pass

try:
    FOTOS_COLAB_DIR.mkdir(parents=True, exist_ok=True)
except OSError:
    pass

try:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
except OSError:
    pass

criar_backup_automatico()

logo_image = None
if LOGO_PATH.exists():
    try:
        with Image.open(LOGO_PATH) as logo_raw:
            logo_copy = logo_raw.copy()
        logo_image = ctk.CTkImage(light_image=logo_copy, dark_image=logo_copy, size=(86, 86))
    except OSError:
        logo_image = None

header_gradient_image = ctk.CTkImage(
    light_image=criar_gradiente_horizontal(1200, 80, COR_AZUL, COR_AZUL_GRADIENTE_FIM),
    dark_image=criar_gradiente_horizontal(1200, 80, COR_AZUL, COR_AZUL_GRADIENTE_FIM),
    size=(1200, 80),
)

app = ctk.CTk()
app.title("JR Escala")

# Ajusta a janela para caber em monitores menores sem esconder o final da aba
DEFAULT_SIZE = (1280, 800)
MIN_SIZE = (1100, 720)
MARGEM_JANELA = (40, 80)  # deixa espaço para a barra do Windows e bordas
screen_width = app.winfo_screenwidth()
screen_height = app.winfo_screenheight()
usable_width = max(MIN_SIZE[0], screen_width - MARGEM_JANELA[0])
usable_height = max(MIN_SIZE[1], screen_height - MARGEM_JANELA[1])
janela_largura = min(DEFAULT_SIZE[0], usable_width)
janela_altura = min(DEFAULT_SIZE[1], usable_height)
app.geometry(f"{janela_largura}x{janela_altura}")
app.resizable(True, True)
app.minsize(*MIN_SIZE)
app.configure(fg_color=COR_FUNDO)

DEBOUNCE_HANDLES: dict[str, str] = {}


def debounce(key: str, delay_ms: int, callback):
    handle = DEBOUNCE_HANDLES.get(key)
    if handle:
        try:
            app.after_cancel(handle)
        except Exception:
            pass

    def _run():
        DEBOUNCE_HANDLES.pop(key, None)
        callback()

    DEBOUNCE_HANDLES[key] = app.after(delay_ms, _run)

frame_header = ctk.CTkFrame(app, height=80, fg_color=NAV_BG, corner_radius=0)
frame_header.pack(fill="x", side="top")

header_bg = ctk.CTkLabel(frame_header, text="", image=header_gradient_image)
header_bg.place(relx=0, rely=0, relwidth=1, relheight=1)

header_inner = ctk.CTkFrame(frame_header, fg_color="transparent")
header_inner.pack(fill="x", padx=20, pady=10)

logo = ctk.CTkLabel(
    header_inner,
    text="" if logo_image else "JR",
    image=logo_image,
    text_color="white",
    font=("Segoe UI Semibold", 32),
    compound="center",
    fg_color="transparent",
)
logo.grid(row=0, column=0, sticky="w")

label_tab_titulo = ctk.CTkLabel(
    header_inner,
    text="JR Escala - Carregamentos",
    text_color="white",
    font=("Segoe UI Semibold", 16),
)
label_tab_titulo.grid(row=0, column=1, padx=20, sticky="w")

label_data = ctk.CTkLabel(
    header_inner,
    text=data_por_extenso(),
    text_color="white",
    font=("Segoe UI", 13),
)
label_data.grid(row=0, column=2, sticky="e", padx=10)

header_inner.grid_columnconfigure(1, weight=1)
header_inner.grid_columnconfigure(2, weight=0)

abas = ctk.CTkTabview(app, fg_color="transparent")
abas.pack(fill="both", expand=True, padx=20, pady=20)
aba_atual_nome = ""


def atualizar_titulo_aba():
    global aba_atual_nome
    nome = abas.get()
    if nome == aba_atual_nome:
        return
    aba_atual_nome = nome
    label_tab_titulo.configure(text=f"JR Escala - {nome}")
    app.title(f"JR Escala - {nome}")
    atualizar_estilo_abas(abas)


abas.configure(command=atualizar_titulo_aba)
atualizar_titulo_aba()
abas.add("Carregamentos")
abas.add("Escala (CD)")
abas.add("Folga")
abas.add("Oficinas")
abas.add("Rotas Semanais")
abas.add("Caminhões")
abas.add("Férias")
abas.add("Colaboradores")
abas.add("LOG")

aba_carregamentos = abas.tab("Carregamentos")
aba_cd = abas.tab("Escala (CD)")
aba_folga = abas.tab("Folga")
aba_oficinas = abas.tab("Oficinas")
aba_rotas = abas.tab("Rotas Semanais")
aba_caminhoes = abas.tab("Caminhões")
aba_ferias = abas.tab("Férias")
aba_colaboradores = abas.tab("Colaboradores")
aba_log = abas.tab("LOG")
personalizar_barra_abas(abas)

carregamentos_container = ctk.CTkScrollableFrame(
    aba_carregamentos, fg_color="transparent"
)
carregamentos_container.pack(fill="both", expand=True, padx=0, pady=0)
carregamentos_container.grid_columnconfigure(0, weight=1)

# ---------- LOG ----------
ctk.CTkLabel(
    aba_log,
    text="LOG DE ESCALAS",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_log,
    text="Acompanhe rotas em andamento, ajustes e histórico de duração.",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).pack(anchor="w", padx=10, pady=(0, 10))

frame_log_filtros = ctk.CTkFrame(aba_log, fg_color="#FFFFFF", corner_radius=12)
frame_log_filtros.pack(fill="x", padx=20, pady=(5, 12))
estilizar_card(frame_log_filtros, pad_x=0, pad_y=0)

log_data_inicio_var = ctk.StringVar(value="")
log_data_fim_var = ctk.StringVar(value="")
log_status_var = ctk.StringVar(value="Em andamento")
log_motorista_var = ctk.StringVar(value="Todos")
log_placa_var = ctk.StringVar(value="Todos")

ctk.CTkLabel(
    frame_log_filtros,
    text="Data inicial (dd/mm/aaaa)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=15, pady=(18, 0), sticky="w")
entry_log_data_inicio = ctk.CTkEntry(frame_log_filtros, textvariable=log_data_inicio_var, width=150)
entry_log_data_inicio.grid(row=1, column=0, padx=15, pady=5, sticky="w")
estilizar_entry(entry_log_data_inicio)

ctk.CTkLabel(
    frame_log_filtros,
    text="Data final (dd/mm/aaaa)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=15, pady=(18, 0), sticky="w")
entry_log_data_fim = ctk.CTkEntry(frame_log_filtros, textvariable=log_data_fim_var, width=150)
entry_log_data_fim.grid(row=1, column=1, padx=15, pady=5, sticky="w")
estilizar_entry(entry_log_data_fim)

ctk.CTkLabel(
    frame_log_filtros,
    text="Status",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=15, pady=(18, 0), sticky="w")
combo_log_status = ctk.CTkOptionMenu(
    frame_log_filtros,
    values=LOG_STATUS_OPCOES,
    variable=log_status_var,
    width=170,
)
combo_log_status.grid(row=1, column=2, padx=15, pady=5, sticky="w")
estilizar_optionmenu(combo_log_status)

ctk.CTkLabel(
    frame_log_filtros,
    text="Motorista",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=3, padx=15, pady=(18, 0), sticky="w")
combo_log_motorista = ctk.CTkOptionMenu(
    frame_log_filtros,
    values=["Todos"],
    variable=log_motorista_var,
    width=220,
)
combo_log_motorista.grid(row=1, column=3, padx=15, pady=5, sticky="w")
estilizar_optionmenu(combo_log_motorista)

ctk.CTkLabel(
    frame_log_filtros,
    text="Placa",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=4, padx=15, pady=(18, 0), sticky="w")
combo_log_placa = ctk.CTkOptionMenu(
    frame_log_filtros,
    values=["Todos"],
    variable=log_placa_var,
    width=140,
)
combo_log_placa.grid(row=1, column=4, padx=15, pady=5, sticky="w")
estilizar_optionmenu(combo_log_placa)


def limpar_lista_log():
    for widget in lista_log_registros.winfo_children():
        widget.destroy()


def renderizar_log(registros: list[dict]):
    limpar_lista_log()
    if not registros:
        mostrar_msg_lista(lista_log_registros, "Nenhuma escala encontrada para o filtro.")
        return
    vistos: set[int] = set()
    for item in registros:
        registro_id = item.get("id")
        if registro_id in vistos:
            continue
        vistos.add(registro_id)
        bloco = ctk.CTkFrame(lista_log_registros, fg_color="#FFFFFF", corner_radius=12)
        bloco.pack(fill="x", padx=10, pady=6)
        titulo = f"{item.get('data_br')} • {item.get('rota')} • {item.get('placa')}"
        ctk.CTkLabel(
            bloco,
            text=titulo,
            text_color=COR_AZUL,
            font=("Inter", 15, "bold"),
        ).pack(anchor="w", padx=12, pady=(10, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Motorista: {item.get('motorista')}  |  Ajudante: {item.get('ajudante')}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Saída: {item.get('data_saida_br')}  •  Previsto: {item.get('data_fim_br')}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Duração planejada: {item.get('duracao_planejada')}  |  Efetiva: {item.get('duracao_efetiva')}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 2))
        status_color = "#2E7D32" if item.get("status") == "Finalizado" else "#D35400"
        ctk.CTkLabel(
            bloco,
            text=f"Status: {item.get('status')}",
            text_color=status_color,
            font=("Inter", 13, "bold"),
        ).pack(anchor="w", padx=12, pady=(2, 0))
        if item.get("status_texto"):
            ctk.CTkLabel(
                bloco,
                text=item.get("status_texto"),
                text_color=COR_TEXTO,
            ).pack(anchor="w", padx=12, pady=(0, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Resumo: {item.get('resumo')}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 8))

        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="e", padx=12, pady=(0, 10))
        if item.get("status") == "Em andamento":
            ctk.CTkButton(
                botoes,
                text="Editar",
                fg_color=COR_AZUL_CLARO,
                hover_color=COR_AZUL_HOVER,
                command=lambda registro=item: abrir_modal_ajuste_log(registro),
            ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            fg_color=COR_VERMELHA,
            hover_color=COR_VERMELHA_HOVER,
            command=lambda registro=item: handle_excluir_log(registro),
        ).pack(side="left", padx=4)


def handle_excluir_log(registro: dict):
    global carregamento_dia_selecionado_id
    if not messagebox.askyesno(
        "Excluir escala",
        f"Deseja remover a rota '{registro.get('rota')}' do dia {registro.get('data_br')}?",
    ):
        return
    try:
        remover_carregamento_completo(registro["id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    carregamentos_revisados.discard(registro["id"])
    if carregamento_dia_selecionado_id == registro["id"]:
        carregamento_dia_selecionado_id = None
    aplicar_filtros_log()
    recarregar_disponibilidade()
    refresh_carregamentos_ui()


def validar_intervalo_datas():
    inicio_txt = log_data_inicio_var.get().strip()
    fim_txt = log_data_fim_var.get().strip()
    inicio_iso = data_br_para_iso(inicio_txt) if inicio_txt else None
    if inicio_txt and not inicio_iso:
        messagebox.showerror("Data inválida", "Informe a data inicial no formato dd/mm/aaaa.")
        return None, None
    fim_iso = data_br_para_iso(fim_txt) if fim_txt else None
    if fim_txt and not fim_iso:
        messagebox.showerror("Data inválida", "Informe a data final no formato dd/mm/aaaa.")
        return None, None
    if inicio_iso and fim_iso and inicio_iso > fim_iso:
        messagebox.showerror("Período inválido", "A data inicial não pode ser posterior à data final.")
        return None, None
    return inicio_iso, fim_iso


def aplicar_filtros_log():
    inicio_iso, fim_iso = validar_intervalo_datas()
    if inicio_iso is None and log_data_inicio_var.get().strip():
        return
    if fim_iso is None and log_data_fim_var.get().strip():
        return
    motor_display = log_motorista_var.get()
    placa_display = log_placa_var.get()
    filtros = {
        "data_inicio": inicio_iso,
        "data_fim": fim_iso,
        "status": log_status_var.get(),
        "motorista_id": log_motorista_map.get(motor_display),
        "placa": log_placa_map.get(placa_display),
    }
    registros = consultar_log_carregamentos(filtros)
    log_registros_cache.clear()
    log_registros_cache.extend(registros)
    renderizar_log(registros)


def handle_exportar_log():
    exportar_log_para_excel(log_registros_cache)


btn_log_filtrar = ctk.CTkButton(
    frame_log_filtros,
    text="Filtrar",
    command=aplicar_filtros_log,
)
estilizar_botao(btn_log_filtrar, "primary")
btn_log_filtrar.grid(row=1, column=5, padx=15, pady=5)

btn_log_exportar = ctk.CTkButton(
    frame_log_filtros,
    text="Exportar para Excel",
    command=handle_exportar_log,
)
estilizar_botao(btn_log_exportar, "ghost")
btn_log_exportar.grid(row=1, column=6, padx=15, pady=5)

btn_log_backup = ctk.CTkButton(
    frame_log_filtros,
    text="Restaurar backup",
    command=restaurar_ultimo_backup,
)
estilizar_botao(btn_log_backup, "danger-ghost", small=True)
btn_log_backup.grid(row=1, column=7, padx=10, pady=5)

frame_log_filtros.grid_columnconfigure(7, weight=1)

lista_log_registros = ctk.CTkScrollableFrame(
    aba_log, fg_color="#FFFFFF", corner_radius=12, height=360
)
lista_log_registros.pack(fill="both", expand=True, padx=20, pady=(0, 20))
estilizar_scrollable(lista_log_registros)


def atualizar_filtros_log():
    global log_motorista_map, log_placa_map
    colaboradores = listar_colaboradores()
    opcoes_motoristas = ["Todos"]
    log_motorista_map = {"Todos": None}
    for colaborador in colaboradores:
        funcao = colaborador.get("funcao") or colaborador.get("função") or ""
        label = f"{colaborador['nome']} ({funcao})"
        opcoes_motoristas.append(label)
        log_motorista_map[label] = colaborador["id"]
    combo_log_motorista.configure(values=opcoes_motoristas)
    if log_motorista_var.get() not in opcoes_motoristas:
        combo_log_motorista.set(opcoes_motoristas[0])

    caminhoes = listar_caminhoes()
    opcoes_placas = ["Todos"]
    log_placa_map = {"Todos": None}
    for cam in caminhoes:
        label = cam["placa"].upper()
        opcoes_placas.append(label)
        log_placa_map[label] = label
    combo_log_placa.configure(values=opcoes_placas)
    if log_placa_var.get() not in opcoes_placas:
        combo_log_placa.set(opcoes_placas[0])


def abrir_modal_ajuste_log(registro: dict):
    modal = ctk.CTkToplevel(app)
    modal.title("Ajustar duração")
    modal.geometry("480x420")
    modal.resizable(False, False)
    modal.grab_set()

    container = ctk.CTkFrame(modal, fg_color="#FFFFFF", corner_radius=12)
    container.pack(fill="both", expand=True, padx=16, pady=16)

    ctk.CTkLabel(
        container,
        text=f"Rota: {registro.get('rota')}",
        text_color=COR_AZUL,
        font=("Segoe UI Semibold", 18),
    ).pack(anchor="w", padx=20, pady=(18, 4))
    ctk.CTkLabel(
        container,
        text=f"Planejado: {registro.get('duracao_planejada')}  •  Atual: {registro.get('duracao_efetiva')}",
        text_color=COR_TEXTO,
    ).pack(anchor="w", padx=20, pady=(0, 2))
    ctk.CTkLabel(
        container,
        text=f"Saída: {registro.get('data_saida_br')}  |  Fim previsto: {registro.get('data_fim_br')}",
        text_color=COR_TEXTO,
    ).pack(anchor="w", padx=20, pady=(0, 12))

    ctk.CTkLabel(
        container,
        text="Nova duração efetiva (dias)",
        text_color=COR_TEXTO,
    ).pack(anchor="w", padx=20, pady=(0, 4))
    nova_duracao_var = ctk.StringVar(value=str(registro.get("duracao_efetiva", 0)))
    entry_nova_duracao = ctk.CTkEntry(container, textvariable=nova_duracao_var, width=160)
    entry_nova_duracao.pack(anchor="w", padx=20, pady=(0, 10))
    estilizar_entry(entry_nova_duracao)

    ctk.CTkLabel(
        container,
        text="Observação do ajuste",
        text_color=COR_TEXTO,
    ).pack(anchor="w", padx=20, pady=(8, 4))
    box_obs_ajuste = ctk.CTkTextbox(container, height=110)
    box_obs_ajuste.pack(fill="both", expand=True, padx=20, pady=(0, 10))

    try:
        data_inicio_dt = datetime.strptime(registro["data_saida"], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        data_inicio_dt = date.today()

    def executar_ajuste(nova_dur: int, observacao: str, liberar_agora: bool = False):
        atual = registro.get("duracao_efetiva", 0)
        registrar_ajuste_rota(registro["id"], atual, nova_dur, observacao)
        novo_fim = (data_inicio_dt + timedelta(days=nova_dur)).isoformat()
        atualizar_bloqueios_para_ajuste(
            registro["id"],
            novo_fim,
            liberar_imediato=liberar_agora,
        )
        modal.destroy()
        aplicar_filtros_log()
        recarregar_disponibilidade()

    def confirmar_ajuste():
        try:
            nova_dur = int(nova_duracao_var.get().strip())
        except ValueError:
            messagebox.showerror("Entrada inválida", "Informe um número inteiro para a nova duração.")
            return
        if nova_dur < 0:
            messagebox.showerror("Entrada inválida", "A duração não pode ser negativa.")
            return
        atual = registro.get("duracao_efetiva", 0)
        if nova_dur == atual:
            messagebox.showinfo("Sem alterações", "A duração permanece a mesma.")
            return
        observacao = box_obs_ajuste.get("0.0", "end").strip()
        executar_ajuste(nova_dur, observacao, liberar_agora=False)

    def finalizar_agora():
        dias_passados = max((date.today() - data_inicio_dt).days, 0)
        nova_duracao_var.set(str(dias_passados))
        if not box_obs_ajuste.get("0.0", "end").strip():
            box_obs_ajuste.insert("0.0", "Finalizado manualmente.")
        executar_ajuste(dias_passados, box_obs_ajuste.get("0.0", "end").strip(), liberar_agora=True)

    botoes = ctk.CTkFrame(container, fg_color="transparent")
    botoes.pack(fill="x", padx=20, pady=(0, 10))
    ctk.CTkButton(
        botoes,
        text="Salvar ajuste",
        fg_color=COR_AZUL_CLARO,
        hover_color=COR_AZUL_HOVER,
        command=confirmar_ajuste,
    ).pack(side="left", padx=4)
    ctk.CTkButton(
        botoes,
        text="Finalizar agora",
        fg_color="#00A36C",
        hover_color="#008053",
        command=finalizar_agora,
    ).pack(side="left", padx=4)
    ctk.CTkButton(
        botoes,
        text="Cancelar",
        fg_color=COR_VERMELHA,
        hover_color=COR_VERMELHA_HOVER,
        command=modal.destroy,
    ).pack(side="right", padx=4)


atualizar_filtros_log()
aplicar_filtros_log()

# ---------- CARREGAMENTOS ----------
ctk.CTkLabel(
    carregamentos_container,
    text="CARREGAMENTOS",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    carregamentos_container,
    text="Monte a escala diária e acompanhe motoristas e ajudantes disponíveis.",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).pack(anchor="w", padx=10, pady=(0, 10))

btn_exportar_carreg = ctk.CTkButton(
    carregamentos_container,
    text="Exportar relatório",
    command=lambda: None,
)
estilizar_botao(btn_exportar_carreg, "primary")
btn_exportar_carreg.pack(anchor="e", padx=20, pady=(0, 5))

frame_carreg_data = ctk.CTkFrame(
    carregamentos_container, fg_color=COR_PAINEL, corner_radius=12
)
frame_carreg_data.pack(fill="x", padx=20, pady=(5, 10))
estilizar_card(frame_carreg_data, pad_x=20, pad_y=(5, 10))
label_rotas_auto_info: ctk.CTkLabel | None = None

ctk.CTkLabel(
    frame_carreg_data,
    text="Data (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 5), sticky="w")

carreg_data_var = ctk.StringVar(value=date.today().strftime("%d/%m/%Y"))
entry_carreg_data = ctk.CTkEntry(
    frame_carreg_data, textvariable=carreg_data_var, width=180
)
entry_carreg_data.grid(row=1, column=0, padx=20, pady=(0, 20), sticky="w")
estilizar_entry(entry_carreg_data)

label_carreg_dia_semana = ctk.CTkLabel(
    frame_carreg_data,
    text="",
    text_color=COR_TEXTO,
    font=FONT_TEXTO,
)
label_carreg_dia_semana.grid(row=2, column=0, padx=20, pady=(0, 10), sticky="w")

ctk.CTkLabel(
    frame_carreg_data,
    text="Data de saida (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 5), sticky="w")

carreg_saida_var = ctk.StringVar(value="")
entry_carreg_saida = ctk.CTkEntry(
    frame_carreg_data, textvariable=carreg_saida_var, width=180
)
entry_carreg_saida.grid(row=1, column=1, padx=20, pady=(0, 20), sticky="w")
estilizar_entry(entry_carreg_saida)


def atualizar_data_saida_sugerida(force: bool = False):
    data_iso = validar_data(carreg_data_var.get().strip())
    sugestao = calcular_data_saida_carregamento(data_iso)
    atual_saida = carreg_saida_var.get().strip()
    sugestao_br = data_iso_para_br_entrada(sugestao)
    if force:
        carreg_saida_var.set(sugestao_br)
        return
    if not atual_saida and sugestao:
        carreg_saida_var.set(sugestao_br)


def _on_carreg_data_change(*_):
    atualizar_alerta_carregamento()
    atualizar_data_saida_sugerida()
    atualizar_label_dia_semana_carreg()


carreg_data_var.trace_add("write", _on_carreg_data_change)
atualizar_data_saida_sugerida(force=True)
atualizar_label_dia_semana_carreg()


def aplicar_data_carregamentos():
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        messagebox.showerror(
            "Data inválida", "Informe uma data no formato DD/MM/AAAA."
        )
        return
    recarregar_disponibilidade()
    refresh_carregamentos_ui()
    atualizar_alerta_carregamento()
    preencher_carregamentos_automaticos()


btn_atualizar_data = ctk.CTkButton(
    frame_carreg_data,
    text="Atualizar data",
    command=aplicar_data_carregamentos,
)
estilizar_botao(btn_atualizar_data, "primary")
btn_atualizar_data.grid(row=1, column=2, padx=20, pady=(0, 20))


def handle_duplicar_carregamentos():
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        messagebox.showerror(
            "Data inválida", "Informe uma data no formato DD/MM/AAAA."
        )
        return
    try:
        origem = duplicar_carregamentos_para_data(data_iso)
    except ValueError as exc:
        messagebox.showwarning("Duplicação não realizada", str(exc))
        return
    except LookupError as exc:
        messagebox.showwarning("Duplicação não realizada", str(exc))
        return
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao duplicar", str(exc))
        return

    messagebox.showinfo(
        "Duplicação concluída",
        f"Carregamentos de {data_iso_para_br_entrada(origem)} foram copiados para {data_iso_para_br_entrada(data_iso)}.",
    )
    refresh_carregamentos_ui()
    atualizar_alerta_carregamento()
    recarregar_disponibilidade()


btn_duplicar = ctk.CTkButton(
    frame_carreg_data,
    text="Duplicar dia anterior",
    command=handle_duplicar_carregamentos,
)
estilizar_botao(btn_duplicar, "danger")
btn_duplicar.grid(row=1, column=3, padx=20, pady=(0, 20))
btn_rotas_auto = ctk.CTkButton(
    frame_carreg_data,
    text="Recarregar rotas automáticas",
    command=preencher_carregamentos_automaticos,
)
estilizar_botao(btn_rotas_auto, "primary")
btn_rotas_auto.grid(row=1, column=4, padx=20, pady=(0, 20))
frame_carreg_data.grid_columnconfigure(4, weight=1)
label_rotas_auto_info = ctk.CTkLabel(
    frame_carreg_data,
    text="",
    text_color=COR_TEXTO,
    font=FONT_TEXTO,
)
label_rotas_auto_info.grid(row=3, column=0, columnspan=5, padx=20, pady=(0, 5), sticky="w")
atualizar_mensagem_rotas_auto("")
adicionar_divisor_horizontal(carregamentos_container)
frame_carreg_form = ctk.CTkFrame(
    carregamentos_container, fg_color=COR_PAINEL, corner_radius=12
)
frame_carreg_form.pack(fill="x", padx=20, pady=(0, 15))
estilizar_card(frame_carreg_form, pad_x=20, pad_y=(0, 15))

ctk.CTkLabel(
    frame_carreg_form,
    text="Novo carregamento",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).grid(row=0, column=0, columnspan=6, padx=20, pady=(20, 10), sticky="w")

ctk.CTkLabel(
    frame_carreg_form, text="Nº Rota", text_color=COR_TEXTO, font=FONT_SUBTITULO
).grid(row=1, column=0, padx=20, pady=(5, 0), sticky="w")
entry_rota_num = ctk.CTkEntry(
    frame_carreg_form, placeholder_text="Ex.: Rota 01", width=160
)
entry_rota_num.grid(row=2, column=0, padx=20, pady=5, sticky="w")
estilizar_entry(entry_rota_num)

ctk.CTkLabel(
    frame_carreg_form,
    text="Rota/Destino",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=1, column=1, padx=20, pady=(5, 0), sticky="w")
entry_rota_destino = ctk.CTkEntry(
    frame_carreg_form, placeholder_text="Ex.: Curitiba", width=220
)
entry_rota_destino.grid(row=2, column=1, padx=20, pady=5, sticky="w")
estilizar_entry(entry_rota_destino)

ctk.CTkLabel(
    frame_carreg_form, text="Placa", text_color=COR_TEXTO, font=FONT_SUBTITULO
).grid(row=1, column=2, padx=20, pady=(5, 0), sticky="w")
combo_carreg_placa = ctk.CTkOptionMenu(
    frame_carreg_form,
    values=["Cadastre caminhões"],
    state="disabled",
    width=180,
    command=lambda _: atualizar_alerta_carregamento(),
)
combo_carreg_placa.grid(row=2, column=2, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_carreg_placa)

entry_carreg_placa = ctk.CTkEntry(
    frame_carreg_form, placeholder_text="ABC1D23", width=160
)
entry_carreg_placa.grid(row=2, column=2, padx=20, pady=5, sticky="w")
entry_carreg_placa.grid_remove()
entry_carreg_placa.bind("<KeyRelease>", lambda _: atualizar_alerta_carregamento())
estilizar_entry(entry_carreg_placa)

label_carreg_placa_info = ctk.CTkLabel(
    frame_carreg_form,
    text="Selecione um caminhão ou deixe em branco.",
    text_color=COR_TEXTO,
    font=FONT_TEXTO,
)
label_carreg_placa_info.grid(row=5, column=2, padx=20, pady=(0, 0), sticky="w")

label_carreg_placa_alert = ctk.CTkLabel(
    frame_carreg_form,
    text="",
    text_color=COR_VERMELHA,
    font=FONT_SUBTITULO,
)
label_carreg_placa_alert.grid(row=6, column=2, padx=20, pady=(0, 5), sticky="w")

ctk.CTkLabel(
    frame_carreg_form,
    text="Motorista",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=3, column=0, padx=20, pady=(15, 0), sticky="w")
combo_motorista = ctk.CTkOptionMenu(
    frame_carreg_form, values=["Atualize a data"], state="disabled", width=220
)
combo_motorista.grid(row=4, column=0, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_motorista)

ctk.CTkLabel(
    frame_carreg_form,
    text="Ajudante",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=3, column=1, padx=20, pady=(15, 0), sticky="w")
combo_ajudante = ctk.CTkOptionMenu(
    frame_carreg_form, values=["Atualize a data"], state="disabled", width=220
)
combo_ajudante.grid(row=4, column=1, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_ajudante)
preview_motorista_carreg_label = ctk.CTkLabel(
    frame_carreg_form,
    text="Sem foto",
    width=FOTO_PREVIEW_PADRAO[0],
    height=FOTO_PREVIEW_PADRAO[1],
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=12,
    anchor="center",
)
preview_motorista_carreg_label.grid(row=5, column=0, padx=20, pady=(0, 10), sticky="w")
preview_ajudante_carreg_label = ctk.CTkLabel(
    frame_carreg_form,
    text="Sem foto",
    width=FOTO_PREVIEW_PADRAO[0],
    height=FOTO_PREVIEW_PADRAO[1],
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=12,
    anchor="center",
)
preview_ajudante_carreg_label.grid(row=5, column=1, padx=20, pady=(0, 10), sticky="w")
btn_motorista_ajudante = ctk.CTkButton(
    frame_carreg_form,
    text="Motorista como ajudante: OFF",
    command=alternar_motorista_como_ajudante,
)
btn_motorista_ajudante.grid(row=6, column=1, padx=20, pady=(0, 10), sticky="w")
atualizar_botao_motorista_ajudante()
combo_motorista.configure(
    command=lambda valor: atualizar_preview_combo_colaborador(
        combo_motorista,
        motorista_option_map,
        preview_motorista_carreg_label,
        selecionado=valor,
    )
)
combo_ajudante.configure(
    command=lambda valor: atualizar_preview_combo_colaborador(
        combo_ajudante,
        ajudante_option_map,
        preview_ajudante_carreg_label,
        selecionado=valor,
    )
)

ctk.CTkLabel(
    frame_carreg_form,
    text="Observações",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=3, column=2, padx=20, pady=(15, 0), sticky="w")
combo_observacao = ctk.CTkOptionMenu(
    frame_carreg_form,
    values=OBSERVACAO_OPCOES,
)
combo_observacao.set(OBSERVACAO_OPCOES[0])
combo_observacao.grid(row=4, column=2, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_observacao)

ctk.CTkLabel(
    frame_carreg_form,
    text="Cor da marcacao",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=3, column=4, padx=20, pady=(15, 0), sticky="w")
combo_cor_observacao = ctk.CTkOptionMenu(
    frame_carreg_form,
    values=[label for label, _ in OBS_MARCADORES],
)
combo_cor_observacao.set(OBS_MARCADORES[0][0])
combo_cor_observacao.grid(row=4, column=4, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_cor_observacao)

ctk.CTkLabel(
    frame_carreg_form,
    text="Obs. adicional",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=3, column=3, padx=20, pady=(15, 0), sticky="w")
entry_observacao_extra = ctk.CTkEntry(
    frame_carreg_form,
    placeholder_text="Detalhes complementares",
    width=260,
)
entry_observacao_extra.grid(row=4, column=3, padx=20, pady=5, sticky="w")
estilizar_entry(entry_observacao_extra)

motorista_option_map: dict[str, int | None] = {}
ajudante_option_map: dict[str, int | None] = {}
caminhao_option_map: dict[str, str] = {}


def _caminhoes_disponiveis(indisponiveis: set[str]) -> list[str]:
    valores = [VALOR_SEM_CAMINHAO]
    indis_placas = {placa.upper() for placa in indisponiveis if placa}
    for display, placa in caminhao_option_map.items():
        if display == VALOR_SEM_CAMINHAO:
            continue
        placa_upper = (placa or "").upper()
        if placa_upper and placa_upper in indis_placas:
            continue
        valores.append(display)
    return valores


def atualizar_combo_caminhao_carregamento(indisponiveis: set[str]):
    valores = _caminhoes_disponiveis(indisponiveis)
    if len(valores) > 1:
        atual = combo_carreg_placa.get()
        combo_carreg_placa.configure(values=valores, state="normal")
        combo_carreg_placa.grid()
        entry_carreg_placa.grid_remove()
        combo_carreg_placa.set(atual if atual in valores else valores[0])
        label_carreg_placa_info.configure(
            text="Selecione um caminhão disponível ou deixe o campo em branco."
        )
    else:
        combo_carreg_placa.configure(values=valores, state="disabled")
        combo_carreg_placa.grid_remove()
        entry_carreg_placa.grid()
        label_carreg_placa_info.configure(
            text="Nenhum caminhão disponível. Digite a placa manualmente."
        )


def atualizar_combo_caminhao_oficina(indisponiveis: set[str]):
    valores = _caminhoes_disponiveis(indisponiveis)
    if len(valores) > 1:
        atual = combo_oficina_placa.get()
        combo_oficina_placa.configure(values=valores, state="normal")
        combo_oficina_placa.grid()
        entry_oficina_placa.grid_remove()
        combo_oficina_placa.set(atual if atual in valores else valores[0])
        label_oficina_placa_info.configure(
            text="Selecione um caminhão cadastrado ou deixe em branco."
        )
    else:
        combo_oficina_placa.configure(values=valores, state="disabled")
        combo_oficina_placa.grid_remove()
        entry_oficina_placa.grid()
        label_oficina_placa_info.configure(
            text="Nenhum caminhão disponível. Digite a placa manualmente."
        )


def obter_placa_carregamento() -> str:
    if combo_carreg_placa.winfo_manager() and combo_carreg_placa.cget("state") != "disabled":
        selecionado = combo_carreg_placa.get()
        return caminhao_option_map.get(selecionado, "").upper()
    return entry_carreg_placa.get().strip().upper()


def atualizar_alerta_carregamento():
    data_iso = validar_data(carreg_data_var.get().strip())
    placa = obter_placa_carregamento()
    if data_iso and placa and placa_em_manutencao(placa, data_iso):
        label_carreg_placa_alert.configure(text="Em manutenção")
    else:
        label_carreg_placa_alert.configure(text="")


def obter_placa_oficina() -> str:
    if combo_oficina_placa.winfo_manager() and combo_oficina_placa.cget("state") != "disabled":
        selecionado = combo_oficina_placa.get()
        return caminhao_option_map.get(selecionado, "").upper()
    return entry_oficina_placa.get().strip().upper()


def atualizar_alerta_oficina():
    data_iso = validar_data(oficina_data_var.get().strip())
    placa = obter_placa_oficina()
    if data_iso and placa and placa_em_manutencao(placa, data_iso):
        label_oficina_placa_alert.configure(text="Em manutenção")
    else:
        label_oficina_placa_alert.configure(text="")


def refresh_disponibilidade_carregamentos():
    data_iso = validar_data(carreg_data_var.get().strip())
    valores_motoristas: list[str] = []
    valores_ajudantes: list[str] = []

    ignorar = {"carregamento_id": carregamento_em_edicao_id} if carregamento_em_edicao_id else None

    if not data_iso:
        combo_motorista.configure(values=["Informe uma data"], state="disabled")
        combo_motorista.set("Informe uma data")
        atualizar_preview_combo_colaborador(
            combo_motorista, motorista_option_map, preview_motorista_carreg_label
        )
        combo_ajudante.configure(values=["Informe uma data"], state="disabled")
        combo_ajudante.set("Informe uma data")
        atualizar_preview_combo_colaborador(
            combo_ajudante, ajudante_option_map, preview_ajudante_carreg_label
        )
        motorista_option_map.clear()
        ajudante_option_map.clear()
        atualizar_combo_caminhao_carregamento(set())
        return

    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )

    motoristas = listar_colaboradores_por_funcao("Motorista")
    ajudantes = listar_colaboradores_por_funcao("Ajudante")
    motorista_option_map.clear()
    ajudante_option_map.clear()
    motorista_option_map[VALOR_SEM_MOTORISTA] = None
    ajudante_option_map[VALOR_SEM_AJUDANTE] = None

    for col in motoristas:
        if col["id"] in indis_colaboradores:
            continue
        display = f"{col['nome']} (#{col['id']})"
        motorista_option_map[display] = col["id"]
        valores_motoristas.append(display)

    if permitir_motorista_ajudante:
        for col in motoristas:
            if col["id"] in indis_colaboradores:
                continue
            display = f"{col['nome']} (#{col['id']}) {MOTORISTA_AJUDANTE_TAG}"
            ajudante_option_map[display] = col["id"]
            valores_ajudantes.append(display)
    else:
        for col in ajudantes:
            if col["id"] in indis_colaboradores:
                continue
            display = f"{col['nome']} (#{col['id']})"
            ajudante_option_map[display] = col["id"]
            valores_ajudantes.append(display)

    if valores_motoristas:
        valores_motoristas_combo = [VALOR_SEM_MOTORISTA] + valores_motoristas
        atual_motorista = combo_motorista.get()
        combo_motorista.configure(values=valores_motoristas_combo, state="normal")
        combo_motorista.set(
            atual_motorista if atual_motorista in valores_motoristas_combo else VALOR_SEM_MOTORISTA
        )
        atualizar_preview_combo_colaborador(
            combo_motorista, motorista_option_map, preview_motorista_carreg_label
        )
    else:
        motorista_option_map[AVISO_NENHUM_COLAB] = None
        combo_motorista.configure(values=[AVISO_NENHUM_COLAB], state="normal")
        combo_motorista.set(AVISO_NENHUM_COLAB)
        atualizar_preview_combo_colaborador(
            combo_motorista, motorista_option_map, preview_motorista_carreg_label
        )

    if valores_ajudantes:
        valores_ajudantes_combo = [VALOR_SEM_AJUDANTE] + valores_ajudantes
        atual_ajudante = combo_ajudante.get()
        combo_ajudante.configure(values=valores_ajudantes_combo, state="normal")
        combo_ajudante.set(
            atual_ajudante if atual_ajudante in valores_ajudantes_combo else VALOR_SEM_AJUDANTE
        )
        atualizar_preview_combo_colaborador(
            combo_ajudante, ajudante_option_map, preview_ajudante_carreg_label
        )
    else:
        ajudante_option_map[AVISO_NENHUM_COLAB] = None
        combo_ajudante.configure(values=[AVISO_NENHUM_COLAB], state="normal")
        combo_ajudante.set(AVISO_NENHUM_COLAB)
        atualizar_preview_combo_colaborador(
            combo_ajudante, ajudante_option_map, preview_ajudante_carreg_label
        )

    atualizar_combo_caminhao_carregamento(disponibilidade.get("caminhoes", set()))
    atualizar_alerta_carregamento()


def atualizar_combo_carregamento_dia(registros: list[dict] | None, data_valida: bool):
    if combo_carregamento_dia is None:
        return
    carregamento_dia_map.clear()
    if not data_valida:
        combo_carregamento_dia.configure(values=[CARREGAMENTO_DIA_SEM_DATA], state="disabled")
        combo_carregamento_dia.set(CARREGAMENTO_DIA_SEM_DATA)
        return
    if not registros:
        combo_carregamento_dia.configure(values=[CARREGAMENTO_DIA_VAZIO], state="disabled")
        combo_carregamento_dia.set(CARREGAMENTO_DIA_VAZIO)
        return

    valores: list[str] = []
    for item in registros:
        rota = item.get("rota") or DISPLAY_VAZIO
        placa = item.get("placa") or DISPLAY_VAZIO
        motorista = item.get("motorista_nome") or DISPLAY_VAZIO
        item_id = item.get("id")
        marcador = f"{CARREGAMENTO_DIA_MARCADOR} " if item_id in carregamentos_revisados else ""
        display = f"{marcador}{rota} | {placa} | {motorista}"
        if display in carregamento_dia_map:
            display = f"{display} (#{item_id})"
        carregamento_dia_map[display] = item
        valores.append(display)

    valores = [CARREGAMENTO_DIA_PLACEHOLDER] + valores
    atual = combo_carregamento_dia.get()
    combo_carregamento_dia.configure(values=valores, state="normal")
    selecionado = None
    if carregamento_dia_selecionado_id is not None:
        for display, item in carregamento_dia_map.items():
            if item.get("id") == carregamento_dia_selecionado_id:
                selecionado = display
                break
    if selecionado:
        combo_carregamento_dia.set(selecionado)
    elif atual in valores:
        combo_carregamento_dia.set(atual)
    else:
        combo_carregamento_dia.set(CARREGAMENTO_DIA_PLACEHOLDER)


def selecionar_carregamento_dia(valor: str):
    global carregamento_dia_selecionado_id
    if valor in (
        CARREGAMENTO_DIA_PLACEHOLDER,
        CARREGAMENTO_DIA_SEM_DATA,
        CARREGAMENTO_DIA_VAZIO,
    ):
        carregamento_dia_selecionado_id = None
        return
    item = carregamento_dia_map.get(valor)
    if not item:
        return
    carregamento_dia_selecionado_id = item.get("id")
    preparar_edicao_carregamento(item)


def limpar_modificacoes_carregamentos_dia():
    global carregamento_dia_selecionado_id
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        messagebox.showerror("Data invalida", "Informe uma data no formato DD/MM/AAAA.")
        return
    if not messagebox.askyesno(
        "Limpar alteracoes",
        "Isso vai remover todos os carregamentos deste dia e recarregar as rotas semanais. Deseja continuar?",
    ):
        return
    registros = listar_carregamentos(data_iso)
    for item in registros:
        try:
            remover_carregamento_completo(item["id"])
        except sqlite3.Error:
            continue
    carregamentos_revisados.clear()
    carregamento_dia_selecionado_id = None
    limpar_form_carregamento()
    refresh_carregamentos_ui()
    recarregar_disponibilidade()
    preencher_carregamentos_automaticos()
    if "aplicar_filtros_log" in globals():
        aplicar_filtros_log()


def refresh_carregamentos_ui():
    data_iso = validar_data(carreg_data_var.get().strip())
    for widget in lista_carregamentos.winfo_children():
        widget.destroy()

    if not data_iso:
        atualizar_combo_carregamento_dia(None, data_valida=False)
        mostrar_msg_lista(
            lista_carregamentos,
            "Informe uma data válida para listar os carregamentos.",
        )
        return

    registros = listar_carregamentos(data_iso)
    atualizar_combo_carregamento_dia(registros, data_valida=True)
    if not registros:
        mostrar_msg_lista(lista_carregamentos)
        return

    data_legenda = datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    for item in registros:
        placa_texto = (
            item["placa"].upper()
            if item.get("placa")
            else DISPLAY_VAZIO
        )
        motorista_nome = item.get("motorista_nome") or DISPLAY_VAZIO
        ajudante_nome = formatar_ajudante_nome(
            item.get("ajudante_nome") or DISPLAY_VAZIO,
            item.get("ajudante_id"),
        )
        observacao_texto = combinar_observacoes(item.get("observacao"), item.get("observacao_extra"))
        cor_marcador = ajustar_cor_marcador(item.get("observacao_cor"))
        bloco = ctk.CTkFrame(lista_carregamentos, fg_color=COR_PAINEL, corner_radius=12)
        bloco.pack(fill="x", padx=10, pady=6)
        aplicar_hover_cartao(bloco, COR_PAINEL)
        ctk.CTkLabel(
            bloco,
            text=f"{item['rota']} • {placa_texto}",
            text_color=COR_AZUL,
            font=FONT_TITULO,
        ).pack(anchor="w", padx=12, pady=(8, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Motorista: {motorista_nome} | Ajudante: {ajudante_nome} | Data: {data_legenda}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 4))
        obs_frame = ctk.CTkFrame(
            bloco,
            fg_color=cor_marcador or "transparent",
            corner_radius=8,
        )
        obs_frame.pack(fill="x", padx=8, pady=(0, 8))
        ctk.CTkLabel(
            obs_frame,
            text=f"Obs.: {observacao_texto}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
            anchor="w",
        ).pack(anchor="w", padx=10, pady=6)
        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="e", padx=10, pady=(0, 6))
        ctk.CTkButton(
            botoes,
            text="Editar",
            width=80,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda item=item: preparar_edicao_carregamento(item),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Duplicar",
            width=90,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda item=item: duplicar_carregamento_dia(item),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            width=80,
            fg_color="#FFFFFF",
            text_color=COR_VERMELHA,
            border_color=COR_VERMELHA,
            border_width=1,
            hover_color="#FCE4E4",
            command=lambda item=item: handle_excluir_carregamento(item),
        ).pack(side="left", padx=4)


def _linha_relatorio_carregamento(item: dict) -> tuple[list[str], str | None]:
    rota_bruta = item.get("rota") or DISPLAY_VAZIO
    numero = DISPLAY_VAZIO
    destino = rota_bruta
    if " - " in rota_bruta:
        numero_part, destino_part = rota_bruta.split(" - ", 1)
        numero = numero_part.strip() or DISPLAY_VAZIO
        destino = destino_part.strip() or DISPLAY_VAZIO
    placa_texto = item.get("placa")
    placa_formatada = placa_texto.upper() if placa_texto else DISPLAY_VAZIO
    ajudante_nome = formatar_ajudante_nome(
        item.get("ajudante_nome") or DISPLAY_VAZIO,
        item.get("ajudante_id"),
    )
    obs_padrao = (item.get("observacao") or "").strip()
    obs_extra = (item.get("observacao_extra") or "").strip()
    if obs_padrao == "0":
        obs_padrao = ""
    if obs_padrao and obs_extra:
        observacao_relatorio = f"{obs_padrao} - {obs_extra}"
    elif obs_padrao:
        observacao_relatorio = obs_padrao
    elif obs_extra:
        observacao_relatorio = obs_extra
    else:
        observacao_relatorio = ""
    valores = [
        numero,
        placa_formatada,
        destino,
        item.get("motorista_nome") or DISPLAY_VAZIO,
        ajudante_nome,
        observacao_relatorio,
    ]
    return valores, (item.get("observacao_cor") or "").strip() or None


def gerar_relatorio_moderno(
    arquivo_stub: str,
    titulo_header: str,
    linha_principal_rotulo: str,
    data_principal_iso: str | None,
    linha_secundaria_rotulo: str | None,
    data_secundaria_iso: str | None,
    total_legenda: str,
    colunas: list[str],
    col_widths: list[float],
    linhas: list[list[str]],
    col_align_center: set[int] | None = None,
    highlight_col: int | None = None,
    highlight_colors: list[str | None] | None = None,
    arquivo_data_iso: str | None = None,
    fallback_highlight: bool = True,
):
    if not linhas:
        linhas = [[DISPLAY_VAZIO for _ in colunas]]

    col_align_center = col_align_center or set()
    highlight_colors = highlight_colors or []

    largura = 1920
    header_altura = 140
    rodape_altura = 40
    margem = 40
    table_header_altura = 48
    linha_altura_base = 48
    table_width = largura - 2 * margem
    col_px = [int(table_width * proporcao) for proporcao in col_widths]

    font_titulo = carregar_fonte(29, bold=True)
    font_info = carregar_fonte(21, bold=True)
    font_info_secundario = carregar_fonte(19, bold=True)
    font_header_table = carregar_fonte(18, bold=True)
    font_table = carregar_fonte(16, bold=True)
    font_footer = carregar_fonte(13, bold=True)
    line_height = getattr(font_table, "size", 13) + 6

    dummy_draw = ImageDraw.Draw(Image.new("RGB", (1, 1)))

    def quebrar_texto(texto: str | None, largura_coluna: int) -> list[str]:
        if texto is None:
            conteudo = DISPLAY_VAZIO
        else:
            conteudo = str(texto)
        conteudo = conteudo.strip()
        if conteudo == "":
            return [""]
        palavras = conteudo.split()
        linhas_resultado: list[str] = []
        atual = ""
        while palavras:
            palavra = palavras.pop(0)
            tentativa = (atual + " " + palavra).strip()
            largura_t, _ = medir_texto(dummy_draw, tentativa, font_table)
            if largura_t <= largura_coluna - 24:
                atual = tentativa
            else:
                if atual:
                    linhas_resultado.append(atual)
                atual = palavra
        if atual:
            linhas_resultado.append(atual)
        return linhas_resultado or [DISPLAY_VAZIO]

    linhas_processadas: list[tuple[int, list[list[str]], str | None]] = []
    total_altura_tabela = table_header_altura
    for idx_linha, linha in enumerate(linhas):
        celulas = []
        altura_linha = linha_altura_base
        for col_idx in range(len(colunas)):
            valor = linha[col_idx] if col_idx < len(linha) else ""
            largura_coluna = col_px[col_idx]
            linhas_texto = quebrar_texto(str(valor) if valor is not None else DISPLAY_VAZIO, largura_coluna)
            altura_linha = max(altura_linha, len(linhas_texto) * line_height + 16)
            celulas.append(linhas_texto)
        cor_marcador = highlight_colors[idx_linha] if idx_linha < len(highlight_colors) else None
        linhas_processadas.append((altura_linha, celulas, cor_marcador))
        total_altura_tabela += altura_linha

    altura_minima = 1080
    altura_calculada = header_altura + 25 + total_altura_tabela + rodape_altura + 25
    altura = max(altura_minima, int(altura_calculada))
    tabela_top = header_altura + 25

    imagem = Image.new("RGB", (largura, altura), "#F1F3F6")
    draw = ImageDraw.Draw(imagem)
    grad = criar_gradiente_horizontal(largura, header_altura, COR_AZUL, COR_AZUL_GRADIENTE_FIM)
    imagem.paste(grad, (0, 0))

    logo_pos = (margem, 35)
    if LOGO_PATH.exists():
        try:
            with Image.open(LOGO_PATH) as logo_img:
                logo_rel = logo_img.convert("RGBA").resize((80, 80))
            imagem.paste(logo_rel, logo_pos, logo_rel)
        except OSError:
            draw.text((logo_pos[0], logo_pos[1] + 20), "JR", fill="#FFFFFF", font=font_titulo)
    else:
        draw.text((logo_pos[0], logo_pos[1] + 20), "JR", fill="#FFFFFF", font=font_titulo)

    titulo_x = logo_pos[0] + 110
    draw.text((titulo_x, 32), titulo_header, fill="#FFFFFF", font=font_titulo)

    label_principal = (linha_principal_rotulo or "Informação").strip().upper()
    texto_principal = f"{label_principal}: {data_iso_para_extenso(data_principal_iso)}"
    draw.text((titulo_x, 70), texto_principal, fill="#FFFFFF", font=font_info)

    if linha_secundaria_rotulo:
        label_secundario = linha_secundaria_rotulo.strip().upper()
        texto_secundario = f"{label_secundario}: {data_iso_para_extenso(data_secundaria_iso or data_principal_iso)}"
        draw.text((titulo_x, 98), texto_secundario, fill="#FFFFFF", font=font_info)

    if total_legenda:
        texto_total = total_legenda
        texto_total_w, _ = medir_texto(draw, texto_total, font_info_secundario)
        draw.text(
            (largura - margem - texto_total_w, header_altura - 30),
            texto_total,
            fill="#E8EEF8",
            font=font_info_secundario,
        )

    y = tabela_top
    x = margem
    header_bg = "#0A375F"
    for idx, coluna in enumerate(colunas):
        largura_coluna = col_px[idx]
        draw.rectangle([x, y, x + largura_coluna, y + table_header_altura], fill=header_bg)
        texto_header = coluna.upper()
        texto_w, texto_h = medir_texto(draw, texto_header, font_header_table)
        if idx in col_align_center:
            texto_x = x + (largura_coluna - texto_w) / 2
        else:
            texto_x = x + 12
        texto_y = y + (table_header_altura - texto_h) / 2
        draw.text((texto_x, texto_y), texto_header, fill="#FFFFFF", font=font_header_table)
        x += largura_coluna
    y += table_header_altura

    linha_bg_1 = "#D9E7FB"
    linha_bg_2 = "#CFE0F7"
    borda_cor = "#6B84A6"
    borda_largura = 2

    for idx_linha, (altura_linha, celulas_processadas, cor_marcador_raw) in enumerate(linhas_processadas):
        x = margem
        bg = linha_bg_1 if idx_linha % 2 == 0 else linha_bg_2
        cor_marcador = ajustar_cor_marcador(cor_marcador_raw)
        for col_idx, linhas_texto in enumerate(celulas_processadas):
            largura_coluna = col_px[col_idx]
            cel_bg = bg
            if highlight_col is not None and col_idx == highlight_col:
                if cor_marcador:
                    cel_bg = cor_marcador
                elif fallback_highlight:
                    primeira_linha = (linhas_texto[0] or "").strip()
                    if primeira_linha and primeira_linha not in ("0", DISPLAY_VAZIO):
                        cel_bg = "#FFF2B2"
            draw.rectangle([x, y, x + largura_coluna, y + altura_linha], fill=cel_bg)
            for offset in range(borda_largura):
                draw.rectangle(
                    [
                        x + offset,
                        y + offset,
                        x + largura_coluna - offset,
                        y + altura_linha - offset,
                    ],
                    outline=borda_cor,
                )
            text_y = y + 8
            for linha_texto in linhas_texto:
                texto = linha_texto
                texto_w, _ = medir_texto(draw, texto, font_table)
                if col_idx in col_align_center:
                    texto_x = x + (largura_coluna - texto_w) / 2
                else:
                    texto_x = x + 12
                draw.text((texto_x, text_y), texto, fill=COR_TEXTO, font=font_table)
                text_y += line_height
            x += largura_coluna
        y += altura_linha

    draw.text(
        (margem, altura - rodape_altura + 6),
        "JR Ferragens e Madeiras | Sistema JR Escala",
        fill=COR_TEXTO,
        font=font_footer,
    )

    arquivo_base = arquivo_data_iso or data_principal_iso or date.today().isoformat()
    nome_arquivo = f"relatorio_JR_{arquivo_stub}_{arquivo_base}.png"
    caminho = RELATORIOS_DIR / nome_arquivo
    imagem.save(caminho)
    messagebox.showinfo(
        "Relatório gerado",
        f"Relatório salvo em: {caminho}",
    )


def _desenhar_relatorio_carregamentos(
    data_carreg_iso: str,
    data_saida_iso: str,
    linhas: list[list[str]],
    total_registros: int,
    cores_obs: list[str | None] | None = None,
):
    colunas = ["Nº", "Placa", "Rota", "Motorista", "Ajudante", "Obs."]
    col_widths = [0.08, 0.12, 0.25, 0.2, 0.2, 0.15]
    gerar_relatorio_moderno(
        arquivo_stub="carregamentos",
        titulo_header="JR Escala - Carregamentos",
        linha_principal_rotulo="Carregamento",
        data_principal_iso=data_carreg_iso,
        linha_secundaria_rotulo="Rotas",
        data_secundaria_iso=data_saida_iso,
        total_legenda=f"Total de rotas: {total_registros}",
        colunas=colunas,
        col_widths=col_widths,
        linhas=linhas,
        col_align_center={0, 1},
        highlight_col=len(colunas) - 1,
        highlight_colors=cores_obs or [],
        arquivo_data_iso=data_carreg_iso,
        fallback_highlight=False,
    )


def exportar_relatorio_carregamentos():
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        messagebox.showerror(
            "Data invalida", "Informe uma data no formato DD/MM/AAAA."
        )
        return
    registros = listar_carregamentos(data_iso)
    linhas: list[list[str]] = []
    cores_obs: list[str | None] = []
    for item in registros:
        valores, cor = _linha_relatorio_carregamento(item)
        linhas.append(valores)
        cores_obs.append(cor)
    total_registros = len(registros)

    data_saida_raw = carreg_saida_var.get().strip()
    if data_saida_raw:
        data_saida_iso = validar_data(data_saida_raw)
        if not data_saida_iso:
            messagebox.showerror(
                "Data invalida",
                "Informe uma data de saida valida no formato DD/MM/AAAA.",
            )
            return
    else:
        data_saida_iso = calcular_data_saida_carregamento(data_iso) or data_iso
        carreg_saida_var.set(data_iso_para_br_entrada(data_saida_iso))

    if data_saida_iso and data_saida_iso < data_iso:
        messagebox.showerror(
            "Periodo invalido",
            "A data de saida nao pode ser anterior a data.",
        )
        return

    _desenhar_relatorio_carregamentos(data_iso, data_saida_iso, linhas, total_registros, cores_obs)


btn_exportar_carreg.configure(command=exportar_relatorio_carregamentos)


def handle_salvar_carregamento():
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        messagebox.showerror(
            "Data invalida", "Informe uma data no formato DD/MM/AAAA."
        )
        return

    data_saida_bruta = carreg_saida_var.get().strip()
    if data_saida_bruta:
        data_saida_iso = validar_data(data_saida_bruta)
        if not data_saida_iso:
            messagebox.showerror(
                "Data de saída inválida",
                "Informe uma data válida no formato DD/MM/AAAA para a saída.",
            )
            return
    else:
        data_saida_iso = calcular_data_saida_carregamento(data_iso) or data_iso
        carreg_saida_var.set(data_iso_para_br_entrada(data_saida_iso))

    if data_saida_iso and data_saida_iso < data_iso:
        messagebox.showerror(
            "Periodo invalido",
            "A data de saida nao pode ser anterior a data.",
        )
        return

    rota_num = entry_rota_num.get().strip()
    rota_destino = entry_rota_destino.get().strip()
    placa = obter_placa_carregamento()
    observacao = combo_observacao.get().strip()
    observacao_extra = entry_observacao_extra.get().strip()
    observacao_cor_label = combo_cor_observacao.get()
    observacao_cor = OBS_MARCADORES_MAP.get(observacao_cor_label, "")
    motorista_label = combo_motorista.get()
    ajudante_label = combo_ajudante.get()

    if not rota_num or not rota_destino:
        messagebox.showerror(
            "Campos obrigatorios", "Informe rota e destino."
        )
        return

    motorista_id = motorista_option_map.get(motorista_label)
    ajudante_id = ajudante_option_map.get(ajudante_label)

    campos_vazios: list[str] = []
    if not motorista_id:
        campos_vazios.append("Motorista")
    if not ajudante_id:
        campos_vazios.append("Ajudante")
    if not placa:
        campos_vazios.append("Caminhão/Placa")
    if not observacao:
        campos_vazios.append("Obs. padrão")

    if not confirmar_campos_vazios(campos_vazios):
        return

    rota_texto = f"{rota_num} - {rota_destino}"
    placa_db = placa if placa else None

    ignorar = {"carregamento_id": carregamento_em_edicao_id} if carregamento_em_edicao_id else None
    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )
    if motorista_id and motorista_id in indis_colaboradores:
        messagebox.showwarning(
            "Colaborador indisponivel",
            "Este motorista ja esta alocado em outra atividade nesta data.",
        )
        return

    if ajudante_id and ajudante_id in indis_colaboradores:
        messagebox.showwarning(
            "Colaborador indisponivel",
            "Este ajudante ja esta alocado em outra atividade nesta data.",
        )
        return

    indis_caminhoes = {placa.upper() for placa in disponibilidade.get("caminhoes", set())}
    if placa_db and placa_db.upper() in indis_caminhoes:
        messagebox.showwarning(
            "Caminhao indisponivel",
            "Este caminhao ja esta alocado em outra atividade nesta data.",
        )
        return

    try:
        if carregamento_em_edicao_id:
            atualizar_carregamento(
                carregamento_em_edicao_id,
                data_iso,
                data_saida_iso,
                rota_texto,
                placa_db,
                motorista_id,
                ajudante_id,
                observacao,
                observacao_extra,
                observacao_cor,
            )
            remover_bloqueios_por_carregamento(carregamento_em_edicao_id)
            criar_bloqueios_para_carregamento(
                carregamento_em_edicao_id,
                data_iso,
                [motorista_id, ajudante_id],
                observacao,
            )
            carregamentos_revisados.add(carregamento_em_edicao_id)
            messagebox.showinfo("Atualizado", "Carregamento atualizado.")
        else:
            novo_id = salvar_carregamento(
                data_iso,
                rota_texto,
                placa_db,
                motorista_id,
                ajudante_id,
                observacao,
                observacao_extra,
                observacao_cor,
                data_saida=data_saida_iso,
            )
            criar_bloqueios_para_carregamento(
                novo_id,
                data_iso,
                [motorista_id, ajudante_id],
                observacao,
            )
            messagebox.showinfo("Sucesso", "Carregamento salvo.")
    except ValueError as exc:
        messagebox.showerror("Validação", str(exc))
        return
    except sqlite3.IntegrityError:
        messagebox.showwarning(
            "Placa duplicada",
            "Ja existe esta placa para esta data.",
        )
        return
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return

    criar_backup_automatico()
    limpar_form_carregamento()
    atualizar_alerta_carregamento()
    refresh_carregamentos_ui()
    recarregar_disponibilidade()
    aplicar_filtros_log()


btn_salvar_carreg = ctk.CTkButton(
    frame_carreg_form,
    text="Salvar carregamento",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=handle_salvar_carregamento,
)
btn_salvar_carreg.grid(row=5, column=3, padx=20, pady=5)
btn_cancelar_edicao_carreg = ctk.CTkButton(
    frame_carreg_form,
    text="Cancelar edição",
    fg_color=COR_VERMELHA,
    hover_color=COR_AZUL,
    state="disabled",
    command=lambda: limpar_form_carregamento(),
)
btn_cancelar_edicao_carreg.grid(row=5, column=4, padx=20, pady=5)
frame_carreg_form.grid_columnconfigure(0, weight=1)
frame_carreg_form.grid_columnconfigure(1, weight=1)
frame_carreg_form.grid_columnconfigure(2, weight=1)
frame_carreg_form.grid_columnconfigure(3, weight=1)
frame_carreg_form.grid_columnconfigure(4, weight=1)

ctk.CTkLabel(
    carregamentos_container,
    text="CARREGAMENTOS DO DIA",
    text_color=COR_TEXTO,
    font=("Inter", 18, "bold"),
).pack(anchor="w", padx=20, pady=(5, 0))

frame_carreg_selecao = ctk.CTkFrame(carregamentos_container, fg_color="transparent")
frame_carreg_selecao.pack(fill="x", padx=20, pady=(4, 0))
ctk.CTkLabel(
    frame_carreg_selecao,
    text="Selecionar carregamento",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, pady=(0, 4), sticky="w")
combo_carregamento_dia = ctk.CTkOptionMenu(
    frame_carreg_selecao,
    values=[CARREGAMENTO_DIA_SEM_DATA],
    state="disabled",
    width=460,
    command=selecionar_carregamento_dia,
)
combo_carregamento_dia.grid(row=1, column=0, sticky="w")
combo_carregamento_dia.set(CARREGAMENTO_DIA_SEM_DATA)
estilizar_optionmenu(combo_carregamento_dia)
btn_limpar_revisados = ctk.CTkButton(
    frame_carreg_selecao,
    text="Limpar alterações",
    command=limpar_modificacoes_carregamentos_dia,
)
btn_limpar_revisados.grid(row=1, column=1, padx=(12, 0), sticky="w")
estilizar_botao(btn_limpar_revisados, "ghost", small=True)
frame_carreg_selecao.grid_columnconfigure(0, weight=1)

lista_carregamentos = ctk.CTkScrollableFrame(
    carregamentos_container, fg_color=COR_PAINEL, corner_radius=12, height=260
)
lista_carregamentos.pack(fill="both", expand=True, padx=20, pady=(8, 20))
estilizar_scrollable(lista_carregamentos)
# ---------- OFICINAS ----------
ctk.CTkLabel(
    aba_oficinas,
    text="OFICINAS",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_oficinas,
    text="Controle datas de manutenção e mantenha a frota organizada.",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).pack(anchor="w", padx=10, pady=(0, 10))

btn_exportar_oficinas = ctk.CTkButton(
    aba_oficinas,
    text="Exportar relatório",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=lambda: exportar_relatorio_oficinas(),
)
btn_exportar_oficinas.pack(anchor="e", padx=20, pady=(0, 5))

frame_oficina_form = ctk.CTkFrame(
    aba_oficinas, fg_color="#FFFFFF", corner_radius=12
)
frame_oficina_form.pack(fill="x", padx=10, pady=(5, 12))

oficina_data_var = ctk.StringVar(value=date.today().strftime("%d/%m/%Y"))
oficina_data_saida_var = ctk.StringVar(value="")

ctk.CTkLabel(
    frame_oficina_form,
    text="Data (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
entry_oficina_data = ctk.CTkEntry(
    frame_oficina_form, textvariable=oficina_data_var, width=180
)
entry_oficina_data.grid(row=1, column=0, padx=20, pady=5, sticky="w")
estilizar_entry(entry_oficina_data)

ctk.CTkLabel(
    frame_oficina_form,
    text="Data de saida (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
entry_oficina_data_saida = ctk.CTkEntry(
    frame_oficina_form, textvariable=oficina_data_saida_var, width=180
)
entry_oficina_data_saida.grid(row=1, column=1, padx=20, pady=5, sticky="w")
estilizar_entry(entry_oficina_data_saida)

def atualizar_data_saida_oficina(force: bool = False):
    data_iso = validar_data(oficina_data_var.get().strip())
    sugestao = calcular_data_saida_padrao(data_iso)
    atual = oficina_data_saida_var.get().strip()
    if force or not atual:
        oficina_data_saida_var.set(data_iso_para_br_entrada(sugestao))


atualizar_data_saida_oficina(force=True)

ctk.CTkLabel(
    frame_oficina_form,
    text="Motorista",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=20, pady=(20, 0), sticky="w")
combo_oficina_motorista = ctk.CTkOptionMenu(
    frame_oficina_form, values=["Cadastre motoristas"], state="disabled", width=220
)
combo_oficina_motorista.grid(row=1, column=2, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_oficina_motorista)
preview_oficina_motorista_label = ctk.CTkLabel(
    frame_oficina_form,
    text="Sem foto",
    width=FOTO_PREVIEW_PADRAO[0],
    height=FOTO_PREVIEW_PADRAO[1],
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=12,
    anchor="center",
)
preview_oficina_motorista_label.grid(row=2, column=2, padx=20, pady=(0, 10), sticky="w")
combo_oficina_motorista.configure(
    command=lambda valor: atualizar_preview_combo_colaborador(
        combo_oficina_motorista,
        oficina_motoristas_map,
        preview_oficina_motorista_label,
        selecionado=valor,
    )
)

ctk.CTkLabel(
    frame_oficina_form,
    text="Placa",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=3, padx=20, pady=(20, 0), sticky="w")
combo_oficina_placa = ctk.CTkOptionMenu(
    frame_oficina_form,
    values=["Cadastre caminhões"],
    state="disabled",
    width=180,
    command=lambda _: atualizar_alerta_oficina(),
)
combo_oficina_placa.grid(row=1, column=3, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_oficina_placa)

entry_oficina_placa = ctk.CTkEntry(
    frame_oficina_form, placeholder_text="ABC1D23", width=160
)
entry_oficina_placa.grid(row=1, column=3, padx=20, pady=5, sticky="w")
entry_oficina_placa.grid_remove()
entry_oficina_placa.bind("<KeyRelease>", lambda _: atualizar_alerta_oficina())
estilizar_entry(entry_oficina_placa)

label_oficina_placa_info = ctk.CTkLabel(
    frame_oficina_form,
    text="Selecione um caminhao cadastrado.",
    text_color=COR_TEXTO,
    font=FONT_TEXTO,
)
label_oficina_placa_info.grid(row=2, column=3, padx=20, pady=(0, 0), sticky="w")

label_oficina_placa_alert = ctk.CTkLabel(
    frame_oficina_form,
    text="",
    text_color=COR_VERMELHA,
    font=FONT_SUBTITULO,
)
label_oficina_placa_alert.grid(row=3, column=3, padx=20, pady=(0, 5), sticky="w")

ctk.CTkLabel(
    frame_oficina_form,
    text="Observações",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=4, padx=20, pady=(20, 0), sticky="w")
entry_oficina_obs = ctk.CTkEntry(
    frame_oficina_form, placeholder_text="MANUTENCAO / AUTO FREIOS", width=260
)
entry_oficina_obs.grid(row=1, column=4, padx=20, pady=5, sticky="w")
estilizar_entry(entry_oficina_obs)
entry_oficina_obs_extra = ctk.CTkEntry(
    frame_oficina_form, placeholder_text="Detalhes adicionais", width=260
)
ctk.CTkLabel(
    frame_oficina_form,
    text="Obs. adicional",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=2, column=4, padx=20, pady=(10, 0), sticky="w")
entry_oficina_obs_extra.grid(row=3, column=4, padx=20, pady=5, sticky="w")
estilizar_entry(entry_oficina_obs_extra)
ctk.CTkLabel(
    frame_oficina_form,
    text="Cor da marcacao",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=5, padx=20, pady=(20, 0), sticky="w")
combo_cor_oficina = ctk.CTkOptionMenu(
    frame_oficina_form,
    values=[label for label, _ in OBS_MARCADORES],
    width=160,
)
combo_cor_oficina.set(OBS_MARCADORES[0][0])
combo_cor_oficina.grid(row=1, column=5, padx=20, pady=5, sticky="w")
estilizar_optionmenu(combo_cor_oficina)

oficina_motoristas_map: dict[str, int | None] = {}


def refresh_oficina_motoristas():
    data_iso = validar_data(oficina_data_var.get().strip())
    ignorar = {"oficina_id": oficina_em_edicao_id} if oficina_em_edicao_id else None
    oficina_motoristas_map.clear()

    if not data_iso:
        combo_oficina_motorista.configure(values=["Informe uma data"], state="disabled")
        combo_oficina_motorista.set("Informe uma data")
        atualizar_preview_combo_colaborador(
            combo_oficina_motorista, oficina_motoristas_map, preview_oficina_motorista_label
        )
        return

    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )

    motoristas = listar_colaboradores_por_funcao("Motorista")
    valores = [VALOR_SEM_MOTORISTA]
    oficina_motoristas_map[VALOR_SEM_MOTORISTA] = None
    for col in motoristas:
        if col["id"] in indis_colaboradores:
            continue
        display = f"{col['nome']} (#{col['id']})"
        valores.append(display)
        oficina_motoristas_map[display] = col["id"]

    if len(valores) == 1:
        valores = [AVISO_NENHUM_COLAB]
        oficina_motoristas_map[AVISO_NENHUM_COLAB] = None
        combo_oficina_motorista.configure(values=valores, state="normal")
        combo_oficina_motorista.set(AVISO_NENHUM_COLAB)
        atualizar_preview_combo_colaborador(
            combo_oficina_motorista, oficina_motoristas_map, preview_oficina_motorista_label
        )
        return

    combo_oficina_motorista.configure(values=valores, state="normal")
    atual = combo_oficina_motorista.get()
    combo_oficina_motorista.set(atual if atual in valores else VALOR_SEM_MOTORISTA)
    atualizar_preview_combo_colaborador(
        combo_oficina_motorista, oficina_motoristas_map, preview_oficina_motorista_label
    )


def refresh_oficina_caminhoes():
    data_iso = validar_data(oficina_data_var.get().strip())
    ignorar = {"oficina_id": oficina_em_edicao_id} if oficina_em_edicao_id else None
    if not data_iso:
        atualizar_combo_caminhao_oficina(set())
        return
    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    atualizar_combo_caminhao_oficina(disponibilidade.get("caminhoes", set()))


oficina_data_var.trace_add(
    "write",
    lambda *_: (
        atualizar_data_saida_oficina(),
        atualizar_alerta_oficina(),
        refresh_oficina_motoristas(),
        refresh_oficina_caminhoes(),
    ),
)


def refresh_oficinas_ui():
    data_iso = validar_data(oficina_data_var.get().strip())
    for widget in lista_oficinas.winfo_children():
        widget.destroy()

    if not data_iso:
        mostrar_msg_lista(
            lista_oficinas, "Informe uma data válida para listar oficinas."
        )
        return

    registros = listar_oficinas(data_iso)
    if not registros:
        mostrar_msg_lista(lista_oficinas)
        return

    data_legenda = datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    for item in registros:
        placa_texto = item["placa"].upper() if item.get("placa") else DISPLAY_VAZIO
        motorista_nome = item.get("motorista_nome") or DISPLAY_VAZIO
        observacao_bruta = (item.get("observacao") or "").strip()
        observacao = observacao_bruta if observacao_bruta else DISPLAY_VAZIO
        em_edicao = oficina_em_edicao_id == item["id"]
        bloco = ctk.CTkFrame(
            lista_oficinas,
            fg_color=CARD_BG_EDICAO if em_edicao else CARD_BG_DEFAULT,
            corner_radius=10,
        )
        bloco.pack(fill="x", padx=10, pady=6)
        ctk.CTkLabel(
            bloco,
            text=f"{placa_texto} | Motorista: {motorista_nome}",
            text_color=COR_AZUL,
            font=("Inter", 15, "bold"),
        ).pack(anchor="w", padx=12, pady=(8, 2))
        cor_texto = (
            COR_VERMELHA
            if observacao_bruta and "manutenção" in observacao_bruta.lower()
            else COR_TEXTO
        )
        cor_marcador = ajustar_cor_marcador(item.get("observacao_cor"))
        obs_frame = ctk.CTkFrame(
            bloco,
            fg_color=cor_marcador or "transparent",
            corner_radius=8,
        )
        obs_frame.pack(fill="x", padx=10, pady=(0, 4))
        ctk.CTkLabel(
            obs_frame,
            text=f"{data_legenda} | {observacao}",
            text_color=cor_texto,
            anchor="w",
        ).pack(anchor="w", padx=10, pady=6)
        if em_edicao:
            ctk.CTkLabel(
                bloco,
                text="Em edição",
                text_color=COR_VERMELHA,
                font=FONT_SUBTITULO,
            ).pack(anchor="w", padx=12, pady=(0, 6))
        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="e", padx=8, pady=(0, 6))
        ctk.CTkButton(
            botoes,
            text="Editar",
            width=80,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda item=item: preparar_edicao_oficina(item),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            width=90,
            fg_color="#FFFFFF",
            text_color=COR_VERMELHA,
            border_color=COR_VERMELHA,
            border_width=1,
            hover_color="#FCE4E4",
            command=lambda item=item: handle_excluir_oficina(item),
        ).pack(side="left", padx=4)

    atualizar_alerta_oficina()


def handle_salvar_oficina():
    global oficina_em_edicao_id
    data_iso = validar_data(oficina_data_var.get().strip())
    if not data_iso:
        messagebox.showerror("Data invalida", "Informe uma data no formato DD/MM/AAAA.")
        return

    motorista_label = combo_oficina_motorista.get()
    placa = obter_placa_oficina()
    observacao = entry_oficina_obs.get().strip()
    observacao_extra = entry_oficina_obs_extra.get().strip()
    observacao_cor_label = combo_cor_oficina.get()
    observacao_cor = OBS_MARCADORES_MAP.get(observacao_cor_label, "")
    data_saida_texto = oficina_data_saida_var.get().strip()
    data_saida_iso = None
    if data_saida_texto:
        data_saida_iso = validar_data(data_saida_texto)
        if not data_saida_iso:
            messagebox.showerror("Data de saida invalida", "Use o formato DD/MM/AAAA na data de saida.")
            return

    motorista_id = oficina_motoristas_map.get(motorista_label)
    if motorista_id is None and motorista_label not in oficina_motoristas_map:
        messagebox.showerror("Motorista", "Selecione um motorista cadastrado ou deixe o campo em branco.")
        return
    if not placa:
        messagebox.showerror("Placa", "Informe a placa do veiculo.")
        return

    campos_vazios = []
    if motorista_id is None:
        campos_vazios.append("Motorista")
    if not observacao:
        campos_vazios.append("Observacoes")
    if not confirmar_campos_vazios(campos_vazios):
        return

    placa_upper = placa.upper()

    ignorar = {"oficina_id": oficina_em_edicao_id} if oficina_em_edicao_id else None
    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )
    if motorista_id and motorista_id in indis_colaboradores:
        messagebox.showwarning("Colaborador indisponivel", "Este motorista ja esta alocado em outra atividade nesta data.")
        return
    indis_caminhoes = {plac.upper() for plac in disponibilidade.get("caminhoes", set())}
    if placa_upper and placa_upper in indis_caminhoes:
        messagebox.showwarning("Caminhao indisponivel", "Este caminhao ja esta alocado em outra atividade nesta data.")
        return

    if not data_saida_iso:
        data_saida_iso = calcular_data_saida_padrao(data_iso)

    if data_saida_iso and data_saida_iso < data_iso:
        messagebox.showerror(
            "Periodo invalido",
            "A data de saida nao pode ser anterior a data.",
        )
        return

    try:
        if oficina_em_edicao_id:
            editar_oficina(
                oficina_em_edicao_id,
                motorista_id,
                placa_upper,
                observacao,
                observacao_extra,
                data_saida_iso,
                observacao_cor,
            )
            messagebox.showinfo("Atualizado", "Registro atualizado.")
        else:
            salvar_oficina(
                data_iso,
                motorista_id,
                placa_upper,
                observacao,
                observacao_extra,
                data_saida_iso,
                observacao_cor,
            )
            messagebox.showinfo("Sucesso", "Oficina registrada.")
    except ValueError as exc:
        messagebox.showerror("Validacao", str(exc))
        return
    except sqlite3.IntegrityError:
        messagebox.showwarning(
            "Placa duplicada",
            "Ja existe esta placa registrada na oficina para a data escolhida.",
        )
        return
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return

    criar_backup_automatico()
    limpar_form_oficina()
    refresh_oficinas_ui()
    recarregar_disponibilidade()

def limpar_form_oficina():
    global oficina_em_edicao_id
    oficina_em_edicao_id = None
    entry_oficina_obs.delete(0, "end")
    entry_oficina_obs_extra.delete(0, "end")
    entry_oficina_placa.delete(0, "end")
    oficina_data_saida_var.set("")
    atualizar_data_saida_oficina(force=True)
    if combo_oficina_motorista.cget("state") == "normal":
        combo_oficina_motorista.set(combo_oficina_motorista.cget("values")[0])
        atualizar_preview_combo_colaborador(
            combo_oficina_motorista, oficina_motoristas_map, preview_oficina_motorista_label
        )
    if combo_oficina_placa.winfo_manager() and combo_oficina_placa.cget("state") != "disabled":
        valores = combo_oficina_placa.cget("values")
        if valores:
            combo_oficina_placa.set(valores[0])
    combo_cor_oficina.set(OBS_MARCADORES[0][0])
    btn_salvar_oficina.configure(text="Salvar oficina")
    btn_cancelar_oficina.configure(state="disabled")
    refresh_oficina_motoristas()
    refresh_oficina_caminhoes()


def preparar_edicao_oficina(item: dict):
    global oficina_em_edicao_id
    oficina_em_edicao_id = item["id"]
    refresh_oficina_motoristas()
    refresh_oficina_caminhoes()
    selecionar_combo_colaborador(
        combo_oficina_motorista,
        oficina_motoristas_map,
        item.get("motorista_id"),
        item.get("motorista_nome"),
        VALOR_SEM_MOTORISTA,
        preview_oficina_motorista_label,
    )
    set_oficina_placa_field(item.get("placa") or "")
    entry_oficina_obs.delete(0, "end")
    entry_oficina_obs.insert(0, item.get("observacao") or "")
    entry_oficina_obs_extra.delete(0, "end")
    entry_oficina_obs_extra.insert(0, item.get("observacao_extra") or "")
    oficina_data_saida_var.set(data_iso_para_br_entrada(item.get("data_saida")))
    if not item.get("data_saida"):
        atualizar_data_saida_oficina(force=False)
    combo_cor_oficina.set(label_cor_observacao(item.get("observacao_cor")))
    btn_salvar_oficina.configure(text="Atualizar oficina")
    btn_cancelar_oficina.configure(state="normal")


def handle_excluir_oficina(item: dict):
    if not messagebox.askyesno(
        "Excluir oficina",
        "Deseja remover este registro de oficina?",
    ):
        return
    try:
        excluir_oficina(item["id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    limpar_form_oficina()
    refresh_oficinas_ui()
    recarregar_disponibilidade()


def exportar_relatorio_oficinas():
    data_iso = validar_data(oficina_data_var.get().strip())
    if not data_iso:
        messagebox.showerror(
            "Data inválida", "Informe uma data no formato DD/MM/AAAA."
        )
        return
    registros = listar_oficinas(data_iso)
    data_saida_txt = oficina_data_saida_var.get().strip()
    data_saida_iso = validar_data(data_saida_txt) if data_saida_txt else None
    if data_saida_txt and not data_saida_iso:
        messagebox.showerror(
            "Data inválida",
            "Use o formato DD/MM/AAAA para a data de saída.",
        )
        return
    if not data_saida_iso:
        data_saida_iso = calcular_data_saida_padrao(data_iso)
        if data_saida_iso:
            oficina_data_saida_var.set(data_iso_para_br_entrada(data_saida_iso))

    if data_saida_iso and data_saida_iso < data_iso:
        messagebox.showerror(
            "Periodo invalido",
            "A data de saida nao pode ser anterior a data.",
        )
        return

    linhas: list[list[str]] = []
    cores: list[str | None] = []
    for item in registros:
        linhas.append(
            [
                item.get("motorista_nome") or DISPLAY_VAZIO,
                (item["placa"].upper() if item.get("placa") else DISPLAY_VAZIO),
                combinar_observacoes(item.get("observacao"), item.get("observacao_extra")),
            ]
        )
        cores.append(item.get("observacao_cor"))

    gerar_relatorio_moderno(
        arquivo_stub="oficinas",
        titulo_header="JR Escala - Oficinas",
        linha_principal_rotulo="Oficina",
        data_principal_iso=data_iso,
        linha_secundaria_rotulo="Saída",
        data_secundaria_iso=data_saida_iso,
        total_legenda=f"Total de oficinas: {len(registros)}",
        colunas=["Motorista", "Placa", "Observações"],
        col_widths=[0.35, 0.2, 0.45],
        linhas=linhas,
        col_align_center={1},
        highlight_col=2,
        highlight_colors=cores,
        arquivo_data_iso=data_iso,
        fallback_highlight=False,
    )


btn_exportar_oficinas.configure(command=exportar_relatorio_oficinas)

def refresh_caminhoes_dropdowns(caminhoes: list[dict] | None = None):
    if caminhoes is None:
        caminhoes = listar_caminhoes_ativos()
    caminhao_option_map.clear()
    caminhao_option_map[VALOR_SEM_CAMINHAO] = ""
    for caminhao in caminhoes:
        display = caminhao["placa"].upper()
        if caminhao.get("modelo"):
            display += f" - {caminhao['modelo']}"
        caminhao_option_map[display] = caminhao["placa"].upper()
    atualizar_alerta_carregamento()
    atualizar_alerta_oficina()
    recarregar_disponibilidade()


btn_salvar_oficina = ctk.CTkButton(
    frame_oficina_form,
    text="Salvar oficina",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=handle_salvar_oficina,
)
btn_salvar_oficina.grid(row=2, column=5, padx=20, pady=5, sticky="e")
btn_cancelar_oficina = ctk.CTkButton(
    frame_oficina_form,
    text="Cancelar edição",
    fg_color=COR_VERMELHA,
    hover_color=COR_AZUL,
    state="disabled",
    command=limpar_form_oficina,
)
btn_cancelar_oficina.grid(row=3, column=5, padx=20, pady=5, sticky="w")
frame_oficina_form.grid_columnconfigure(0, weight=1)
frame_oficina_form.grid_columnconfigure(1, weight=1)
frame_oficina_form.grid_columnconfigure(2, weight=1)
frame_oficina_form.grid_columnconfigure(3, weight=1)
frame_oficina_form.grid_columnconfigure(4, weight=1)
frame_oficina_form.grid_columnconfigure(5, weight=1)

btn_listar_oficinas = ctk.CTkButton(
    aba_oficinas,
    text="Ver oficinas do dia",
    fg_color=COR_VERMELHA,
    hover_color=COR_AZUL,
    command=refresh_oficinas_ui,
)
btn_listar_oficinas.pack(anchor="e", padx=30, pady=(0, 8))

lista_oficinas = ctk.CTkScrollableFrame(
    aba_oficinas, fg_color="#FFFFFF", corner_radius=12, height=260
)
lista_oficinas.pack(fill="both", expand=True, padx=20, pady=(0, 20))

# ---------- ROTAS SEMANAIS ----------
ctk.CTkLabel(
    aba_rotas,
    text="ROTAS SEMANAIS JR",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_rotas,
    text="Configure rotas padrão para cada dia da semana.",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).pack(anchor="w", padx=10, pady=(0, 10))

frame_rotas_form = ctk.CTkFrame(aba_rotas, fg_color="#FFFFFF", corner_radius=12)
frame_rotas_form.pack(fill="x", padx=10, pady=(5, 12))

dia_atual_label = DIA_CHAVE_PARA_LABEL.get(
    obter_dia_semana_por_data(date.today().isoformat()), DIAS_SEMANA[0][1]
)
dia_semana_var = ctk.StringVar(value=dia_atual_label)

ctk.CTkLabel(
    frame_rotas_form,
    text="Dia da semana",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
combo_dia_semana = ctk.CTkOptionMenu(
    frame_rotas_form,
    values=DIAS_SEMANA_LABELS,
    command=lambda _: atualizar_lista_rotas_semanais(),
    variable=dia_semana_var,
)
combo_dia_semana.grid(row=1, column=0, padx=20, pady=5, sticky="w")
ctk.CTkLabel(
    frame_rotas_form,
    text="Data (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
entry_data_semana = ctk.CTkEntry(
    frame_rotas_form, width=150, placeholder_text="DD/MM/AAAA"
)
entry_data_semana.grid(row=1, column=1, padx=20, pady=5, sticky="w")
label_data_semana_erro = ctk.CTkLabel(
    frame_rotas_form,
    text="Data inválida (use DD/MM/AAAA)",
    text_color=COR_VERMELHA,
    font=("Inter", 12),
)
label_data_semana_erro.grid(row=2, column=1, padx=20, pady=(0, 5), sticky="w")
label_data_semana_erro.grid_remove()


def obter_dia_semana_rotas():
    return combo_dia_semana.get().strip()


def definir_dia_semana_pela_data(_event=None):
    valor = entry_data_semana.get().strip()
    if not valor:
        label_data_semana_erro.grid_remove()
        return
    data_iso = validar_data(valor)
    if not data_iso:
        label_data_semana_erro.configure(text="Data inválida (use DD/MM/AAAA)")
        label_data_semana_erro.grid()
        return
    data_selecionada = datetime.strptime(data_iso, "%Y-%m-%d").date()
    label_data_semana_erro.grid_remove()
    dia_chave = obter_dia_semana_por_data(data_selecionada.isoformat())
    dia_label = DIA_CHAVE_PARA_LABEL.get(dia_chave, DIAS_SEMANA[0][1])
    combo_dia_semana.set(dia_label)
    atualizar_lista_rotas_semanais()


entry_data_semana.bind("<Return>", definir_dia_semana_pela_data)
entry_data_semana.bind("<FocusOut>", definir_dia_semana_pela_data)

ctk.CTkLabel(
    frame_rotas_form,
    text="Rota",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=20, pady=(20, 0), sticky="w")
entry_rota_padrao = ctk.CTkEntry(
    frame_rotas_form, placeholder_text="Ex.: 10", width=240
)
entry_rota_padrao.grid(row=1, column=2, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_rotas_form,
    text="Destino",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=3, padx=20, pady=(20, 0), sticky="w")
entry_destino_padrao = ctk.CTkEntry(
    frame_rotas_form, placeholder_text="Ex.: Divinópolis", width=240
)
entry_destino_padrao.grid(row=1, column=3, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_rotas_form,
    text="Observações",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=4, padx=20, pady=(20, 0), sticky="w")
entry_obs_padrao = ctk.CTkEntry(
    frame_rotas_form, placeholder_text="Opcional", width=240
)
entry_obs_padrao.grid(row=1, column=4, padx=20, pady=5, sticky="w")

rota_semana_em_edicao_id: int | None = None


def limpar_form_rotas_semanais():
    global rota_semana_em_edicao_id
    rota_semana_em_edicao_id = None
    entry_rota_padrao.delete(0, "end")
    entry_destino_padrao.delete(0, "end")
    entry_obs_padrao.delete(0, "end")
    btn_salvar_rota_semana.configure(text="Salvar rota")
    btn_cancelar_rota_semana.configure(state="disabled")


def atualizar_lista_rotas_semanais():
    for widget in lista_rotas_semanais.winfo_children():
        widget.destroy()
    dia_label = obter_dia_semana_rotas()
    dia_chave = DIA_LABEL_PARA_CHAVE.get(dia_label, DIAS_SEMANA[0][0])
    rotas = listar_rotas_semanais(dia_chave)
    if not rotas:
        mostrar_msg_lista(lista_rotas_semanais, "Nenhuma rota cadastrada para este dia.")
        return
    for rota in rotas:
        bloco = ctk.CTkFrame(lista_rotas_semanais, fg_color="#F7F9FC", corner_radius=8)
        bloco.pack(fill="x", padx=10, pady=5)
        destino = rota.get("destino") or DISPLAY_VAZIO
        ctk.CTkLabel(
            bloco,
            text=f"{rota['rota']} - {destino}",
            text_color=COR_AZUL,
            font=("Inter", 15, "bold"),
        ).pack(anchor="w", padx=12, pady=(8, 2))
        obs = rota.get("observacao")
        if obs:
            ctk.CTkLabel(
                bloco,
                text=f"Obs.: {obs}",
                text_color=COR_TEXTO,
            ).pack(anchor="w", padx=12, pady=(0, 6))
        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="w", padx=10, pady=(0, 8))
        ctk.CTkButton(
            botoes,
            text="Editar",
            width=90,
            height=30,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda rota=rota: preparar_edicao_rota_semana(rota),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            width=90,
            height=30,
            fg_color="#FFFFFF",
            text_color=COR_VERMELHA,
            border_color=COR_VERMELHA,
            border_width=1,
            hover_color="#FCE4E4",
            command=lambda rota_id=rota["id"]: handle_excluir_rota_semana(rota_id),
        ).pack(side="left", padx=4)


def preparar_edicao_rota_semana(rota: dict):
    global rota_semana_em_edicao_id
    rota_semana_em_edicao_id = rota["id"]
    entry_rota_padrao.delete(0, "end")
    entry_rota_padrao.insert(0, rota["rota"])
    entry_destino_padrao.delete(0, "end")
    entry_destino_padrao.insert(0, rota.get("destino") or "")
    entry_obs_padrao.delete(0, "end")
    entry_obs_padrao.insert(0, rota.get("observacao") or "")
    btn_salvar_rota_semana.configure(text="Atualizar rota")
    btn_cancelar_rota_semana.configure(state="normal")


def handle_salvar_rota_semana():
    global rota_semana_em_edicao_id
    rota_texto = entry_rota_padrao.get().strip()
    if not rota_texto:
        messagebox.showerror("Campos obrigatórios", "Informe o nome da rota.")
        return
    destino = entry_destino_padrao.get().strip()
    if not destino:
        messagebox.showerror("Campos obrigatorios", "Informe o destino.")
        return
    observacao = entry_obs_padrao.get().strip()
    dia_label = obter_dia_semana_rotas()
    dia_chave = DIA_LABEL_PARA_CHAVE.get(dia_label, DIAS_SEMANA[0][0])
    em_edicao = rota_semana_em_edicao_id is not None
    try:
        if em_edicao:
            editar_rota_semana(rota_semana_em_edicao_id, dia_chave, rota_texto, destino, observacao)
            messagebox.showinfo("Atualizado", "Rota semanal atualizada.")
        else:
            adicionar_rota_semana(dia_chave, rota_texto, destino, observacao)
            messagebox.showinfo("Salvo", "Rota semanal adicionada.")
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return
    criar_backup_automatico()
    limpar_form_rotas_semanais()
    atualizar_lista_rotas_semanais()
    if not em_edicao:
        sincronizar_rota_semana_com_carregamentos(dia_chave, rota_texto, destino, observacao)


def sincronizar_rota_semana_com_carregamentos(
    dia_chave: str,
    rota_texto: str,
    destino: str,
    observacao: str,
):
    if "carreg_data_var" not in globals():
        return
    data_iso = validar_data(carreg_data_var.get().strip())
    if not data_iso:
        return
    if obter_dia_semana_por_data(data_iso) != dia_chave:
        return
    texto_rota = rota_texto.strip()
    destino = destino.strip()
    if destino:
        texto_rota = f"{texto_rota} - {destino}"
    if carregamento_existe_para_rota(data_iso, texto_rota):
        return
    data_saida_iso = None
    if "carreg_saida_var" in globals():
        data_saida_raw = carreg_saida_var.get().strip()
        if data_saida_raw:
            data_saida_iso = validar_data(data_saida_raw)
    try:
        salvar_carregamento(
            data_iso,
            texto_rota,
            None,
            None,
            None,
            OBSERVACAO_OPCOES[0],
            observacao_extra=observacao.strip() or None,
            data_saida=data_saida_iso,
        )
    except sqlite3.Error:
        return
    refresh_carregamentos_ui()
    recarregar_disponibilidade()


def handle_excluir_rota_semana(rota_id: int):
    if not messagebox.askyesno("Excluir rota", "Deseja realmente excluir esta rota semanal?"):
        return
    try:
        remover_rota_semana(rota_id)
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    limpar_form_rotas_semanais()
    atualizar_lista_rotas_semanais()


btn_salvar_rota_semana = ctk.CTkButton(
    frame_rotas_form,
    text="Salvar rota",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=handle_salvar_rota_semana,
)
btn_salvar_rota_semana.grid(row=1, column=5, padx=20, pady=5)
btn_cancelar_rota_semana = ctk.CTkButton(
    frame_rotas_form,
    text="Cancelar edição",
    fg_color=COR_VERMELHA,
    hover_color=COR_AZUL,
    state="disabled",
    command=limpar_form_rotas_semanais,
)
btn_cancelar_rota_semana.grid(row=1, column=6, padx=20, pady=5)

lista_rotas_semanais = ctk.CTkScrollableFrame(
    aba_rotas, fg_color="#FFFFFF", corner_radius=12, height=320
)
lista_rotas_semanais.pack(fill="both", expand=True, padx=20, pady=(5, 20))
atualizar_lista_rotas_semanais()
limpar_form_rotas_semanais()

# ---------- ESCALA (CD) ----------
ctk.CTkLabel(
    aba_cd,
    text="ESCALA (CD) — Centro de Distribuição",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_cd,
    text="Registre quem permanece no CD em cada dia.",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).pack(anchor="w", padx=10, pady=(0, 10))

btn_exportar_cd = ctk.CTkButton(
    aba_cd,
    text="Exportar relatório",
    command=lambda: None,
)
estilizar_botao(btn_exportar_cd, "primary")
btn_exportar_cd.pack(anchor="e", padx=20, pady=(0, 10))

frame_cd_form = ctk.CTkFrame(aba_cd, fg_color=COR_PAINEL, corner_radius=12)
frame_cd_form.pack(fill="x", padx=10, pady=(5, 12))

escala_cd_data_var = ctk.StringVar(value=date.today().strftime("%d/%m/%Y"))
escala_cd_data_saida_var = ctk.StringVar(value="")

ctk.CTkLabel(
    frame_cd_form,
    text="Data (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
entry_cd_data = ctk.CTkEntry(frame_cd_form, textvariable=escala_cd_data_var, width=160)
entry_cd_data.grid(row=1, column=0, padx=20, pady=5, sticky="w")
estilizar_entry(entry_cd_data)

ctk.CTkLabel(
    frame_cd_form,
    text="Data de saída (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
entry_cd_data_saida = ctk.CTkEntry(
    frame_cd_form, textvariable=escala_cd_data_saida_var, width=160
)
entry_cd_data_saida.grid(row=1, column=1, padx=20, pady=5, sticky="w")
estilizar_entry(entry_cd_data_saida)


def atualizar_data_saida_escala_cd(force: bool = False):
    data_iso = validar_data(escala_cd_data_var.get().strip())
    sugestao = calcular_data_saida_padrao(data_iso)
    atual = escala_cd_data_saida_var.get().strip()
    if force or not atual:
        escala_cd_data_saida_var.set(data_iso_para_br_entrada(sugestao))


escala_cd_data_var.trace_add("write", lambda *_: atualizar_data_saida_escala_cd())
atualizar_data_saida_escala_cd(force=True)


def aplicar_data_escala_cd():
    data_iso = validar_data(escala_cd_data_var.get().strip())
    if not data_iso:
        messagebox.showerror("Data inválida", "Informe uma data no formato DD/MM/AAAA.")
        return
    refresh_escala_cd_dropdowns()
    refresh_escala_cd_lista()


btn_cd_atualizar_data = ctk.CTkButton(
    frame_cd_form,
    text="Atualizar data",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=aplicar_data_escala_cd,
)
btn_cd_atualizar_data.grid(row=1, column=2, padx=20, pady=5)

ctk.CTkLabel(
    frame_cd_form,
    text="Motorista",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=2, column=0, padx=20, pady=(10, 0), sticky="w")
combo_cd_motorista = ctk.CTkOptionMenu(
    frame_cd_form, values=["Atualize a data"], state="disabled", width=220
)
combo_cd_motorista.grid(row=3, column=0, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_cd_form,
    text="Ajudante",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=2, column=1, padx=20, pady=(10, 0), sticky="w")
combo_cd_ajudante = ctk.CTkOptionMenu(
    frame_cd_form, values=["Atualize a data"], state="disabled", width=220
)
combo_cd_ajudante.grid(row=3, column=1, padx=20, pady=5, sticky="w")
preview_cd_motorista_label = ctk.CTkLabel(
    frame_cd_form,
    text="Sem foto",
    width=FOTO_PREVIEW_PADRAO[0],
    height=FOTO_PREVIEW_PADRAO[1],
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=12,
    anchor="center",
)
preview_cd_motorista_label.grid(row=4, column=0, padx=20, pady=(0, 10), sticky="w")
preview_cd_ajudante_label = ctk.CTkLabel(
    frame_cd_form,
    text="Sem foto",
    width=FOTO_PREVIEW_PADRAO[0],
    height=FOTO_PREVIEW_PADRAO[1],
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=12,
    anchor="center",
)
preview_cd_ajudante_label.grid(row=4, column=1, padx=20, pady=(0, 10), sticky="w")
combo_cd_motorista.configure(
    command=lambda valor: atualizar_preview_combo_colaborador(
        combo_cd_motorista,
        escala_cd_motorista_map,
        preview_cd_motorista_label,
        selecionado=valor,
    )
)
combo_cd_ajudante.configure(
    command=lambda valor: atualizar_preview_combo_colaborador(
        combo_cd_ajudante,
        escala_cd_ajudante_map,
        preview_cd_ajudante_label,
        selecionado=valor,
    )
)

ctk.CTkLabel(
    frame_cd_form,
    text="Observações",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=2, column=2, padx=20, pady=(10, 0), sticky="w")
entry_cd_obs = ctk.CTkEntry(
    frame_cd_form, placeholder_text="Observações (opcional)", width=240
)
entry_cd_obs.grid(row=3, column=2, padx=20, pady=5, sticky="w")

escala_cd_motorista_map: dict[str, int | None] = {}
escala_cd_ajudante_map: dict[str, int | None] = {}
escala_cd_em_edicao_id: int | None = None


def refresh_escala_cd_dropdowns():
    data_iso = validar_data(escala_cd_data_var.get().strip())
    ignorar = {"escala_cd_id": escala_cd_em_edicao_id} if escala_cd_em_edicao_id else None
    if not data_iso:
        combo_cd_motorista.configure(values=["Informe uma data"], state="disabled")
        combo_cd_motorista.set("Informe uma data")
        atualizar_preview_combo_colaborador(
            combo_cd_motorista, escala_cd_motorista_map, preview_cd_motorista_label
        )
        combo_cd_ajudante.configure(values=["Informe uma data"], state="disabled")
        combo_cd_ajudante.set("Informe uma data")
        atualizar_preview_combo_colaborador(
            combo_cd_ajudante, escala_cd_ajudante_map, preview_cd_ajudante_label
        )
        escala_cd_motorista_map.clear()
        escala_cd_ajudante_map.clear()
        return

    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )

    motoristas = [
        col for col in listar_colaboradores_por_funcao("Motorista") if col["id"] not in indis_colaboradores
    ]
    ajudantes = [
        col for col in listar_colaboradores_por_funcao("Ajudante") if col["id"] not in indis_colaboradores
    ]
    escala_cd_motorista_map.clear()
    escala_cd_ajudante_map.clear()
    escala_cd_motorista_map[VALOR_SEM_MOTORISTA] = None
    escala_cd_ajudante_map[VALOR_SEM_AJUDANTE] = None

    valores_motoristas = [VALOR_SEM_MOTORISTA]
    for col in motoristas:
        display = f"{col['nome']} (#{col['id']})"
        escala_cd_motorista_map[display] = col["id"]
        valores_motoristas.append(display)

    valores_ajudantes = [VALOR_SEM_AJUDANTE]
    for col in ajudantes:
        display = f"{col['nome']} (#{col['id']})"
        escala_cd_ajudante_map[display] = col["id"]
        valores_ajudantes.append(display)

    if len(valores_motoristas) == 1:
        valores_motoristas = [AVISO_NENHUM_COLAB]
        escala_cd_motorista_map[AVISO_NENHUM_COLAB] = None
    combo_cd_motorista.configure(values=valores_motoristas, state="normal")
    combo_cd_motorista.set(valores_motoristas[0])
    atualizar_preview_combo_colaborador(
        combo_cd_motorista, escala_cd_motorista_map, preview_cd_motorista_label
    )

    if len(valores_ajudantes) == 1:
        valores_ajudantes = [AVISO_NENHUM_COLAB]
        escala_cd_ajudante_map[AVISO_NENHUM_COLAB] = None
    combo_cd_ajudante.configure(values=valores_ajudantes, state="normal")
    combo_cd_ajudante.set(valores_ajudantes[0])
    atualizar_preview_combo_colaborador(
        combo_cd_ajudante, escala_cd_ajudante_map, preview_cd_ajudante_label
    )


def refresh_escala_cd_lista():
    data_iso = validar_data(escala_cd_data_var.get().strip())
    for widget in lista_escala_cd.winfo_children():
        widget.destroy()

    if not data_iso:
        mostrar_msg_lista(lista_escala_cd, "Informe uma data válida.")
        return

    registros = listar_escala_cd(data_iso)
    if not registros:
        mostrar_msg_lista(lista_escala_cd, "Nenhum registro encontrado para esta data.")
        return

    for item in registros:
        bloco = ctk.CTkFrame(lista_escala_cd, fg_color="#F7F9FC", corner_radius=10)
        bloco.pack(fill="x", padx=10, pady=5)
        texto_motorista = item.get("motorista_nome") or DISPLAY_VAZIO
        texto_ajudante = item.get("ajudante_nome") or DISPLAY_VAZIO
        ctk.CTkLabel(
            bloco,
            text=f"Motorista: {texto_motorista}",
            text_color=COR_TEXTO,
            font=("Inter", 15, "bold"),
        ).pack(anchor="w", padx=12, pady=(8, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Ajudante: {texto_ajudante}",
            text_color=COR_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 2))
        ctk.CTkLabel(
            bloco,
            text=f"Obs.: {item.get('observacao') or DISPLAY_VAZIO}",
            text_color=COR_TEXTO,
            font=FONT_TEXTO,
        ).pack(anchor="w", padx=12, pady=(0, 6))
        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="e", padx=8, pady=(0, 8))
        ctk.CTkButton(
            botoes,
            text="Editar",
            width=90,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda item=item: preparar_edicao_escala_cd(item),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            width=90,
            fg_color="#FFFFFF",
            text_color=COR_VERMELHA,
            border_color=COR_VERMELHA,
            border_width=1,
            hover_color="#FCE4E4",
            command=lambda item=item: handle_excluir_escala_cd(item),
        ).pack(side="left", padx=4)


def preparar_edicao_escala_cd(item: dict):
    global escala_cd_em_edicao_id
    escala_cd_em_edicao_id = item["id"]
    selecionar_combo_colaborador(
        combo_cd_motorista,
        escala_cd_motorista_map,
        item.get("motorista_id"),
        item.get("motorista_nome"),
        VALOR_SEM_MOTORISTA,
        preview_cd_motorista_label,
    )
    selecionar_combo_colaborador(
        combo_cd_ajudante,
        escala_cd_ajudante_map,
        item.get("ajudante_id"),
        item.get("ajudante_nome"),
        VALOR_SEM_AJUDANTE,
        preview_cd_ajudante_label,
    )
    obs = item.get("observacao") or ""
    entry_cd_obs.delete(0, "end")
    entry_cd_obs.insert(0, obs)
    btn_cd_salvar.configure(text="Atualizar escala")
    btn_cd_cancelar.configure(state="normal")


def limpar_form_escala_cd():
    global escala_cd_em_edicao_id
    escala_cd_em_edicao_id = None
    refresh_escala_cd_dropdowns()
    entry_cd_obs.delete(0, "end")
    escala_cd_data_saida_var.set("")
    atualizar_data_saida_escala_cd(force=True)
    btn_cd_salvar.configure(text="Salvar escala")
    btn_cd_cancelar.configure(state="disabled")


def handle_excluir_escala_cd(item: dict):
    if not messagebox.askyesno(
        "Excluir registro",
        "Deseja remover esta escala do CD?",
    ):
        return
    try:
        excluir_escala_cd(item["id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    limpar_form_escala_cd()
    refresh_escala_cd_lista()
    recarregar_disponibilidade()


def handle_salvar_escala_cd():
    global escala_cd_em_edicao_id
    data_iso = validar_data(escala_cd_data_var.get().strip())
    if not data_iso:
        messagebox.showerror("Data inválida", "Informe uma data válida.")
        return
    motorista_label = combo_cd_motorista.get()
    ajudante_label = combo_cd_ajudante.get()
    motorista_id = escala_cd_motorista_map.get(motorista_label)
    ajudante_id = escala_cd_ajudante_map.get(ajudante_label)
    observacao = entry_cd_obs.get().strip()
    ignorar = {"escala_cd_id": escala_cd_em_edicao_id} if escala_cd_em_edicao_id else None
    disponibilidade = verificar_disponibilidade(data_iso, ignorar)
    indis_colaboradores = disponibilidade.get("motoristas", set()).union(
        disponibilidade.get("ajudantes", set())
    )

    if motorista_id and motorista_id in indis_colaboradores:
        messagebox.showwarning(
            "Colaborador indisponível",
            "Este motorista já está escalado para outra atividade neste dia.",
        )
        return

    if ajudante_id and ajudante_id in indis_colaboradores:
        messagebox.showwarning(
            "Colaborador indisponível",
            "Este ajudante já está escalado para outra atividade neste dia.",
        )
        return

    try:
        if escala_cd_em_edicao_id:
            editar_escala_cd(escala_cd_em_edicao_id, motorista_id, ajudante_id, observacao)
            messagebox.showinfo("Atualizado", "Escala atualizada.")
        else:
            adicionar_escala_cd(data_iso, motorista_id, ajudante_id, observacao)
            messagebox.showinfo("Salvo", "Escala registrada.")
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return

    criar_backup_automatico()
    limpar_form_escala_cd()
    refresh_escala_cd_lista()
    recarregar_disponibilidade()


def exportar_relatorio_escala_cd():
    data_iso = validar_data(escala_cd_data_var.get().strip())
    if not data_iso:
        messagebox.showerror("Data inválida", "Informe uma data no formato DD/MM/AAAA.")
        return

    data_saida_txt = escala_cd_data_saida_var.get().strip()
    data_saida_iso = validar_data(data_saida_txt) if data_saida_txt else None
    if data_saida_txt and not data_saida_iso:
        messagebox.showerror(
            "Data inválida",
            "Use o formato DD/MM/AAAA para a data de saída.",
        )
        return
    if not data_saida_iso:
        data_saida_iso = calcular_data_saida_padrao(data_iso)
        if data_saida_iso:
            escala_cd_data_saida_var.set(data_iso_para_br_entrada(data_saida_iso))

    if data_saida_iso and data_saida_iso < data_iso:
        messagebox.showerror(
            "Periodo invalido",
            "A data de saida nao pode ser anterior a data.",
        )
        return

    registros = listar_escala_cd(data_iso)
    linhas = [
        [
            item.get("motorista_nome") or DISPLAY_VAZIO,
            item.get("ajudante_nome") or DISPLAY_VAZIO,
            item.get("observacao") or DISPLAY_VAZIO,
        ]
        for item in registros
    ]

    gerar_relatorio_moderno(
        arquivo_stub="escala_cd",
        titulo_header="JR Escala - Escala (CD)",
        linha_principal_rotulo="Escala (CD)",
        data_principal_iso=data_iso,
        linha_secundaria_rotulo="Saída",
        data_secundaria_iso=data_saida_iso or data_iso,
        total_legenda=f"Total de equipes: {len(registros)}",
        colunas=["Motorista", "Ajudante", "Observações"],
        col_widths=[0.4, 0.3, 0.3],
        linhas=linhas,
        highlight_col=None,
        arquivo_data_iso=data_iso,
    )

btn_exportar_cd.configure(command=exportar_relatorio_escala_cd)


btn_cd_salvar = ctk.CTkButton(
    frame_cd_form,
    text="Salvar escala",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=handle_salvar_escala_cd,
)
btn_cd_salvar.grid(row=3, column=3, padx=20, pady=5)
btn_cd_cancelar = ctk.CTkButton(
    frame_cd_form,
    text="Cancelar edição",
    fg_color=COR_VERMELHA,
    hover_color=COR_AZUL,
    state="disabled",
    command=limpar_form_escala_cd,
)
btn_cd_cancelar.grid(row=3, column=4, padx=20, pady=5)

lista_escala_cd = ctk.CTkScrollableFrame(
    aba_cd, fg_color="#FFFFFF", corner_radius=12, height=320
)
lista_escala_cd.pack(fill="both", expand=True, padx=20, pady=(5, 20))
refresh_escala_cd_dropdowns()
refresh_escala_cd_lista()
# ---------- CAMINHÕES ----------
ctk.CTkLabel(
    aba_caminhoes,
    text="CAMINHÕES",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_caminhoes,
    text="Cadastre e gerencie a frota disponível para as rotas e oficinas.",
    text_color=COR_TEXTO,
).pack(anchor="w", padx=10, pady=(0, 10))

frame_caminhao_form = ctk.CTkFrame(aba_caminhoes, fg_color="#FFFFFF", corner_radius=12)
frame_caminhao_form.pack(fill="x", padx=10, pady=(5, 12))

ctk.CTkLabel(
    frame_caminhao_form,
    text="Placa",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
entry_caminhao_placa = ctk.CTkEntry(
    frame_caminhao_form, placeholder_text="ABC1D23", width=160
)
entry_caminhao_placa.grid(row=1, column=0, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_caminhao_form,
    text="Modelo",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
entry_caminhao_modelo = ctk.CTkEntry(
    frame_caminhao_form, placeholder_text="Ex.: VW Delivery 13.180", width=280
)
entry_caminhao_modelo.grid(row=1, column=1, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_caminhao_form,
    text="Observação",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=20, pady=(20, 0), sticky="w")
entry_caminhao_obs = ctk.CTkEntry(
    frame_caminhao_form, placeholder_text="Opcional (ex.: plataforma)", width=280
)
entry_caminhao_obs.grid(row=1, column=2, padx=20, pady=5, sticky="w")

caminhao_em_edicao_id: int | None = None


def limpar_caminhao_form():
    global caminhao_em_edicao_id
    caminhao_em_edicao_id = None
    entry_caminhao_placa.delete(0, "end")
    entry_caminhao_modelo.delete(0, "end")
    entry_caminhao_obs.delete(0, "end")
    btn_salvar_caminhao.configure(text="Salvar caminhão")
    btn_cancelar_edicao_caminhao.configure(state="disabled")


def handle_salvar_caminhao():
    global caminhao_em_edicao_id
    placa = entry_caminhao_placa.get().strip().upper()
    modelo = entry_caminhao_modelo.get().strip()
    observacao = entry_caminhao_obs.get().strip()

    if not placa:
        messagebox.showerror("Campos obrigatórios", "Informe a placa do caminhão.")
        return

    try:
        if caminhao_em_edicao_id:
            editar_caminhao(caminhao_em_edicao_id, placa, modelo, observacao)
            messagebox.showinfo("Atualizado", "Caminhão atualizado com sucesso.")
        else:
            add_caminhao(placa, modelo, observacao)
            messagebox.showinfo("Sucesso", "Caminhão cadastrado com sucesso.")
    except ValueError as exc:
        messagebox.showerror("Validação", str(exc))
        return
    except sqlite3.IntegrityError:
        messagebox.showwarning(
            "Placa duplicada", "Já existe um caminhão cadastrado com essa placa."
        )
        return
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return

    criar_backup_automatico()
    limpar_caminhao_form()
    refresh_caminhoes_ui()


btn_salvar_caminhao = ctk.CTkButton(
    frame_caminhao_form,
    text="Salvar caminhão",
    fg_color=COR_AZUL,
    hover_color=COR_VERMELHA,
    command=handle_salvar_caminhao,
)
btn_salvar_caminhao.grid(row=1, column=3, padx=20, pady=5)

btn_cancelar_edicao_caminhao = ctk.CTkButton(
    frame_caminhao_form,
    text="Cancelar edição",
    fg_color=COR_VERMELHA,
    hover_color=COR_AZUL,
    command=limpar_caminhao_form,
    state="disabled",
)
btn_cancelar_edicao_caminhao.grid(row=1, column=4, padx=20, pady=5)

ctk.CTkLabel(
    aba_caminhoes,
    text="Caminhões cadastrados",
    text_color=COR_TEXTO,
    font=("Inter", 18, "bold"),
).pack(anchor="w", padx=20, pady=(5, 0))

frame_busca_caminhoes = ctk.CTkFrame(aba_caminhoes, fg_color="transparent")
frame_busca_caminhoes.pack(fill="x", padx=20, pady=(4, 0))
ctk.CTkLabel(
    frame_busca_caminhoes,
    text="Pesquisar caminhao",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, sticky="w")
caminhao_busca_var = ctk.StringVar(value="")
entry_caminhao_busca = ctk.CTkEntry(
    frame_busca_caminhoes,
    textvariable=caminhao_busca_var,
    placeholder_text="Digite a placa ou modelo",
    width=320,
)
entry_caminhao_busca.grid(row=1, column=0, pady=(4, 0), sticky="w")
estilizar_entry(entry_caminhao_busca)
caminhao_busca_var.trace_add(
    "write",
    lambda *_: debounce(
        "caminhao_busca",
        180,
        lambda: refresh_caminhoes_ui(atualizar_dependentes=False),
    ),
)
frame_busca_caminhoes.grid_columnconfigure(0, weight=1)

lista_caminhoes = ctk.CTkScrollableFrame(
    aba_caminhoes, fg_color="#FFFFFF", corner_radius=12, height=320
)
lista_caminhoes.pack(fill="both", expand=True, padx=20, pady=(5, 20))


def preparar_edicao_caminhao(caminhao: dict):
    global caminhao_em_edicao_id
    caminhao_em_edicao_id = caminhao["id"]
    entry_caminhao_placa.delete(0, "end")
    entry_caminhao_placa.insert(0, caminhao["placa"])
    entry_caminhao_modelo.delete(0, "end")
    entry_caminhao_modelo.insert(0, caminhao.get("modelo", ""))
    entry_caminhao_obs.delete(0, "end")
    entry_caminhao_obs.insert(0, caminhao.get("observacao", ""))
    btn_salvar_caminhao.configure(text="Atualizar caminhão")
    btn_cancelar_edicao_caminhao.configure(state="normal")


def handle_excluir_caminhao(caminhao_id: int, placa: str):
    if not messagebox.askyesno(
        "Excluir caminhão", f"Confirma a exclusão do caminhão {placa}?"
    ):
        return
    try:
        remover_caminhao(caminhao_id)
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return

    messagebox.showinfo("Removido", "Caminhão excluído.")
    limpar_caminhao_form()
    refresh_caminhoes_ui()


def refresh_caminhoes_ui(atualizar_dependentes: bool = True):
    for widget in lista_caminhoes.winfo_children():
        widget.destroy()

    caminhoes = listar_caminhoes_ativos()
    filtro = caminhao_busca_var.get().strip().lower()
    if filtro:
        caminhoes_filtrados: list[dict] = []
        for caminhao in caminhoes:
            alvo = " ".join(
                [
                    caminhao.get("placa", ""),
                    caminhao.get("modelo", ""),
                    caminhao.get("observacao", ""),
                ]
            ).lower()
            if filtro in alvo:
                caminhoes_filtrados.append(caminhao)
    else:
        caminhoes_filtrados = caminhoes

    if not caminhoes:
        mostrar_msg_lista(lista_caminhoes)
    elif not caminhoes_filtrados:
        mostrar_msg_lista(lista_caminhoes, "Nenhum caminhao encontrado.")
    else:
        for caminhao in caminhoes_filtrados:
            bloco = ctk.CTkFrame(lista_caminhoes, fg_color="#F7F9FC", corner_radius=10)
            bloco.pack(fill="x", padx=10, pady=6)
            titulo = f"{caminhao['placa'].upper()}"
            if caminhao.get("modelo"):
                titulo += f" - {caminhao['modelo']}"
            ctk.CTkLabel(
                bloco,
                text=titulo,
                text_color=COR_AZUL,
                font=("Inter", 15, "bold"),
            ).pack(anchor="w", padx=12, pady=(8, 2))
            if caminhao.get("observacao"):
                ctk.CTkLabel(
                    bloco,
                    text=caminhao["observacao"],
                    text_color=COR_TEXTO,
                ).pack(anchor="w", padx=12, pady=(0, 4))
            botoes = ctk.CTkFrame(bloco, fg_color="transparent")
            botoes.pack(anchor="w", padx=10, pady=(0, 8))
            ctk.CTkButton(
                botoes,
                text="Editar",
                width=90,
                height=32,
                fg_color=COR_AZUL,
                hover_color=COR_VERMELHA,
                command=lambda data=caminhao: preparar_edicao_caminhao(data),
            ).pack(side="left", padx=5)
            ctk.CTkButton(
                botoes,
                text="Excluir",
                width=105,
                height=32,
                fg_color="#FFFFFF",
                text_color=COR_VERMELHA,
                hover_color="#FCE4E4",
                border_color=COR_VERMELHA,
                border_width=1,
                command=lambda cid=caminhao["id"], placa=caminhao["placa"]: handle_excluir_caminhao(
                    cid, placa
                ),
            ).pack(side="left", padx=5)

    if atualizar_dependentes:
        refresh_caminhoes_dropdowns(caminhoes)

# ---------- FOLGA ----------
ctk.CTkLabel(
    aba_folga,
    text="FOLGAS",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_folga,
    text="Cadastre e acompanhe folgas e afastamentos dos colaboradores.",
    text_color=COR_TEXTO,
).pack(anchor="w", padx=10, pady=(0, 10))

btn_exportar_folgas = ctk.CTkButton(
    aba_folga,
    text="Exportar relatório",
    command=lambda: None,
)
estilizar_botao(btn_exportar_folgas, "primary")
btn_exportar_folgas.pack(anchor="e", padx=20, pady=(0, 10))

frame_folga_form = ctk.CTkFrame(aba_folga, fg_color="#FFFFFF", corner_radius=12)
frame_folga_form.pack(fill="x", padx=10, pady=(5, 12))

folga_data_var = ctk.StringVar(value=date.today().strftime("%d/%m/%Y"))
folga_data_saida_var = ctk.StringVar(value="")
folga_obs_padrao_var = ctk.StringVar(value=OBSERVACAO_OPCOES[0])

ctk.CTkLabel(
    frame_folga_form,
    text="Data da folga (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
entry_folga_data = ctk.CTkEntry(frame_folga_form, textvariable=folga_data_var, width=180)
entry_folga_data.grid(row=1, column=0, padx=20, pady=5, sticky="w")
estilizar_entry(entry_folga_data)

label_folga_retorno = ctk.CTkLabel(
    frame_folga_form,
    text="Data de saida (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
)
label_folga_retorno.grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
entry_folga_data_saida = ctk.CTkEntry(
    frame_folga_form, textvariable=folga_data_saida_var, width=180
)
entry_folga_data_saida.grid(row=1, column=1, padx=20, pady=5, sticky="w")
estilizar_entry(entry_folga_data_saida)


def atualizar_data_saida_folga(_event=None):
    data_iso = validar_data(folga_data_var.get().strip())
    if not data_iso:
        return
    data_saida_atual = folga_data_saida_var.get().strip()
    if data_saida_atual:
        if validar_data(data_saida_atual):
            return
    data_saida_iso = calcular_data_saida_padrao(data_iso)
    if not data_saida_iso:
        return
    folga_data_saida_var.set(data_iso_para_br_entrada(data_saida_iso))


entry_folga_data.bind("<Return>", atualizar_data_saida_folga)
entry_folga_data.bind("<FocusOut>", atualizar_data_saida_folga)
atualizar_data_saida_folga()

ctk.CTkLabel(
    frame_folga_form,
    text="Colaborador",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=20, pady=(20, 0), sticky="w")
combo_folga_colaborador = ctk.CTkOptionMenu(
    frame_folga_form, values=["Cadastre um colaborador"], state="disabled", width=260
)
combo_folga_colaborador.grid(row=1, column=2, padx=20, pady=5, sticky="w")
preview_folga_colaborador_label = ctk.CTkLabel(
    frame_folga_form,
    text="Sem foto",
    width=110,
    height=110,
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=55,
    anchor="center",
)
preview_folga_colaborador_label.grid(row=0, column=4, rowspan=3, padx=20, pady=(15, 10), sticky="n")


def atualizar_preview_folga_combo(selecionado: str | None = None):
    atualizar_preview_combo_colaborador(
        combo_folga_colaborador,
        colaborador_option_map,
        preview_folga_colaborador_label,
        size=(110, 110),
        selecionado=selecionado,
        circular=True,
    )


combo_folga_colaborador.configure(
    command=lambda valor: atualizar_preview_folga_combo(valor)
)

btn_atualizar_folgas = ctk.CTkButton(
    frame_folga_form,
    text="Atualizar lista",
    command=lambda: atualizar_lista_folgas(),
)
estilizar_botao(btn_atualizar_folgas, "ghost", small=True)
btn_atualizar_folgas.grid(row=1, column=3, padx=20, pady=5, sticky="w")

btn_salvar_folga = ctk.CTkButton(
    frame_folga_form,
    text="Salvar folga",
    command=lambda: handle_salvar_folga(),
)
estilizar_botao(btn_salvar_folga, "primary")
btn_salvar_folga.grid(row=3, column=3, padx=20, pady=5, sticky="w")
btn_cancelar_folga = ctk.CTkButton(
    frame_folga_form,
    text="Cancelar edição",
    state="disabled",
    command=lambda: limpar_form_folga(),
)
estilizar_botao(btn_cancelar_folga, "danger-ghost")
btn_cancelar_folga.grid(row=3, column=4, padx=20, pady=5, sticky="w")

lista_folgas = ctk.CTkScrollableFrame(
    aba_folga, fg_color="#FFFFFF", corner_radius=12, height=320
)
lista_folgas.pack(fill="both", expand=True, padx=10, pady=(5, 20))


def atualizar_dropdown_folga():
    valores = list(colaborador_option_map.keys())
    if not valores:
        combo_folga_colaborador.configure(values=["Cadastre um colaborador"], state="disabled")
        combo_folga_colaborador.set("Cadastre um colaborador")
        atualizar_preview_folga_combo(combo_folga_colaborador.get())
    else:
        combo_folga_colaborador.configure(values=valores, state="normal")
        combo_folga_colaborador.set(valores[0])
        atualizar_preview_folga_combo(valores[0])


def _linha_relatorio_folga(registro: dict) -> tuple[list[str], str | None]:
    nome = registro.get("nome") or DISPLAY_VAZIO
    funcao = registro.get("funcao") or DISPLAY_VAZIO
    inicio_iso = registro.get("data")
    fim_iso = registro.get("data_fim") or calcular_data_saida_padrao(inicio_iso)
    inicio = data_iso_para_br(inicio_iso)
    fim = data_iso_para_br(fim_iso)
    periodo = inicio or DISPLAY_VAZIO
    if fim and fim != DISPLAY_VAZIO and fim != inicio:
        periodo = f"{inicio} -> {fim}"
    obs = (registro.get("observacao_extra") or "").strip() or DISPLAY_VAZIO
    cor = ajustar_cor_marcador(registro.get("observacao_cor"))
    return [nome, funcao, periodo, obs], cor


def _desenhar_relatorio_folgas(
    data_iso: str, data_saida_iso: str | None, linhas: list[list[str]], total_registros: int
):
    header_iso = data_saida_iso or data_iso
    gerar_relatorio_moderno(
        arquivo_stub="folgas",
        titulo_header="FOLGA:",
        linha_principal_rotulo="Data",
        data_principal_iso=header_iso,
        linha_secundaria_rotulo=None,
        data_secundaria_iso=None,
        total_legenda=f"Total de folgas: {total_registros}",
        colunas=["Motorista", "Ajudante"],
        col_widths=[0.5, 0.5],
        linhas=linhas,
        col_align_center=set(),
        highlight_col=None,
        highlight_colors=[],
        arquivo_data_iso=data_iso,
        fallback_highlight=False,
    )


def exportar_relatorio_folgas():
    data_iso = validar_data(folga_data_var.get().strip())
    if not data_iso:
        messagebox.showerror("Data invalida", "Informe a data da folga no formato DD/MM/AAAA.")
        return
    data_saida_raw = folga_data_saida_var.get().strip()
    if data_saida_raw:
        data_saida_iso = validar_data(data_saida_raw)
        if not data_saida_iso:
            messagebox.showerror(
                "Data invalida",
                "Informe a data de saida no formato DD/MM/AAAA.",
            )
            return
    else:
        data_saida_iso = calcular_data_saida_padrao(data_iso)
    registros = listar_folgas(data_iso)
    if not registros:
        messagebox.showinfo("Relatorio", "Nao ha folgas cadastradas para exportar.")
        return
    motoristas: list[str] = []
    ajudantes: list[str] = []
    for registro in registros:
        nome = (registro.get("nome") or "").strip()
        if not nome:
            continue
        funcao = (registro.get("funcao") or "").strip().lower()
        if funcao.startswith("motor"):
            motoristas.append(nome)
        elif funcao.startswith("ajud"):
            ajudantes.append(nome)
        else:
            motoristas.append(nome)
    motoristas.sort()
    ajudantes.sort()
    linhas: list[list[str]] = []
    max_len = max(len(motoristas), len(ajudantes))
    for idx in range(max_len):
        linhas.append([
            motoristas[idx] if idx < len(motoristas) else "",
            ajudantes[idx] if idx < len(ajudantes) else "",
        ])
    _desenhar_relatorio_folgas(data_iso, data_saida_iso, linhas, len(registros))


btn_exportar_folgas.configure(command=exportar_relatorio_folgas)


def obter_colaborador_id_dropdown(selecionado: str) -> int | None:
    return colaborador_option_map.get(selecionado)


def atualizar_lista_folgas():
    data_txt = folga_data_var.get().strip()
    data_iso = validar_data(data_txt)
    if not data_iso:
        messagebox.showerror("Data inválida", "Informe a data no formato DD/MM/AAAA.")
        return
    exibir_folgas_para_data(data_iso)


def handle_salvar_folga():
    global folga_em_edicao_id
    data_txt = folga_data_var.get().strip()
    data_iso = validar_data(data_txt)
    if not data_iso:
        messagebox.showerror("Data inv?lida", "Informe a data no formato DD/MM/AAAA.")
        return
    data_saida_txt = folga_data_saida_var.get().strip()
    if data_saida_txt:
        retorno_iso = validar_data(data_saida_txt)
        if not retorno_iso:
            messagebox.showerror(
                "Data invalida",
                "Informe a data de saida no formato DD/MM/AAAA.",
            )
            return
    else:
        retorno_iso = calcular_data_saida_padrao(data_iso)
        if retorno_iso:
            folga_data_saida_var.set(data_iso_para_br_entrada(retorno_iso))
    if retorno_iso and retorno_iso < data_iso:
        messagebox.showerror(
            "Periodo invalido",
            "A data de saida nao pode ser anterior a data.",
        )
        return
    colaborador_display = combo_folga_colaborador.get()
    colaborador_id = obter_colaborador_id_dropdown(colaborador_display)
    if not colaborador_id:
        messagebox.showerror("Campos obrigat?rios", "Selecione um colaborador.")
        return
    observacao_padrao = folga_obs_padrao_var.get()
    observacao_extra = None
    observacao_cor = ""
    try:
        if folga_em_edicao_id:
            editar_folga(
                folga_em_edicao_id,
                data_iso,
                retorno_iso,
                colaborador_id,
                observacao_padrao,
                observacao_extra,
                observacao_cor,
            )
            messagebox.showinfo("Atualizado", "Folga atualizada com sucesso.")
        else:
            salvar_folga(
                data_iso,
                colaborador_id,
                data_fim=retorno_iso,
                observacao_padrao=observacao_padrao,
                observacao_extra=observacao_extra,
                observacao_cor=observacao_cor,
            )
            messagebox.showinfo("Salvo", "Folga registrada com sucesso.")
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return
    criar_backup_automatico()
    limpar_form_folga()
    exibir_folgas_para_data(data_iso)
    recarregar_disponibilidade()

def limpar_form_folga():
    global folga_em_edicao_id
    folga_em_edicao_id = None
    folga_data_var.set(date.today().strftime("%d/%m/%Y"))
    data_saida_iso = calcular_data_saida_padrao(date.today().isoformat())
    folga_data_saida_var.set(data_iso_para_br_entrada(data_saida_iso))
    folga_obs_padrao_var.set(OBSERVACAO_OPCOES[0])
    btn_salvar_folga.configure(text="Salvar folga")
    btn_cancelar_folga.configure(state="disabled")
    if combo_folga_colaborador.cget("state") == "normal":
        combo_folga_colaborador.set(combo_folga_colaborador.cget("values")[0])
    atualizar_preview_folga_combo(combo_folga_colaborador.get())


def preparar_edicao_folga(registro: dict):
    global folga_em_edicao_id
    folga_em_edicao_id = registro.get("folga_id")
    folga_data_var.set(
        data_iso_para_br_entrada(registro.get("data") or date.today().isoformat())
    )
    folga_data_saida_var.set(
        data_iso_para_br_entrada(
            registro.get("data_fim") or calcular_data_saida_padrao(registro.get("data"))
        )
    )
    folga_obs_padrao_var.set(registro.get("observacao_padrao") or OBSERVACAO_OPCOES[0])
    colaborador_display = None
    colaborador_id = registro.get("colaborador_id")
    if colaborador_id:
        for display, cid in colaborador_option_map.items():
            if cid == colaborador_id:
                colaborador_display = display
                break
    if colaborador_display:
        combo_folga_colaborador.set(colaborador_display)
    else:
        valores = combo_folga_colaborador.cget("values")
        if valores:
            combo_folga_colaborador.set(valores[0])
    atualizar_preview_folga_combo(combo_folga_colaborador.get())
    btn_salvar_folga.configure(text="Atualizar folga")
    btn_cancelar_folga.configure(state="normal")


def handle_excluir_folga(item: dict):
    if not messagebox.askyesno(
        "Excluir folga", f"Deseja remover a folga de {item['nome']} em {item['data']}?"
    ):
        return
    try:
        remover_folga(item["folga_id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    exibir_folgas_para_data(item["data"])
    recarregar_disponibilidade()


def exibir_folgas_para_data(data_iso: str):
    for widget in lista_folgas.winfo_children():
        widget.destroy()
    registros = listar_folgas(data_iso)
    if not registros:
        mostrar_msg_lista(lista_folgas, "Nenhuma folga nesta data.")
        return
    for registro in registros:
        bloco = ctk.CTkFrame(lista_folgas, fg_color="#F9FAFB", corner_radius=14)
        bloco.pack(fill="x", padx=10, pady=4)
        conteudo = ctk.CTkFrame(bloco, fg_color="transparent")
        conteudo.pack(fill="x", padx=12, pady=8)
        avatar = ctk.CTkLabel(
            conteudo,
            text="Sem foto",
            width=70,
            height=70,
            fg_color="#E7ECF5",
            text_color=COR_CINZA,
            corner_radius=999,
            anchor="center",
        )
        avatar.pack(side="left", padx=(0, 12))
        imagem_avatar = obter_imagem_colaborador_por_id(
            registro.get("colaborador_id"), size=(70, 70), circular=True
        )
        if imagem_avatar:
            avatar.configure(image=imagem_avatar, text="")
            avatar.image = imagem_avatar
        info = ctk.CTkFrame(conteudo, fg_color="transparent")
        info.pack(fill="both", expand=True)
        titulo = f"{registro['nome']} - {registro['funcao']}"
        ctk.CTkLabel(
            info,
            text=titulo,
            text_color=COR_AZUL,
            font=("Inter", 14, "bold"),
        ).pack(anchor="w", pady=(0, 2))
        inicio_iso = registro.get("data")
        fim_iso = registro.get("data_fim") or calcular_data_saida_padrao(inicio_iso)
        inicio = data_iso_para_br(inicio_iso)
        fim = data_iso_para_br(fim_iso)
        periodo = inicio or DISPLAY_VAZIO
        if fim and fim != DISPLAY_VAZIO and fim != inicio:
            periodo = f"{inicio} -> {fim}"
        ctk.CTkLabel(
            info,
            text=periodo,
            text_color=COR_TEXTO,
        ).pack(anchor="w")
        obs_extra = (registro.get("observacao_extra") or "").strip()
        cor_marcador = ajustar_cor_marcador(registro.get("observacao_cor"))
        if obs_extra:
            obs_frame = ctk.CTkFrame(
                info,
                fg_color=cor_marcador or "#EEF2FA",
                corner_radius=12,
            )
            obs_frame.pack(fill="x", pady=(6, 0))
            ctk.CTkLabel(
                obs_frame,
                text=f"Obs.: {obs_extra}",
                text_color=COR_TEXTO,
                font=FONT_TEXTO,
            ).pack(anchor="w", padx=10, pady=8)
        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="e", padx=12, pady=(0, 8))
        ctk.CTkButton(
            botoes,
            text="Editar",
            width=80,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda item=registro: preparar_edicao_folga(item),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            width=90,
            fg_color="#FFFFFF",
            text_color=COR_VERMELHA,
            border_color=COR_VERMELHA,
            border_width=1,
            hover_color="#FCE4E4",
            command=lambda item=registro: handle_excluir_folga(item),
        ).pack(side="left", padx=4)

# ---------- COLABORADORES ----------
ctk.CTkLabel(
    aba_colaboradores,
    text="COLABORADORES",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_colaboradores,
    text="Cadastre motoristas e ajudantes ativos.",
    text_color=COR_TEXTO,
).pack(anchor="w", padx=10, pady=(0, 10))

frame_colab_form = ctk.CTkFrame(aba_colaboradores, fg_color="#FFFFFF", corner_radius=12)
frame_colab_form.pack(fill="x", padx=10, pady=(5, 12))

ctk.CTkLabel(
    frame_colab_form,
    text="Nome completo",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
entry_nome = ctk.CTkEntry(
    frame_colab_form, placeholder_text="Digite o nome completo", width=360
)
entry_nome.grid(row=1, column=0, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_colab_form,
    text="Função",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
combo_funcao = ctk.CTkOptionMenu(frame_colab_form, values=FUNCOES)
combo_funcao.set(FUNCOES[0])
combo_funcao.grid(row=1, column=1, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_colab_form,
    text="Observação",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=20, pady=(20, 0), sticky="w")
entry_obs_colab = ctk.CTkEntry(
    frame_colab_form, placeholder_text="Ex.: Em treinamento", width=300
)
entry_obs_colab.grid(row=1, column=2, padx=20, pady=5, sticky="w")

frame_foto_colab = ctk.CTkFrame(frame_colab_form, fg_color="#F7F8FB", corner_radius=12)
frame_foto_colab.grid(row=0, column=3, rowspan=2, padx=20, pady=10, sticky="n")
label_colaborador_foto_preview = ctk.CTkLabel(
    frame_foto_colab,
    text="Sem foto",
    width=150,
    height=140,
    fg_color="#FFFFFF",
    text_color=COR_CINZA,
    corner_radius=12,
    anchor="center",
    justify="center",
)
label_colaborador_foto_preview.pack(padx=12, pady=(12, 6))
btn_escolher_foto = ctk.CTkButton(
    frame_foto_colab, text="Selecionar foto", command=lambda: selecionar_foto_colaborador()
)
estilizar_botao(btn_escolher_foto, "primary", small=True)
btn_escolher_foto.pack(fill="x", padx=12, pady=(0, 6))
btn_limpar_foto_colaborador = ctk.CTkButton(
    frame_foto_colab,
    text="Remover foto",
    state="disabled",
    command=lambda: limpar_foto_colaborador(),
)
estilizar_botao(btn_limpar_foto_colaborador, "danger-ghost", small=True)
btn_limpar_foto_colaborador.pack(fill="x", padx=12, pady=(0, 12))

btn_salvar_colab = ctk.CTkButton(
    frame_colab_form,
    text="Salvar colaborador",
    command=lambda: handle_salvar_colaborador(),
)
estilizar_botao(btn_salvar_colab, "primary")
btn_salvar_colab.grid(row=2, column=0, padx=20, pady=5, sticky="w")
btn_cancelar_colab = ctk.CTkButton(
    frame_colab_form,
    text="Cancelar edição",
    state="disabled",
    command=lambda: limpar_form_colaborador(),
)
estilizar_botao(btn_cancelar_colab, "danger-ghost")
btn_cancelar_colab.grid(row=2, column=1, padx=20, pady=5, sticky="w")

label_total_colaboradores = ctk.CTkLabel(
    aba_colaboradores,
    text="Colaboradores ativos (0)",
    text_color=COR_TEXTO,
    font=("Inter", 18, "bold"),
)
label_total_colaboradores.pack(anchor="w", padx=20, pady=(5, 0))

frame_busca_colaboradores = ctk.CTkFrame(aba_colaboradores, fg_color="transparent")
frame_busca_colaboradores.pack(fill="x", padx=20, pady=(4, 0))
ctk.CTkLabel(
    frame_busca_colaboradores,
    text="Pesquisar colaborador",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, sticky="w")
colab_busca_var = ctk.StringVar(value="")
entry_colab_busca = ctk.CTkEntry(
    frame_busca_colaboradores,
    textvariable=colab_busca_var,
    placeholder_text="Digite o nome",
    width=320,
)
entry_colab_busca.grid(row=1, column=0, pady=(4, 0), sticky="w")
estilizar_entry(entry_colab_busca)
colab_busca_var.trace_add(
    "write",
    lambda *_: debounce(
        "colab_busca",
        180,
        lambda: refresh_colaboradores_ui(atualizar_dependentes=False),
    ),
)
frame_busca_colaboradores.grid_columnconfigure(0, weight=1)

lista_colaboradores = ctk.CTkScrollableFrame(
    aba_colaboradores, fg_color="#FFFFFF", corner_radius=12, height=320
)
lista_colaboradores.pack(fill="both", expand=True, padx=20, pady=(5, 20))


def atualizar_preview_foto_form():
    if label_colaborador_foto_preview is None:
        return
    caminho = None
    if colaborador_foto_origem_temp:
        caminho = Path(colaborador_foto_origem_temp)
    elif colaborador_foto_rel_atual:
        caminho = caminho_absoluto_foto(colaborador_foto_rel_atual)
    if caminho and caminho.exists():
        imagem = carregar_imagem_em_cache(caminho, (150, 150))
        if imagem:
            label_colaborador_foto_preview.configure(image=imagem, text="")
            label_colaborador_foto_preview.image = imagem
        else:
            label_colaborador_foto_preview.configure(image=None, text="Sem foto")
            label_colaborador_foto_preview.image = None
    else:
        label_colaborador_foto_preview.configure(image=None, text="Sem foto")
        label_colaborador_foto_preview.image = None
    if btn_limpar_foto_colaborador is not None:
        estado = (
            "normal" if (colaborador_foto_origem_temp or colaborador_foto_rel_atual) else "disabled"
        )
        btn_limpar_foto_colaborador.configure(state=estado)


def selecionar_foto_colaborador():
    global colaborador_foto_origem_temp
    arquivo = filedialog.askopenfilename(
        title="Selecione a foto do colaborador",
        filetypes=[
            ("Imagens", ";".join(f"*{ext}" for ext in FOTO_EXTENSOES_SUPORTADAS)),
            ("Todos os arquivos", "*.*"),
        ],
    )
    if not arquivo:
        return
    colaborador_foto_origem_temp = arquivo
    atualizar_preview_foto_form()


def limpar_foto_colaborador():
    global colaborador_foto_origem_temp, colaborador_foto_rel_atual, colaborador_foto_rel_original
    colaborador_foto_origem_temp = None
    colaborador_foto_rel_atual = None
    colaborador_foto_rel_original = None
    atualizar_preview_foto_form()


def handle_salvar_colaborador():
    global colaborador_em_edicao_id, colaborador_foto_rel_atual, colaborador_foto_rel_original, colaborador_foto_origem_temp
    nome = entry_nome.get().strip()
    funcao = combo_funcao.get().strip()
    observacao = entry_obs_colab.get().strip()
    if not nome:
        messagebox.showerror("Campos obrigatórios", "Informe o nome do colaborador.")
        return
    if funcao == FUNCOES[0]:
        messagebox.showerror("Campos obrigatórios", "Selecione a função.")
        return
    foto_relativa = colaborador_foto_rel_atual
    if colaborador_foto_origem_temp:
        foto_relativa = salvar_foto_colaborador_local(colaborador_foto_origem_temp)
        if not foto_relativa:
            messagebox.showerror(
                "Foto", "Não foi possível salvar a foto selecionada. Tente novamente."
            )
            return
    foto_antiga = None
    try:
        if colaborador_em_edicao_id:
            atualizar_colaborador(colaborador_em_edicao_id, nome, funcao, observacao, foto_relativa)
            if colaborador_foto_rel_original and colaborador_foto_rel_original != foto_relativa:
                foto_antiga = colaborador_foto_rel_original
            colaborador_foto_rel_original = foto_relativa
            messagebox.showinfo("Atualizado", "Colaborador atualizado com sucesso.")
        else:
            add_colaborador(nome, funcao, observacao, foto_relativa)
            messagebox.showinfo("Sucesso", "Colaborador cadastrado com sucesso.")
            colaborador_foto_rel_original = foto_relativa
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", f"Não foi possível salvar o colaborador.\n{exc}")
        return
    if foto_antiga:
        remover_arquivo_foto(foto_antiga)
    colaborador_foto_origem_temp = None
    colaborador_foto_rel_atual = foto_relativa
    atualizar_preview_foto_form()
    criar_backup_automatico()
    limpar_form_colaborador()
    refresh_colaboradores_ui()
    recarregar_disponibilidade()


def limpar_form_colaborador():
    global colaborador_em_edicao_id, colaborador_foto_rel_atual, colaborador_foto_rel_original, colaborador_foto_origem_temp
    colaborador_em_edicao_id = None
    colaborador_foto_rel_atual = None
    colaborador_foto_rel_original = None
    colaborador_foto_origem_temp = None
    entry_nome.delete(0, "end")
    combo_funcao.set(FUNCOES[0])
    entry_obs_colab.delete(0, "end")
    btn_salvar_colab.configure(text="Salvar colaborador")
    btn_cancelar_colab.configure(state="disabled")
    atualizar_preview_foto_form()


def preparar_edicao_colaborador(colaborador: dict):
    global colaborador_em_edicao_id, colaborador_foto_rel_atual, colaborador_foto_rel_original, colaborador_foto_origem_temp
    colaborador_em_edicao_id = colaborador["id"]
    entry_nome.delete(0, "end")
    entry_nome.insert(0, colaborador["nome"])
    combo_funcao.set(colaborador["funcao"])
    entry_obs_colab.delete(0, "end")
    entry_obs_colab.insert(0, colaborador.get("observacao") or "")
    colaborador_foto_rel_original = colaborador.get("foto")
    colaborador_foto_rel_atual = colaborador_foto_rel_original
    colaborador_foto_origem_temp = None
    atualizar_preview_foto_form()
    btn_salvar_colab.configure(text="Atualizar colaborador")
    btn_cancelar_colab.configure(state="normal")


def handle_excluir_colaborador(colaborador: dict):
    if not messagebox.askyesno(
        "Excluir colaborador",
        f"Deseja remover {colaborador['nome']} do cadastro?",
    ):
        return
    try:
        desativar_colaborador(colaborador["id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    remover_arquivo_foto(colaborador.get("foto"))
    refresh_colaboradores_ui()
    recarregar_disponibilidade()


def refresh_colaboradores_ui(atualizar_dependentes: bool = True):
    global colaborador_option_map
    colaboradores = listar_colaboradores()
    colaborador_option_map = {}
    for colaborador in colaboradores:
        display_value = (
            f"#{colaborador['id']} - {colaborador['nome']} ({colaborador['funcao']})"
        )
        colaborador_option_map[display_value] = colaborador["id"]

    filtro = colab_busca_var.get().strip().lower()
    colaboradores_filtrados = colaboradores
    if filtro:
        colaboradores_filtrados = [
            c
            for c in colaboradores
            if filtro in (c.get("nome") or "").lower()
        ]
        label_total_colaboradores.configure(
            text=f"Colaboradores ativos ({len(colaboradores_filtrados)} de {len(colaboradores)})"
        )
    else:
        label_total_colaboradores.configure(
            text=f"Colaboradores ativos ({len(colaboradores)})"
        )
    for widget in lista_colaboradores.winfo_children():
        widget.destroy()
    if not colaboradores:
        mostrar_msg_lista(lista_colaboradores)
    elif not colaboradores_filtrados:
        mostrar_msg_lista(lista_colaboradores, "Nenhum colaborador encontrado.")
    else:
        for colaborador in colaboradores_filtrados:
            linha = ctk.CTkFrame(lista_colaboradores, fg_color="#F5F5F5", corner_radius=8)
            linha.pack(fill="x", padx=10, pady=5)
            conteudo = ctk.CTkFrame(linha, fg_color="transparent")
            conteudo.pack(fill="x", padx=10, pady=6)
            conteudo.columnconfigure(1, weight=1)
            foto_label = ctk.CTkLabel(
                conteudo,
                text="Sem foto",
                width=80,
                height=80,
                fg_color="#FFFFFF",
                text_color=COR_CINZA,
                corner_radius=12,
                anchor="center",
                justify="center",
            )
            foto_label.grid(row=0, column=0, rowspan=3, padx=(0, 12), pady=6, sticky="n")
            imagem = None
            if colaborador.get("foto"):
                caminho = caminho_absoluto_foto(colaborador.get("foto"))
                if caminho:
                    imagem = carregar_imagem_em_cache(caminho, (70, 70))
            if imagem:
                foto_label.configure(image=imagem, text="")
                foto_label.image =imagem
            funcao_label = colaborador.get("funcao") or ""
            ctk.CTkLabel(
                conteudo,
                text=f"{colaborador['nome']} - {funcao_label}",
                text_color=COR_TEXTO,
                font=("Inter", 14, "bold"),
            ).grid(row=0, column=1, sticky="w", pady=(4, 2))
            if colaborador.get("observacao"):
                ctk.CTkLabel(
                    conteudo,
                    text=f"Obs.: {colaborador['observacao']}",
                    text_color=COR_TEXTO,
                    font=FONT_TEXTO,
                ).grid(row=1, column=1, sticky="w")
            botoes = ctk.CTkFrame(conteudo, fg_color="transparent")
            botoes.grid(row=2, column=1, sticky="w", pady=(4, 2))
            ctk.CTkButton(
                botoes,
                text="Editar",
                width=80,
                height=30,
                fg_color=COR_AZUL,
                hover_color=COR_VERMELHA,
                command=lambda c=colaborador: preparar_edicao_colaborador(c),
            ).pack(side="left", padx=4)
            ctk.CTkButton(
                botoes,
                text="Excluir",
                width=90,
                height=30,
                fg_color="#FFFFFF",
                text_color=COR_VERMELHA,
                border_color=COR_VERMELHA,
                border_width=1,
                hover_color="#FCE4E4",
                command=lambda c=colaborador: handle_excluir_colaborador(c),
            ).pack(side="left", padx=4)
    if atualizar_dependentes:
        atualizar_dropdown_folga()
        refresh_ferias_colaboradores(colaboradores)

# ---------- FÉRIAS ----------
ctk.CTkLabel(
    aba_ferias,
    text="FÉRIAS DOS COLABORADORES",
    text_color=COR_AZUL,
    font=FONT_TITULO,
).pack(anchor="w", padx=10, pady=(10, 0))

ctk.CTkLabel(
    aba_ferias,
    text="Registre períodos de férias e mantenha a escala atualizada.",
    text_color=COR_TEXTO,
).pack(anchor="w", padx=10, pady=(0, 10))

frame_ferias_form = ctk.CTkFrame(aba_ferias, fg_color="#FFFFFF", corner_radius=12)
frame_ferias_form.pack(fill="x", padx=10, pady=(5, 12))

ferias_colaborador_map: dict[str, int] = {}

ctk.CTkLabel(
    frame_ferias_form,
    text="Colaborador",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")
combo_ferias_colaborador = ctk.CTkOptionMenu(
    frame_ferias_form, values=["Cadastre um colaborador"], state="disabled", width=260
)
combo_ferias_colaborador.grid(row=1, column=0, padx=20, pady=5, sticky="w")
preview_ferias_colaborador_label = ctk.CTkLabel(
    frame_ferias_form,
    text="Sem foto",
    width=110,
    height=110,
    fg_color="#F3F5FA",
    text_color=COR_CINZA,
    corner_radius=55,
    anchor="center",
)
preview_ferias_colaborador_label.grid(row=0, column=4, rowspan=2, padx=20, pady=(15, 10), sticky="n")


def atualizar_preview_ferias_combo(selecionado: str | None = None):
    atualizar_preview_combo_colaborador(
        combo_ferias_colaborador,
        ferias_colaborador_map,
        preview_ferias_colaborador_label,
        size=(110, 110),
        selecionado=selecionado,
        circular=True,
    )


combo_ferias_colaborador.configure(
    command=lambda valor: atualizar_preview_ferias_combo(valor)
)

ctk.CTkLabel(
    frame_ferias_form,
    text="Data início (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=1, padx=20, pady=(20, 0), sticky="w")
ferias_inicio_var = ctk.StringVar(value=date.today().strftime("%d/%m/%Y"))
entry_ferias_inicio = ctk.CTkEntry(
    frame_ferias_form, textvariable=ferias_inicio_var, width=160
)
entry_ferias_inicio.grid(row=1, column=1, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_ferias_form,
    text="Data fim (DD/MM/AAAA)",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=2, padx=20, pady=(20, 0), sticky="w")
ferias_fim_var = ctk.StringVar(value=date.today().strftime("%d/%m/%Y"))
entry_ferias_fim = ctk.CTkEntry(
    frame_ferias_form, textvariable=ferias_fim_var, width=160
)
entry_ferias_fim.grid(row=1, column=2, padx=20, pady=5, sticky="w")

ctk.CTkLabel(
    frame_ferias_form,
    text="Observação",
    text_color=COR_TEXTO,
    font=FONT_SUBTITULO,
).grid(row=0, column=3, padx=20, pady=(20, 0), sticky="w")
entry_ferias_obs = ctk.CTkEntry(
    frame_ferias_form, placeholder_text="Opcional", width=240
)
entry_ferias_obs.grid(row=1, column=3, padx=20, pady=5, sticky="w")

btn_salvar_ferias = ctk.CTkButton(
    frame_ferias_form,
    text="Salvar férias",
    command=lambda: handle_salvar_ferias(),
)
estilizar_botao(btn_salvar_ferias, "primary")
btn_salvar_ferias.grid(row=2, column=3, padx=20, pady=10, sticky="e")

btn_cancelar_ferias = ctk.CTkButton(
    frame_ferias_form,
    text="Cancelar edição",
    state="disabled",
    command=lambda: limpar_form_ferias(),
)
estilizar_botao(btn_cancelar_ferias, "danger-ghost")
btn_cancelar_ferias.grid(row=2, column=2, padx=20, pady=10, sticky="e")

lista_ferias = ctk.CTkScrollableFrame(
    aba_ferias, fg_color="#FFFFFF", corner_radius=12, height=320
)
lista_ferias.pack(fill="both", expand=True, padx=10, pady=(5, 20))
estilizar_scrollable(lista_ferias)


def obter_ferias_colaborador_id(selecionado: str | None = None) -> int | None:
    chave = selecionado if selecionado is not None else combo_ferias_colaborador.get()
    return ferias_colaborador_map.get(chave)


def limpar_form_ferias():
    global ferias_em_edicao_id
    ferias_em_edicao_id = None
    ferias_inicio_var.set(date.today().strftime("%d/%m/%Y"))
    ferias_fim_var.set(date.today().strftime("%d/%m/%Y"))
    entry_ferias_obs.delete(0, "end")
    valores = combo_ferias_colaborador.cget("values")
    if valores:
        combo_ferias_colaborador.set(valores[0])
        atualizar_preview_ferias_combo(valores[0])
    btn_salvar_ferias.configure(text="Salvar férias")
    btn_cancelar_ferias.configure(state="disabled")


def preparar_edicao_ferias(registro: dict):
    global ferias_em_edicao_id
    ferias_em_edicao_id = registro["id"]
    ferias_inicio_var.set(
        data_iso_para_br_entrada(registro.get("data_inicio") or date.today().isoformat())
    )
    ferias_fim_var.set(
        data_iso_para_br_entrada(registro.get("data_fim") or date.today().isoformat())
    )
    entry_ferias_obs.delete(0, "end")
    entry_ferias_obs.insert(0, registro.get("observacao") or "")
    for display, cid in ferias_colaborador_map.items():
        if cid == registro.get("colaborador_id"):
            combo_ferias_colaborador.set(display)
            atualizar_preview_ferias_combo(display)
            break
    btn_salvar_ferias.configure(text="Atualizar férias")
    btn_cancelar_ferias.configure(state="normal")


def handle_salvar_ferias():
    global ferias_em_edicao_id
    colaborador_id = obter_ferias_colaborador_id()
    if not colaborador_id:
        messagebox.showerror("Campos obrigatórios", "Selecione um colaborador.")
        return
    inicio_txt = ferias_inicio_var.get().strip()
    fim_txt = ferias_fim_var.get().strip()
    inicio_iso = validar_data(inicio_txt)
    if not inicio_iso:
        messagebox.showerror("Data inválida", "Informe a data de início no formato DD/MM/AAAA.")
        return
    fim_iso = validar_data(fim_txt) if fim_txt else None
    if not fim_iso:
        fim_iso = inicio_iso
    if fim_iso < inicio_iso:
        messagebox.showerror("Período inválido", "A data fim não pode ser anterior ao início.")
        return
    observacao = entry_ferias_obs.get().strip() or None
    try:
        if ferias_em_edicao_id:
            atualizar_ferias(ferias_em_edicao_id, colaborador_id, inicio_iso, fim_iso, observacao)
            messagebox.showinfo("Atualizado", "Período de férias atualizado.")
        else:
            adicionar_ferias(colaborador_id, inicio_iso, fim_iso, observacao)
            messagebox.showinfo("Salvo", "Férias registradas com sucesso.")
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao salvar", str(exc))
        return
    criar_backup_automatico()
    limpar_form_ferias()
    atualizar_lista_ferias()
    recarregar_disponibilidade()


def handle_excluir_ferias(registro: dict):
    if not messagebox.askyesno(
        "Excluir férias",
        f"Deseja remover o período de {registro.get('nome')}?",
    ):
        return
    try:
        remover_ferias(registro["id"])
    except sqlite3.Error as exc:
        messagebox.showerror("Erro ao excluir", str(exc))
        return
    atualizar_lista_ferias()
    limpar_form_ferias()
    recarregar_disponibilidade()


def atualizar_lista_ferias():
    for widget in lista_ferias.winfo_children():
        widget.destroy()
    registros = listar_ferias()
    if not registros:
        mostrar_msg_lista(lista_ferias, "Nenhum registro de férias.")
        return
    for registro in registros:
        bloco = ctk.CTkFrame(lista_ferias, fg_color="#F9FAFB", corner_radius=14)
        bloco.pack(fill="x", padx=10, pady=4)
        conteudo = ctk.CTkFrame(bloco, fg_color="transparent")
        conteudo.pack(fill="x", padx=12, pady=8)
        avatar = ctk.CTkLabel(
            conteudo,
            text="Sem foto",
            width=70,
            height=70,
            fg_color="#E7ECF5",
            text_color=COR_CINZA,
            corner_radius=999,
            anchor="center",
        )
        avatar.pack(side="left", padx=(0, 12))
        imagem_avatar = obter_imagem_colaborador_por_id(
            registro.get("colaborador_id"), size=(70, 70), circular=True
        )
        if imagem_avatar:
            avatar.configure(image=imagem_avatar, text="")
            avatar.image = imagem_avatar
        info = ctk.CTkFrame(conteudo, fg_color="transparent")
        info.pack(fill="both", expand=True)
        titulo = f"{registro['nome']} - {registro.get('funcao', '')}"
        ctk.CTkLabel(
            info,
            text=titulo,
            text_color=COR_AZUL,
            font=("Inter", 14, "bold"),
        ).pack(anchor="w", pady=(0, 2))
        periodo = f"{registro.get('data_inicio')} -> {registro.get('data_fim')}"
        ctk.CTkLabel(
            info,
            text=periodo,
            text_color=COR_TEXTO,
        ).pack(anchor="w")
        status_text = ""
        data_fim = registro.get("data_fim") or ""
        try:
            if data_fim:
                dt_fim = datetime.strptime(data_fim, "%Y-%m-%d").date()
                if dt_fim < date.today():
                    status_text = "Ja finalizado"
        except ValueError:
            status_text = ""
        if status_text:
            ctk.CTkLabel(
                info,
                text=status_text,
                text_color=COR_VERMELHA,
                font=("Inter", 11, "italic"),
            ).pack(anchor="w")
        obs = registro.get("observacao")
        if obs:
            ctk.CTkLabel(
                info,
                text=obs,
                text_color=COR_TEXTO,
            ).pack(anchor="w", pady=(4, 0))
        botoes = ctk.CTkFrame(bloco, fg_color="transparent")
        botoes.pack(anchor="e", padx=12, pady=(0, 8))
        ctk.CTkButton(
            botoes,
            text="Editar",
            width=80,
            fg_color=COR_AZUL,
            hover_color=COR_VERMELHA,
            command=lambda item=registro: preparar_edicao_ferias(item),
        ).pack(side="left", padx=4)
        ctk.CTkButton(
            botoes,
            text="Excluir",
            width=90,
            fg_color="#FFFFFF",
            text_color=COR_VERMELHA,
            border_color=COR_VERMELHA,
            border_width=1,
            hover_color="#FCE4E4",
            command=lambda item=registro: handle_excluir_ferias(item),
        ).pack(side="left", padx=4)


def refresh_ferias_colaboradores(colaboradores: list[dict] | None = None):
    if colaboradores is None:
        colaboradores = listar_colaboradores()
    ferias_colaborador_map.clear()
    if not colaboradores:
        combo_ferias_colaborador.configure(
            values=["Cadastre um colaborador"], state="disabled"
        )
        combo_ferias_colaborador.set("Cadastre um colaborador")
        atualizar_preview_ferias_combo(combo_ferias_colaborador.get())
        return

    opcoes = []
    for colaborador in colaboradores:
        funcao = colaborador.get("funcao") or ""
        display = f"#{colaborador['id']} - {colaborador['nome']} ({funcao})"
        opcoes.append(display)
        ferias_colaborador_map[display] = colaborador["id"]

    combo_ferias_colaborador.configure(values=opcoes, state="normal")
    combo_ferias_colaborador.set(opcoes[0])
    atualizar_preview_ferias_combo(opcoes[0])
    limpar_form_ferias()
# ---------- DADOS INICIAIS ----------
refresh_colaboradores_ui()
refresh_caminhoes_ui()
aplicar_data_carregamentos()
limpar_form_carregamento()
refresh_oficinas_ui()
exibir_folgas_para_data(date.today().isoformat())
limpar_form_folga()
atualizar_lista_ferias()
limpar_form_ferias()

rodape = ctk.CTkLabel(
    app,
    text="© JR Ferragens & Madeiras | Sistema JR Escala",
    text_color=COR_TEXTO,
    font=("Inter", 10),
)
rodape.pack(side="bottom", pady=5)

app.mainloop()






  







